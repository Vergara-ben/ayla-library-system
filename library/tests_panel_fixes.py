"""Fixes from the panel review: loan extensions, and what a received copy is recorded as."""

import io
import json
from datetime import timedelta
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.utils import timezone

from .auth_utils import hash_password
from .models import (
    Book, BorrowingRule, DueDateExtension, InventoryRecord, Patron, Transaction, User,
)
from .tests import COMPLETE_BOOK, _with_genre, _with_genre_items
from .views import isbn_error


def _staff(role='Staff', modules='transactions,books,inventory,donations'):
    return User.objects.create(fullname='Desk', email='desk-%s@example.invalid' % role.lower(),
                               password_hash=hash_password('DeskTest123'),
                               role=role, account_status='Active', modules=modules)


def _client(**session_values):
    c = Client()
    s = c.session
    for key, value in session_values.items():
        s[key] = value
    s.save()
    return c


class ExtensionTests(TestCase):

    def setUp(self):
        rule = BorrowingRule.current()
        rule.loan_period_days = 2
        rule.save()
        self.today = timezone.localdate()
        self.patron = Patron.objects.create(first_name='Ana', last_name='Cruz',
                                            email='ana@example.invalid', account_status='Active')
        self.book = Book.objects.create(title='Noli', author='Rizal', status='Borrowed')
        self.loan = Transaction.objects.create(
            book=self.book, patron=self.patron, transaction_type='Borrow',
            due_date=self.today + timedelta(days=2))
        self.patron_client = _client(patron_id=self.patron.patron_id)
        self.staff = _staff()
        self.staff_client = _client(admin_id=self.staff.admin_id, admin_role='Staff')

    def _request(self, reason='Still reading'):
        return self.patron_client.post('/patron/request-extension/',
                                       {'transaction_id': self.loan.transaction_id,
                                        'reason': reason}).json()

    def _respond(self, extension, action='approve'):
        return self.staff_client.post('/admin-portal/respond-to-extension/',
                                      {'extension_id': extension.extension_id,
                                       'action': action},
                                      HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()

    def test_borrowed_today_and_extended_today_moves_the_due_date(self):
        # The panel's case: due in 2 days, extended the same day, still due in 2 days.
        self.assertTrue(self._request()['success'])
        ext = DueDateExtension.objects.get(transaction=self.loan)
        self.assertEqual(ext.requested_due_date, self.today + timedelta(days=4))
        self.assertTrue(self._respond(ext)['success'])
        self.loan.refresh_from_db()
        self.assertEqual(self.loan.due_date, self.today + timedelta(days=4))

    def test_an_overdue_loan_cannot_be_extended(self):
        self.loan.due_date = self.today - timedelta(days=1)
        self.loan.overdue_flag = True
        self.loan.save()
        result = self._request()
        self.assertFalse(result['success'])
        self.assertIn('overdue', result['error'])
        self.assertFalse(DueDateExtension.objects.exists())

    def test_a_request_made_before_the_fix_still_adds_a_period_when_approved(self):
        stale = DueDateExtension.objects.create(
            transaction=self.loan, requested_by_patron=True,
            previous_due_date=self.loan.due_date, requested_due_date=self.loan.due_date)
        self.assertTrue(self._respond(stale)['success'])
        self.loan.refresh_from_db()
        self.assertEqual(self.loan.due_date, self.today + timedelta(days=4))



class ReceivingConditionTests(TestCase):

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin').admin_id)

    def _receive(self, items, source='Purchase', **extra):
        data = {'source': source, 'supplier': 'Test Supplier', 'items': json.dumps(_with_genre_items(items))}
        data.update(extra)
        return self.client.post('/admin-portal/inventory/receive/', data, follow=True)

    def test_condition_defaults_to_good(self):
        self._receive([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 1}])
        self.assertEqual(InventoryRecord.objects.get().condition, 'Good')

    def test_lost_or_withdrawn_is_not_a_way_to_arrive(self):
        for condition in ('Lost', 'Withdrawn'):
            self._receive([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 1,
                            'condition': condition}])
        self.assertFalse(Book.objects.filter(title='Ibong Adarna').exists())

    def test_minor_damage_arrives_on_the_copy_and_the_book(self):
        self._receive([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 2, 'condition': 'Minor Damage'}])
        books = Book.objects.filter(title='Ibong Adarna')
        self.assertEqual(set(books.values_list('condition', flat=True)), {'Minor Damage'})
        self.assertEqual(set(InventoryRecord.objects.filter(book__in=books)
                             .values_list('condition', flat=True)), {'Minor Damage'})

    def test_minor_damage_copies_are_still_on_hand(self):
        self._receive([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 1, 'condition': 'Minor Damage'}])
        self.assertTrue(InventoryRecord.objects.get().counts_as_held)

    def test_a_mistyped_isbn_is_refused(self):
        page = self._receive([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 1,
                               'condition': 'Good', 'isbn': '978-971-23-1234-0'}])
        self.assertContains(page, 'wrong check digit')
        self.assertFalse(Book.objects.filter(title='Ibong Adarna').exists())

    def test_the_form_offers_only_arrival_conditions(self):
        page = self.client.get('/admin-portal/inventory/').content.decode()
        picker = page.split('id="rcvCondition"')[1].split('</select>')[0]
        for value in ('Good', 'Minor Damage', 'Damaged', 'Beyond Repair'):
            self.assertIn(value, picker)
        self.assertNotIn('Lost', picker)
        self.assertNotIn('Withdrawn', picker)


class DonationImportConditionTests(TestCase):

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin').admin_id)

    def _import(self, rows):
        import openpyxl
        wb = openpyxl.Workbook()
        for row in _with_genre([['Donor', 'Title', 'Author', 'Condition']] + list(rows)):
            wb.active.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile('d.xlsx', buf.getvalue(), content_type=(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))
        return self.client.post('/admin-portal/import-donations/', {'excel_file': upload, 'new_genres': '1'},
                                HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()

    def test_imported_condition_is_kept(self):
        self._import([['Barangay', 'Florante at Laura', 'Balagtas', 'Damaged']])
        self.assertEqual(InventoryRecord.objects.get().condition, 'Damaged')
        self.assertEqual(Book.objects.get(title='Florante at Laura').condition, 'Damaged')

    def test_a_blank_condition_is_good_and_a_wrong_one_is_reported(self):
        result = self._import([['Barangay', 'Florante at Laura', 'Balagtas', ''],
                               ['Barangay', 'Noli Me Tangere', 'Rizal', 'Lost']])
        self.assertIn('condition must be Good, Minor Damage, Damaged or Beyond Repair (got "Lost")', result['message'])
        self.assertEqual(InventoryRecord.objects.get().condition, 'Good')


class IsbnTests(TestCase):

    def test_valid_isbns(self):
        for isbn in ('978-0-306-40615-7', '0306406152', '0-8044-2957-X', '9780306406157'):
            self.assertIsNone(isbn_error(isbn), isbn)

    def test_invalid_isbns(self):
        self.assertIn('check digit', isbn_error('9780306406158'))
        self.assertIn('10 or 13 digits', isbn_error('12345'))
        self.assertIn('10 or 13 digits', isbn_error('ABCDEFGHIJ'))


class BeaconMinorTests(TestCase):
    UUID = 'FDA50693-A4E2-4FB1-AFCF-C6EB07647825'

    def setUp(self):
        from .models import FloorPlan
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')

    def _place(self, **extra):
        data = {'floor_plan_id': self.plan.floor_plan_id, 'beacon_uuid': self.UUID,
                'map_x': 10, 'map_y': 10}
        data.update(extra)
        return self.client.post('/admin-portal/add-beacon/', data).json()

    def test_minor_and_label_are_filled_in(self):
        first = self._place()
        second = self._place()
        self.assertEqual((first['beacon']['minor'], second['beacon']['minor']), (1, 2))
        self.assertEqual(second['beacon']['label'], 'Beacon 2')
        self.assertEqual(second['beacon']['major'], 1)

    def test_a_freed_minor_is_offered_again(self):
        self._place(minor=1)
        self._place(minor=3)
        r = self.client.get('/admin-portal/next-beacon-minor/', {'uuid': self.UUID.lower(), 'major': 1})
        self.assertEqual(r.json()['minor'], 2)

    def test_a_duplicate_minor_is_refused(self):
        self._place(minor=5)
        result = self._place(minor=5, label='Other')
        self.assertFalse(result['success'])
        self.assertIn('already uses these identifiers', result['error'])


class LabelTests(TestCase):

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.shelved = Book.objects.create(title='Shelved', author='A', shelf_level=self.level, qr_code='q1')
        self.loose = Book.objects.create(title='Loose', author='B', qr_code='q2')
        self.client = _client(admin_id=_staff(modules='books').admin_id, admin_role='Staff')

    def _print(self, *books):
        return self.client.get('/admin-portal/book-qr-labels/',
                               {'ids': ','.join(str(b.book_id) for b in books)})

    def test_shelved_and_unshelved_books_are_printed_and_marked(self):
        # Labelled as they arrive: no shelf yet only leaves out the shelf line.
        response = self._print(self.shelved, self.loose)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertEqual(response['X-Labels-Skipped'], '0')
        self.shelved.refresh_from_db()
        self.loose.refresh_from_db()
        self.assertIsNotNone(self.shelved.label_printed_at)
        self.assertIsNotNone(self.loose.label_printed_at)

    def test_a_sheet_of_only_unshelved_books_prints(self):
        self.assertEqual(self._print(self.loose)['Content-Type'], 'application/pdf')

    def test_books_needing_a_label_are_counted_and_listed(self):
        page = self.client.get('/library-staff/books/')
        self.assertEqual(page.context['unlabelled_count'], 1)
        self.assertEqual(page.context['unshelved_count'], 1)
        ids = self.client.get('/library-staff/books/', {'label': 'missing', 'ids': 'all'}).json()['ids']
        self.assertEqual(ids, [self.shelved.book_id])
        self._print(self.shelved)
        self.assertEqual(self.client.get('/library-staff/books/').context['unlabelled_count'], 0)


class RegistrationFixTests(TestCase):

    def setUp(self):
        from django.core import mail  # noqa: F401
        self.applicant = Patron.objects.create(
            first_name='Ana', last_name='Cruz', email='ana@example.invalid', account_status='Pending',
            registration_channel='Online', credential_document='credentials/x.png', otp_verified=True)
        self.client = _client(admin_id=_staff(modules='patrons').admin_id, admin_role='Staff')

    def _reject(self, **data):
        return self.client.post('/admin-portal/reject-patron/%d/' % self.applicant.patron_id, data)

    def test_a_reason_is_required(self):
        self._reject(mode='reject')
        self.applicant.refresh_from_db()
        self.assertIsNone(self.applicant.archived_at)

    def test_ask_to_fix_keeps_it_pending_and_the_link_updates_it(self):
        self._reject(mode='fix', reason='The ID has expired')
        self.applicant.refresh_from_db()
        self.assertEqual(self.applicant.account_status, 'Pending')
        token = self.applicant.fix_token
        self.assertTrue(token)
        page = Client().get('/patron/registration/fix/%s/' % token)
        self.assertContains(page, 'The ID has expired')
        Client().post('/patron/registration/fix/%s/' % token,
                      {'first_name': 'Ana', 'last_name': 'Santos', 'patron_type': 'Student'})
        self.applicant.refresh_from_db()
        self.assertEqual(self.applicant.last_name, 'Santos')
        self.assertIsNone(self.applicant.fix_token)
        self.assertIsNotNone(self.applicant.resubmitted_at)
        self.assertEqual(Client().get('/patron/registration/fix/%s/' % token).status_code, 404)

    def test_reject_completely_with_a_reason(self):
        self._reject(mode='reject', reason='This is not a valid ID')
        self.assertFalse(Patron.objects.filter(patron_id=self.applicant.patron_id).exists())


class AlertTests(TestCase):

    def _alerts(self, user):
        c = _client(admin_id=user.admin_id, admin_role=user.role)
        return {i['key']: i for i in c.get('/portal/alerts/').json()['items']}

    def test_staff_get_alerts_for_their_modules_only(self):
        desk = _staff(modules='transactions')
        items = self._alerts(desk)
        self.assertEqual(set(items), {'extensions', 'overdue', 'loan_on_shelf'})

    def test_a_new_registration_is_reported_but_not_one_waiting_on_the_applicant(self):
        admin = _staff(role='Admin', modules='')
        Patron.objects.create(first_name='Ana', last_name='Cruz', email='a@example.invalid',
                              account_status='Pending', otp_verified=True)
        Patron.objects.create(first_name='Ben', last_name='Lim', email='b@example.invalid',
                              account_status='Pending', otp_verified=True, fix_token='t')
        item = self._alerts(admin)['registrations']
        self.assertEqual(item['count'], 1)
        self.assertIn('Ana', item['text'])

    def test_signed_out_gets_nothing(self):
        self.assertEqual(Client().get('/portal/alerts/').status_code, 403)


class DonorTests(TestCase):

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')

    def _give(self, donor, day, title):
        return self.client.post('/admin-portal/inventory/receive/', {
            'source': 'Donation', 'donor_name': donor, 'donated_date': day,
            'items': json.dumps(_with_genre_items([{'title': title, 'author': 'Anon', 'quantity': 2}]))})

    def test_a_donor_is_saved_once_and_matched_however_it_is_typed(self):
        from .models import Donor
        self._give('Brgy. Sala Council', '2026-06-01', 'Noli')
        self._give('  brgy. sala   council ', '2026-09-01', 'Fili')
        self.assertEqual(Donor.objects.count(), 1)
        page = self.client.get('/admin-portal/donors/')
        donor = page.context['donors'][0]
        self.assertEqual((donor.gifts, donor.titles, donor.copies), (2, 2, 4))
        self.assertEqual(donor.every_days, 92)

    def test_donor_details_can_be_saved_and_history_shown(self):
        from .models import Donor
        self._give('Rotary Club', '2026-09-01', 'Noli')
        donor = Donor.objects.get()
        self.client.post('/admin-portal/donors/', {'donor_id': donor.donor_id, 'name': 'Rotary Club',
                                                   'donor_type': 'Organization', 'contact_number': '0917'})
        donor.refresh_from_db()
        self.assertEqual((donor.donor_type, donor.contact_number), ('Organization', '0917'))
        page = self.client.get('/admin-portal/donors/', {'d': donor.donor_id})
        self.assertContains(page, 'Noli')

    def test_the_donors_report(self):
        from datetime import date
        from .reports import build_report
        self._give('Rotary Club', '2026-09-01', 'Noli')
        report = build_report('donors', date(2026, 1, 1), date(2026, 12, 31))
        self.assertEqual(report['rows'][0][0], 'Rotary Club')
        self.assertEqual(report['rows'][0][5], 2)

    def test_receiving_suggests_past_donors(self):
        self._give('Rotary Club', '2026-09-01', 'Noli')
        page = self.client.get('/admin-portal/inventory/')
        self.assertContains(page, '<script id="donorNamesData" type="application/json">["Rotary Club"]</script>', html=False)


class IsbnLookupTests(TestCase):

    def setUp(self):
        self.client = _client(admin_id=_staff(modules='books').admin_id, admin_role='Staff')

    def _lookup(self, isbn):
        return self.client.get('/admin-portal/isbn-lookup/', {'isbn': isbn}).json()

    def test_a_book_we_hold_is_filled_from_the_catalogue(self):
        Book.objects.create(title='Noli Me Tangere', author='Rizal, Jose', ISBN='978-0-306-40615-7',
                            genre='Fiction', publication_year=1887)
        result = self._lookup('9780306406157')
        self.assertEqual(result['source'], 'catalogue')
        self.assertEqual(result['book']['title'], 'Noli Me Tangere')
        self.assertEqual(result['copies'], 1)

    def test_otherwise_open_library_is_asked(self):
        from unittest import mock
        found = {'title': 'Opticks', 'author': 'Isaac Newton', 'publication_year': '1979',
                 'cover_img_url': ''}
        with mock.patch('library.views._open_library', return_value=found) as asked:
            result = self._lookup('0-306-40615-2')
        asked.assert_called_once_with('0306406152')
        self.assertEqual(result['source'], 'openlibrary')
        self.assertEqual(result['book']['author'], 'Isaac Newton')

    def test_a_bad_isbn_is_not_looked_up(self):
        from unittest import mock
        with mock.patch('library.views._open_library') as asked:
            result = self._lookup('12345')
        asked.assert_not_called()
        self.assertFalse(result['success'])

    def test_the_add_form_shelves_the_book_and_skips_status(self):
        page = self.client.get('/library-staff/books/').content.decode()
        form = page.split('id="addBookForm"')[1].split('</form>')[0]
        self.assertIn('name="shelf_level"', form)
        self.assertIn('<input type="hidden" name="status" value="Available">', form)
        self.assertNotIn('cover_img_url', form)
        self.assertLess(form.index('id="addIsbn"'), form.index('id="addTitle"'))


class AuditTrailTests(TestCase):

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.here = Book.objects.create(title='Here', author='A', shelf_level=self.level)
        self.gone = Book.objects.create(title='Gone', author='B', shelf_level=self.level)
        self.out = Book.objects.create(title='Out', author='C', shelf_level=self.level, status='Borrowed')
        self.reading = Book.objects.create(title='Reading', author='D', shelf_level=self.level,
                                           status='Being Read')
        reader = Patron.objects.create(first_name='Ana', last_name='Cruz', email='a@example.invalid')
        Transaction.objects.create(book=self.out, patron=reader, transaction_type='Borrow',
                                   due_date=timezone.localdate())
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')

    def test_books_away_are_recorded_without_being_marked(self):
        r = self.client.post('/admin-portal/inventory/audit/file/', {
            'level': self.level.shelf_level_id,
            'found_ids': [self.here.book_id], 'missing_ids': [self.gone.book_id]}).json()
        self.assertTrue(r['success'], r)
        from .models import StockAuditLine
        results = {l.title: (l.result, l.status_then, l.detail) for l in StockAuditLine.objects.all()}
        self.assertEqual(results['Here'][0], 'Found')
        # "Then" is how the count left it, so the count's own change is not "changed since".
        self.assertEqual(results['Gone'][:2], ('Missing', 'Missing'))
        self.assertEqual(results['Out'][0], 'On loan')
        self.assertIn('Ana', results['Out'][2])
        self.assertEqual(results['Reading'][0], 'Being read')

        # Later the borrowed book comes back: the record shows then and now side by side.
        self.out.status = 'Available'
        self.out.save()
        detail = self.client.get('/admin-portal/inventory/audit/%d/' % r['audit_id']).json()
        out = next(l for l in detail['lines'] if l['title'] == 'Out')
        self.assertEqual((out['then']['status'], out['now']['status']), ('Borrowed', 'Available'))
        self.assertIn('status', out['changed'])


class ReviewFollowUpTests(TestCase):
    """Fixes from re-checking the panel work."""

    def test_an_accession_number_is_not_mistaken_for_a_bad_isbn(self):
        from .views import isbn_typo
        self.assertIsNone(isbn_typo('ACC-2026-0042'))
        self.assertIsNone(isbn_typo('12345'))
        self.assertIn('check digit', isbn_typo('9780306406158'))

    def test_books_can_be_marked_as_already_labelled(self):
        book = Book.objects.create(title='Old', author='A')
        c = _client(admin_id=_staff(modules='books').admin_id, admin_role='Staff')
        result = c.post('/admin-portal/mark-labelled/', {'ids': str(book.book_id)}).json()
        self.assertEqual(result['marked'], 1)
        book.refresh_from_db()
        self.assertIsNotNone(book.label_printed_at)

    def test_a_desk_sign_up_cannot_be_sent_a_fix_link(self):
        applicant = Patron.objects.create(first_name='Ben', last_name='Lim', account_status='Pending',
                                          registration_channel='On-site')
        c = _client(admin_id=_staff(modules='patrons').admin_id, admin_role='Staff')
        c.post('/admin-portal/reject-patron/%d/' % applicant.patron_id,
               {'mode': 'fix', 'reason': 'The ID has expired'})
        applicant.refresh_from_db()
        self.assertIsNone(applicant.fix_token)

    def test_a_resubmission_raises_an_alert_naming_that_applicant(self):
        admin = _staff(role='Admin', modules='')
        Patron.objects.create(first_name='Ana', last_name='Cruz', email='a@example.invalid',
                              account_status='Pending', otp_verified=True)
        back = Patron.objects.create(first_name='Ben', last_name='Lim', email='b@example.invalid',
                                     account_status='Pending', otp_verified=True,
                                     resubmitted_at=timezone.now())
        c = _client(admin_id=admin.admin_id, admin_role='Admin')
        item = {i['key']: i for i in c.get('/portal/alerts/').json()['items']}['registrations']
        self.assertIn('Ben', item['text'])
        self.assertIn('resubmitted', item['text'])
        self.assertEqual(item['stamp'], back.resubmitted_at.isoformat())

    def test_a_donor_can_be_added_before_their_first_gift(self):
        from .models import Donor
        c = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')
        c.post('/admin-portal/donors/', {'action': 'add', 'name': 'Rotary Club'})
        c.post('/admin-portal/donors/', {'action': 'add', 'name': 'rotary  club'})
        self.assertEqual(Donor.objects.count(), 1)

    def test_renaming_a_donated_copys_donor_moves_the_donation(self):
        from .models import Donor
        c = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')
        c.post('/admin-portal/inventory/receive/', {
            'source': 'Donation', 'donor_name': 'Rotary Clubb', 'donated_date': '2026-09-01',
            'items': json.dumps(_with_genre_items([{'title': 'Noli', 'author': 'Rizal', 'quantity': 1}]))})
        record = InventoryRecord.objects.get()
        from .views import _sync_donation_row
        record.donor_name = 'Rotary Club'
        record.save()
        _sync_donation_row(record)
        record.donation.refresh_from_db()
        self.assertEqual(record.donation.donor.name, 'Rotary Club')


class OneExtensionPerLoanTests(TestCase):

    def setUp(self):
        rule = BorrowingRule.current()
        rule.loan_period_days = 7
        rule.save()
        self.today = timezone.localdate()
        self.patron = Patron.objects.create(first_name='Ana', last_name='Cruz',
                                            email='ana@example.invalid', account_status='Active')
        self.loans = []
        for title in ('Noli', 'Fili'):
            book = Book.objects.create(title=title, author='Rizal', status='Borrowed')
            self.loans.append(Transaction.objects.create(
                book=book, patron=self.patron, transaction_type='Borrow',
                due_date=self.today + timedelta(days=7)))
        self.client = _client(patron_id=self.patron.patron_id)
        self.staff = _client(admin_id=_staff().admin_id, admin_role='Staff')

    def _ask(self, loan):
        return self.client.post('/patron/request-extension/', {'transaction_id': loan.transaction_id}).json()

    def test_each_loan_can_be_extended_once(self):
        self.assertTrue(self._ask(self.loans[0])['success'])
        ext = DueDateExtension.objects.get()
        self.staff.post('/admin-portal/respond-to-extension/',
                        {'extension_id': ext.extension_id, 'action': 'decline'},
                        HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        again = self._ask(self.loans[0])
        self.assertFalse(again['success'])
        self.assertIn('once', again['error'])

    def test_a_staff_due_date_change_does_not_use_it_up(self):
        self.staff.post('/admin-portal/adjust-due-date/', {
            'transaction_id': self.loans[0].transaction_id,
            'new_due_date': (self.today + timedelta(days=10)).isoformat()})
        self.assertTrue(self._ask(self.loans[0])['success'])

    def test_returning_one_book_leaves_the_other_extendable(self):
        self.loans[0].return_date = timezone.now()
        self.loans[0].save()
        self.assertFalse(self._ask(self.loans[0])['success'])
        self.assertTrue(self._ask(self.loans[1])['success'])

    def test_all_loans_in_one_tap(self):
        page = self.client.get('/patron/account/')
        self.assertContains(page, 'Request extension for all (2)')
        result = self.client.post('/patron/request-extension/all/').json()
        self.assertEqual(len(result['requested']), 2)
        self.assertEqual(DueDateExtension.objects.count(), 2)
        again = self.client.post('/patron/request-extension/all/').json()
        self.assertFalse(again['success'])
        self.assertEqual(len(again['skipped']), 2)

    def test_extend_all_skips_an_overdue_book(self):
        self.loans[1].due_date = self.today - timedelta(days=1)
        self.loans[1].save()
        result = self.client.post('/patron/request-extension/all/').json()
        self.assertEqual([r['title'] for r in result['requested']], ['Noli'])
        self.assertEqual(result['skipped'][0]['title'], 'Fili')


class ReceivingPickerTests(TestCase):

    def test_the_receiving_forms_offer_search_instead_of_a_long_dropdown(self):
        from .models import Donor
        Donor.for_name('Rotary Club')
        Book.objects.create(title='Noli', author='Rizal')
        c = _client(admin_id=_staff(modules='inventory').admin_id, admin_role='Staff')
        page = c.get('/library-staff/receiving/').content.decode()
        self.assertIn('id="rcvBookSearch"', page)
        self.assertIn('js/picker.js', page)
        self.assertIn('<script id="donorNamesData" type="application/json">["Rotary Club"]</script>', page)


class SimplerInventoryTests(TestCase):
    """Donations go Received -> Shelved, and shelving the book is the only step."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.client = _client(admin_id=_staff(role='Admin').admin_id)

    def _donate(self, quantity=1, **extra):
        data = {'source': 'Donation', 'donor_name': 'Rotary Club',
                'items': json.dumps(_with_genre_items([{'title': 'Noli', 'author': 'Rizal', 'quantity': quantity}]))}
        data.update(extra)
        self.client.post('/admin-portal/inventory/receive/', data)
        return list(Book.objects.filter(title='Noli').order_by('book_id'))

    def _shelve(self, *books):
        return self.client.post('/admin-portal/assign-books-to-level/', {
            'shelf_level_id': self.level.shelf_level_id,
            'book_ids': ','.join(str(b.book_id) for b in books)}).json()

    def test_a_donation_waits_hidden_and_shelving_it_releases_it(self):
        from .models import Donation
        book, = self._donate(processing_stage='Processing')
        self.assertEqual(book.status, 'Donated')
        self.assertEqual(InventoryRecord.objects.get().processing_stage, 'Received')
        self.assertEqual(Donation.objects.get().status, 'Received')

        self.assertTrue(self._shelve(book)['success'])
        book.refresh_from_db()
        self.assertEqual(book.status, 'Available')
        self.assertEqual(InventoryRecord.objects.get().processing_stage, 'Shelved')
        self.assertEqual(Donation.objects.get().status, 'Shelved')

    def test_a_donation_is_shelved_only_when_every_copy_is(self):
        from .models import Donation
        first, second = self._donate(quantity=2)
        self._shelve(first)
        self.assertEqual(Donation.objects.get().status, 'Received')
        second.refresh_from_db()
        self.assertEqual(second.status, 'Donated')
        self._shelve(second)
        self.assertEqual(Donation.objects.get().status, 'Shelved')

    def test_placing_a_book_by_saving_it_also_releases_it(self):
        book, = self._donate()
        book.shelf_level = self.level
        book.save()
        book.refresh_from_db()
        self.assertEqual(book.status, 'Available')

    def test_summary_cards_filter_the_list(self):
        from .views import _inventory_stats
        self._donate()
        Book.objects.create(title='Torn', author='A', qr_code='torn')
        torn = InventoryRecord.objects.create(book=Book.objects.get(title='Torn'), source='Existing',
                                              condition='Damaged', qr_label='torn-copy')
        InventoryRecord.objects.filter(pk=torn.pk).update(status='In Stock')
        gone = InventoryRecord.objects.create(source='Existing', title_hint='Mistake', qr_label='gone')
        InventoryRecord.objects.filter(pk=gone.pk).update(status='Removed')

        stats = _inventory_stats()
        self.assertEqual((stats['repair_count'], stats['removed_count']), (1, 1))

        def titles(state=''):
            page = self.client.get('/admin-portal/inventory/', {'state': state} if state else {})
            return [r.display_title for r in page.context['records']]
        self.assertNotIn('Mistake', titles())
        self.assertEqual(titles('repair'), ['Torn'])
        self.assertEqual(titles('removed'), ['Mistake'])
        self.assertIn('Noli', titles('unshelved'))

    def test_no_processing_stage_is_offered_anywhere(self):
        self._donate()
        inventory = self.client.get('/admin-portal/inventory/').content.decode()
        self.assertNotIn('processing_stage', inventory)
        self.assertNotIn('Accessioning stage', inventory)
        staff = _client(admin_id=_staff(modules='inventory,donations').admin_id, admin_role='Staff')
        for url in ('/library-staff/receiving/', '/library-staff/donations/'):
            page = staff.get(url).content.decode()
            self.assertNotIn('Processing', page, url)
            self.assertNotIn('processing_stage', page, url)
        self.assertIn('Waiting for a shelf', staff.get('/library-staff/donations/').content.decode())


class BeaconCheckTests(BeaconMinorTests):
    """Duplicate or impossible beacon values are caught while typing and refused on save."""

    def _check(self, **params):
        data = {'beacon_uuid': self.UUID, 'advertisement_type': 'iBeacon', 'major': 1}
        data.update(params)
        return self.client.get('/admin-portal/check-beacon/', data).json()

    def test_a_taken_minor_is_flagged_before_saving(self):
        self._place(minor=5, label='Door')
        result = self._check(minor=5)
        self.assertFalse(result['ok'])
        self.assertEqual(result['field'], 'minor')
        self.assertIn('"Door"', result['error'])
        self.assertTrue(self._check(minor=6)['ok'])

    def test_editing_a_beacon_does_not_clash_with_itself(self):
        mine = self._place(minor=5)['beacon']
        self._place(minor=6)
        self.assertTrue(self._check(minor=5, beacon_id=mine['beacon_id'])['ok'])
        self.assertFalse(self._check(minor=6, beacon_id=mine['beacon_id'])['ok'])

    def test_numbers_a_beacon_cannot_broadcast_are_refused(self):
        self.assertFalse(self._check(minor=70000)['ok'])
        self.assertFalse(self._check(major=-1, minor=2)['ok'])
        self.assertFalse(self._place(minor=70000)['success'])
        mine = self._place(minor=3)['beacon']
        r = self.client.post('/admin-portal/update-beacon/', {
            'beacon_id': mine['beacon_id'], 'beacon_uuid': self.UUID, 'major': 1, 'minor': -4}).json()
        self.assertFalse(r['success'])

    def test_the_form_has_a_live_check(self):
        page = self.client.get('/admin-portal/floorplan-management/').content.decode()
        self.assertIn('id="beaconCheckMsg"', page)
        self.assertIn('/admin-portal/check-beacon/', page)


class RenovationLiveTests(TestCase):
    """The patron map learns of a renovation, and of its end, from the data it refetches."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf
        self.ground = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        self.upper = FloorPlan.objects.create(name='Upper', floor_number=2, is_active=True,
                                              renovation_notice='Ceiling repairs',
                                              renovation_message='Closed until Friday.')
        room = Room.objects.create(floor_plan=self.upper, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf B', map_x=5, map_y=5)
        patron = Patron.objects.create(first_name='Ana', last_name='Cruz',
                                       email='ana@example.invalid', account_status='Active')
        self.client = _client(patron_id=patron.patron_id)

    def _data(self, **params):
        return self.client.get('/patron/map-data/', params).json()

    def test_the_floor_and_its_message_are_sent(self):
        d = self._data(floor=self.upper.floor_plan_id)
        self.assertEqual((d['renovation_notice'], d['renovation_message']), ('Ceiling repairs', 'Closed until Friday.'))
        tags = {f['floor_plan_id']: f['renovation'] for f in d['floors']}
        self.assertEqual(tags, {self.ground.floor_plan_id: '', self.upper.floor_plan_id: 'Ceiling repairs'})

    def test_the_target_floor_state_follows_the_toggle(self):
        d = self._data(floor=self.ground.floor_plan_id, shelf=self.shelf.shelf_id)
        self.assertEqual(d['target_renovation'], 'Ceiling repairs')
        self.upper.renovation_notice = None
        self.upper.save()
        d = self._data(floor=self.ground.floor_plan_id, shelf=self.shelf.shelf_id)
        self.assertEqual(d['target_renovation'], '')
        self.assertFalse(d['target_floor_closed'])

    def test_the_map_and_kiosk_listen_for_map_changes(self):
        page = self.client.get('/patron/map/').content.decode()
        self.assertIn('id="renoBanner"', page)
        kiosk = open('templates/desk/kiosk.html', encoding='utf-8').read()
        self.assertIn("topics='books map'", kiosk)


class CallNumberGuardTests(TestCase):
    """A bare number is a sticker number, not a call number: refused when typed, replaced when imported."""

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin', modules='books').admin_id, admin_role='Admin')
        self.ajax = {'HTTP_X_REQUESTED_WITH': 'XMLHttpRequest'}

    def test_the_check(self):
        from .views import call_number_error
        for bad in ('739', ' 14 ', '#22', 'No. 5', '739.0'):
            self.assertTrue(call_number_error(bad), bad)
        for good in ('', None, 'FIC D12a 1992', '813.54 B19h 2010', 'LB 1050 L34 1985', 'VAL C33e'):
            self.assertIsNone(call_number_error(good), good)

    def test_add_book_refuses_a_plain_number(self):
        r = self.client.post('/admin-portal/add-book/', dict(COMPLETE_BOOK, title='Noli', author='Rizal, Jose',
                                                              genre='Fiction', new_genre='1', call_number='739'),
                             **self.ajax).json()
        self.assertFalse(r['success'])
        self.assertIn('just a number', r['error'])
        self.assertFalse(Book.objects.filter(title='Noli').exists())

    def test_editing_refuses_a_plain_number_and_keeps_the_old_one(self):
        book = Book.objects.create(title='Noli', author='Rizal, Jose', genre='Fiction', call_number='FIC R49n')
        r = self.client.post('/admin-portal/edit-book/%d/' % book.book_id, {
            'title': 'Noli', 'author': 'Rizal, Jose', 'call_number': '14'}, **self.ajax).json()
        self.assertFalse(r['success'])
        book.refresh_from_db()
        self.assertEqual(book.call_number, 'FIC R49n')

    def test_import_replaces_a_plain_number_and_says_so(self):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.append(['Title', 'Author', 'Genre', 'Publication Year', 'Code Label', 'ISBN', 'Price'])
        wb.active.append(['Noli Me Tangere', 'Rizal, Jose', 'Fiction', 1887, 739, 'none', 300])
        wb.active.append(['El Filibusterismo', 'Rizal, Jose', 'Fiction', 1891, 'FIC R49e 1891', 'none', 300])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile('books.xlsx', buf.getvalue())
        r = self.client.post('/admin-portal/import-books/', {'excel_file': upload, 'new_genres': '1'}).json()
        self.assertIn('code label "739"', r['message'])
        noli = Book.objects.get(title='Noli Me Tangere')
        self.assertNotEqual(noli.call_number, '739')
        self.assertTrue(noli.call_number.startswith('FIC R'))
        self.assertEqual(Book.objects.get(title='El Filibusterismo').call_number, 'FIC R49e 1891')


class BeaconDefaultsTests(TestCase):
    """Placing a beacon needs no typing: everything comes from the library's own beacons."""
    UUID = 'fda50693a4e24fb1afcfc6eb07647825'

    def setUp(self):
        from .models import BLEBeacon, FloorPlan, Room
        self.upper = FloorPlan.objects.create(name='Upper', floor_number=2, is_active=True)
        self.ground = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        Room.objects.create(floor_plan=self.ground, name="Children's Room", map_x=0, map_y=0,
                            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        for minor in (1, 2):
            BLEBeacon.objects.create(floor_plan=self.upper, beacon_uuid=self.UUID, major=10011, minor=minor,
                                     map_x=5, map_y=5, tx_power=-59, path_loss_n=2.5, label='Hall %d' % minor)
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')

    def _defaults(self, **q):
        return self.client.get('/admin-portal/beacon-defaults/', q).json()

    def test_a_floor_without_beacons_still_gets_the_library_uuid(self):
        d = self._defaults(floor_plan_id=self.ground.floor_plan_id, x=50, y=50)
        self.assertEqual(d['uuid'], self.UUID)
        self.assertEqual(d['uuids'], [{'value': self.UUID, 'count': 2}])
        self.assertEqual((d['major'], d['minor']), (10011, 3))
        self.assertEqual((d['tx_power'], d['path_loss_n']), (-59, 2.5))

    def test_the_label_is_the_room_clicked_in_and_numbered(self):
        from .models import BLEBeacon
        d = self._defaults(floor_plan_id=self.ground.floor_plan_id, x=50, y=50)
        self.assertEqual((d['room'], d['label']), ("Children's Room", "Children's Room"))
        BLEBeacon.objects.create(floor_plan=self.ground, beacon_uuid=self.UUID, major=10011, minor=3,
                                 map_x=50, map_y=50, label="Children's Room")
        self.assertEqual(self._defaults(floor_plan_id=self.ground.floor_plan_id, x=60, y=60)['label'],
                         "Children's Room 2")
        outside = self._defaults(floor_plan_id=self.ground.floor_plan_id, x=500, y=500)
        self.assertEqual((outside['room'], outside['label']), (None, 'Beacon 4'))

    def test_with_no_beacons_at_all_nothing_is_invented_on_the_server(self):
        from .models import BLEBeacon
        BLEBeacon.objects.all().delete()
        d = self._defaults(floor_plan_id=self.ground.floor_plan_id)
        self.assertEqual((d['uuids'], d['uuid'], d['major'], d['minor']), ([], '', 1, 1))

    def test_the_form_offers_a_picker_and_a_scan(self):
        page = self.client.get('/admin-portal/floorplan-management/').content.decode()
        for marker in ('id="beaconUuidPick"', 'id="beaconScanBtn"', '/admin-portal/beacon-defaults/', 'requestLEScan'):
            self.assertIn(marker, page)


class RejectionReasonTests(TestCase):
    """Every way of turning a patron down needs a reason, and the reason can be found later."""

    def setUp(self):
        self.applicant = Patron.objects.create(
            first_name='Ana', last_name='Cruz', email='ana@example.invalid', account_status='Pending',
            registration_channel='Online', credential_document='credentials/x.png', otp_verified=True)
        self.member = Patron.objects.create(first_name='Ben', last_name='Reyes', email='ben@example.invalid',
                                            account_status='Inactive', otp_verified=True)
        self.admin = _client(admin_id=_staff(role='Admin', modules='patrons').admin_id, admin_role='Admin')
        self.staff = _client(admin_id=_staff(modules='patrons').admin_id, admin_role='Staff')

    def test_a_rejection_is_emailed_and_listed_with_its_reason(self):
        from django.core import mail
        self.staff.post('/admin-portal/reject-patron/%d/' % self.applicant.patron_id,
                        {'mode': 'reject', 'reason': 'Already a member (duplicate application)'})
        self.assertIn('Already a member (duplicate application)', mail.outbox[-1].body)
        page = self.staff.get('/library-staff/patrons/')
        self.assertEqual(page.context['rejected_count'], 1)
        row = page.context['rejected_applications'][0]
        self.assertEqual((row['name'], row['reason']), ('Ana Cruz', 'Already a member (duplicate application)'))
        self.assertContains(page, 'Rejected Applications')

    def test_an_application_cannot_be_deleted_edited_or_id_checked_around_the_review(self):
        pid = self.applicant.patron_id
        self.admin.post('/admin-portal/delete-patron/%d/' % pid)
        self.admin.post('/admin-portal/edit-patron/%d/' % pid, {
            'first_name': 'Ana', 'last_name': 'Cruz', 'patron_type': 'Student',
            'email': 'ana@example.invalid', 'account_status': 'Inactive'})
        self.staff.post('/admin-portal/patron-id-check/%d/' % pid)
        self.applicant.refresh_from_db()
        self.assertIsNone(self.applicant.archived_at)
        self.assertEqual(self.applicant.account_status, 'Pending')
        self.assertIsNone(self.applicant.identity_verified_at)

    def test_the_patron_list_offers_only_the_review_for_an_application(self):
        page = self.admin.get('/admin-portal/manage-patron/').content.decode()
        row = page.split('Ana Cruz')[-1].split('</tr>')[0]
        self.assertIn('Review ID', row)
        self.assertNotIn('admin-portal/delete-patron/%d/' % self.applicant.patron_id, row)
        for reason in ('The ID has expired', 'Not eligible for library membership'):
            self.assertIn(reason, page)

    def test_rejecting_a_reactivation_needs_a_reason(self):
        from .models import ReactivationRequest
        req = ReactivationRequest.objects.create(patron=self.member)
        url = '/admin-portal/reactivation/%d/respond/' % req.request_id
        self.admin.post(url, {'action': 'reject', 'staff_note': '  '})
        req.refresh_from_db()
        self.assertEqual(req.status, 'Pending')
        self.admin.post(url, {'action': 'reject', 'staff_note': 'Unreturned books or unpaid fines on the account'})
        req.refresh_from_db()
        self.assertEqual((req.status, req.staff_note), ('Rejected', 'Unreturned books or unpaid fines on the account'))

    def test_rejecting_a_card_photo_needs_a_reason(self):
        from .models import PatronPhoto
        photo = PatronPhoto.objects.create(patron=self.member, content_type='image/png', data=b'x')
        url = '/admin-portal/card-photo/%d/respond/' % photo.photo_id
        self.staff.post(url, {'action': 'reject'})
        photo.refresh_from_db()
        self.assertEqual(photo.status, 'Pending')
        self.staff.post(url, {'action': 'reject', 'staff_note': 'The photo is blurry or too dark'})
        photo.refresh_from_db()
        self.assertEqual((photo.status, photo.staff_note), ('Rejected', 'The photo is blurry or too dark'))


class FixLinkTests(TestCase):
    """The link to fix a registration works once, for two weeks, and can be renewed."""

    def setUp(self):
        self.applicant = Patron.objects.create(
            first_name='Ana', last_name='Cruz', email='ana@example.invalid', account_status='Pending',
            registration_channel='Online', credential_document='credentials/x.png', otp_verified=True,
            patron_type='Student')
        self.staff = _client(admin_id=_staff(modules='patrons').admin_id, admin_role='Staff')

    def _ask(self, needs_id):
        self.staff.post('/admin-portal/reject-patron/%d/' % self.applicant.patron_id, {
            'mode': 'fix', 'reason': 'The ID photo is blurry or unreadable', 'needs_id': '1' if needs_id else '0'})
        self.applicant.refresh_from_db()
        return '/patron/registration/fix/%s/' % self.applicant.fix_token

    def _details(self, **extra):
        data = {'first_name': 'Ana', 'last_name': 'Cruz', 'patron_type': 'Student'}
        data.update(extra)
        return data

    def test_a_new_id_is_required_when_the_id_was_the_problem(self):
        from django.core import mail
        url = self._ask(needs_id=True)
        self.assertIn('upload a new photo or scan of your ID', mail.outbox[-1].body)
        page = Client().post(url, self._details())
        self.assertContains(page, 'Please attach a new photo or scan of your ID')
        self.applicant.refresh_from_db()
        self.assertIsNone(self.applicant.resubmitted_at)
        png = SimpleUploadedFile('id.png', b'\x89PNG\r\n\x1a\n' + b'0' * 64, content_type='image/png')
        Client().post(url, self._details(credential_document=png))
        self.applicant.refresh_from_db()
        self.assertIsNotNone(self.applicant.resubmitted_at)

    def test_the_link_works_once(self):
        url = self._ask(needs_id=False)
        Client().post(url, self._details())
        self.assertContains(Client().get(url), 'already been used', status_code=404)

    def test_an_expired_link_can_be_renewed_by_the_applicant_and_by_staff(self):
        from django.core import mail
        url = self._ask(needs_id=True)
        Patron.objects.filter(pk=self.applicant.pk).update(fix_requested_at=timezone.now() - timedelta(days=15))
        self.applicant.refresh_from_db()
        self.assertTrue(self.applicant.fix_link_expired)
        self.assertContains(Client().get(url), 'Email me a new link', status_code=410)
        self.assertContains(self.staff.get('/library-staff/patrons/'), 'Fix link expired')

        sent = len(mail.outbox)
        page = Client().post(url, {'action': 'resend'})
        self.assertContains(page, 'New link sent')
        self.assertEqual(len(mail.outbox), sent + 1)
        self.assertEqual(mail.outbox[-1].to, ['ana@example.invalid'])
        self.applicant.refresh_from_db()
        self.assertFalse(self.applicant.fix_link_expired)
        self.assertTrue(self.applicant.fix_needs_id)
        self.assertContains(Client().get(url), 'already been used', status_code=404)

        Patron.objects.filter(pk=self.applicant.pk).update(fix_requested_at=timezone.now() - timedelta(days=15))
        self.staff.post('/admin-portal/resend-fix-link/%d/' % self.applicant.patron_id)
        self.applicant.refresh_from_db()
        self.assertFalse(self.applicant.fix_link_expired)

    def test_registering_again_while_pending_explains_itself(self):
        page = Client().post('/patron/register/', {
            'first_name': 'Ana', 'last_name': 'Cruz', 'email': 'ana@example.invalid',
            'password': 'Secret12345!', 'confirm_password': 'Secret12345!', 'patron_type': 'Student',
            'contact_number': '09170000000', 'address': 'Somewhere', 'accept_terms': 'yes'})
        self.assertContains(page, 'still reviewing it')


class DonorManagementTests(TestCase):
    """Donors are registered with their details, listed in full, and thanked in writing."""

    def setUp(self):
        from .models import Donor
        self.client = _client(admin_id=_staff(role='Admin').admin_id)
        self.bravo = Donor.objects.create(name='Johny Bravo')

    def _register(self, **data):
        data.setdefault('action', 'add')
        return self.client.post('/admin-portal/donors/', data)

    def test_registering_takes_every_detail_at_once(self):
        from .models import Donor
        self._register(name='Rotary Club of Cabuyao', donor_type='Organization',
                       contact_number='0917 111 2222', email='rotary@example.invalid', address='Cabuyao')
        d = Donor.objects.get(name='Rotary Club of Cabuyao')
        self.assertEqual((d.donor_type, d.contact_number, d.email, d.address),
                         ('Organization', '0917 111 2222', 'rotary@example.invalid', 'Cabuyao'))
        page = self._register(name='Rotary 2', email='not-an-email')
        self.assertContains(page, 'is not a valid email address')
        self.assertFalse(Donor.objects.filter(name='Rotary 2').exists())

    def test_a_near_duplicate_needs_confirming(self):
        from .models import Donor
        page = self._register(name='Johnny Bravo')
        self.assertContains(page, 'Is this someone already registered?')
        self.assertFalse(Donor.objects.filter(name='Johnny Bravo').exists())
        self._register(name='Johnny Bravo', confirm_new='1')
        self.assertTrue(Donor.objects.filter(name='Johnny Bravo').exists())
        # The same name, however typed, opens the saved record instead.
        self.assertRedirects(self._register(name='  johny   BRAVO '), '/admin-portal/donors/?d=%d' % self.bravo.donor_id)

    def _donate(self, donor_name, **extra):
        data = {'source': 'Donation', 'donor_name': donor_name, 'donated_date': '2026-10-01',
                'items': json.dumps(_with_genre_items([{'title': 'Noli', 'author': 'Rizal', 'quantity': 2, 'condition': 'Good'}]))}
        data.update(extra)
        return self.client.post('/admin-portal/inventory/receive/', data)

    def test_a_new_donor_is_saved_with_their_contact_when_receiving(self):
        from .models import Donor
        r = self._donate('Brgy. Council', donor_type='Government', donor_contact_number='0917 000 1111',
                         donor_email='brgy@example.invalid', donor_address='Poblacion')
        d = Donor.objects.get(name='Brgy. Council')
        self.assertEqual((d.donor_type, d.email), ('Government', 'brgy@example.invalid'))
        self.assertIn('ack=%d&ackdate=2026-10-01' % d.donor_id, r['Location'])
        page = self.client.get(r['Location'])
        self.assertContains(page, 'Print acknowledgement')

    def test_receiving_never_overwrites_a_saved_donors_details(self):
        self.bravo.contact_number = '0917 555 0000'
        self.bravo.save()
        self._donate('Johny Bravo', donor_contact_number='0999 999 9999', donor_address='Calamba')
        self.bravo.refresh_from_db()
        self.assertEqual((self.bravo.contact_number, self.bravo.address), ('0917 555 0000', 'Calamba'))

    def test_the_acknowledgement_lists_the_gift(self):
        self._donate('Johny Bravo')
        page = self.client.get('/admin-portal/donors/%d/acknowledgement/2026-10-01/' % self.bravo.donor_id)
        self.assertContains(page, 'Acknowledgement of Donation')
        self.assertContains(page, 'Noli')
        self.assertContains(page, 'DON-20261001-%04d' % self.bravo.donor_id)
        self.assertEqual(page.context['total_copies'], 2)
        missing = self.client.get('/admin-portal/donors/%d/acknowledgement/2026-01-01/' % self.bravo.donor_id)
        self.assertEqual(missing.status_code, 404)

    def test_the_donors_report_lists_every_donor_with_contact(self):
        from datetime import date
        from .models import Donor
        from .reports import build_report
        Donor.objects.create(name='Regular Reader', contact_number='0917', email='rr@example.invalid',
                             address='Sta. Rosa')
        self._donate('Johny Bravo')
        report = build_report('donors', date(2026, 10, 1), date(2026, 10, 31))
        names = [r[0] for r in report['rows']]
        self.assertEqual(names[0], 'Johny Bravo')
        row = report['rows'][names.index('Regular Reader')]
        self.assertEqual((row[2], row[3], row[4]), ('0917 · rr@example.invalid', 'Sta. Rosa', 0))
        self.assertIn('Address', report['columns'])


class ConditionScaleTests(TestCase):
    """Good / Minor Damage / Damaged / Beyond Repair, and what each means for lending."""

    def setUp(self):
        self.staff = _staff(role='Admin', modules='transactions,books,inventory,donations')
        self.client = _client(admin_id=self.staff.admin_id)

    def _receive(self, condition, source='Purchase', **extra):
        data = {'source': source, 'supplier': 'Test Supplier', 'items': json.dumps(_with_genre_items([
            {'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 1, 'condition': condition}]))}
        data.update(extra)
        self.client.post('/admin-portal/inventory/receive/', data)
        return Book.objects.get(title='Ibong Adarna')

    def test_damaged_leaves_circulation_and_returns_when_repaired(self):
        book = Book.objects.create(title='Noli', author='Rizal')
        book.condition = 'Damaged'
        book.save(update_fields=['condition'])
        book.refresh_from_db()
        self.assertEqual(book.status, 'Under Repair')
        book.condition = 'Good'
        book.save()
        book.refresh_from_db()
        self.assertEqual(book.status, 'Available')

    def test_beyond_repair_on_arrival_is_recorded_but_never_shelved_or_lent(self):
        from .views import _inventory_stats
        book = self._receive('Beyond Repair')
        self.assertEqual(book.status, 'Beyond Repair')
        self.assertEqual(InventoryRecord.objects.get().condition, 'Beyond Repair')
        stats = _inventory_stats()
        self.assertEqual((stats['disposal_count'], stats['unshelved_count']), (1, 0))
        page = Client().get('/patron/catalog/search/', {'q': 'Ibong'})
        self.assertNotContains(page, 'Ibong Adarna')

    def test_a_donation_that_arrives_unusable_is_not_left_waiting_for_a_shelf(self):
        from .models import Donation
        self._receive('Beyond Repair', source='Donation', donor_name='Rotary Club')
        self.assertEqual(Donation.objects.get().status, 'Not usable')

    def test_a_damaged_book_is_refused_at_the_desk(self):
        patron = Patron.objects.create(first_name='Ana', last_name='Cruz', email='ana@example.invalid',
                                       account_status='Active', identity_verified_at=timezone.now())
        book = self._receive('Damaged')
        self.assertEqual(book.status, 'Under Repair')
        r = self.client.post('/admin-portal/process-transaction/', {
            'transaction_type': 'Borrow', 'patron_id': patron.patron_id, 'book_ids': [book.book_id]},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertFalse(r.get('success'))
        self.assertIn('not lent until it is repaired', json.dumps(r))
        self.assertFalse(Transaction.objects.exists())

    def test_marking_a_copy_damaged_and_repaired_moves_the_book(self):
        book = self._receive('Good')
        record = InventoryRecord.objects.get()
        url = '/admin-portal/inventory/condition/'
        self.client.post(url, {'inventory_id': record.inventory_id, 'condition': 'Damaged',
                               'reason': 'Water damage', 'admin_password': 'DeskTest123'})
        book.refresh_from_db()
        self.assertEqual((book.condition, book.status), ('Damaged', 'Under Repair'))
        self.client.post(url, {'inventory_id': record.inventory_id, 'condition': 'Good',
                               'reason': 'Repaired and back in use', 'admin_password': 'DeskTest123'})
        book.refresh_from_db()
        self.assertEqual((book.condition, book.status), ('Good', 'Available'))

    def test_editing_the_books_condition_moves_its_copy_too(self):
        book = self._receive('Good')
        self.client.post('/admin-portal/edit-book/%d/' % book.book_id, {
            'title': book.title, 'author': book.author, 'condition': 'Minor Damage'},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(InventoryRecord.objects.get().condition, 'Minor Damage')

    def test_a_book_found_damaged_while_out_goes_to_repair_when_returned(self):
        patron = Patron.objects.create(first_name='Ben', last_name='Reyes', email='ben@example.invalid',
                                       account_status='Active', identity_verified_at=timezone.now())
        book = self._receive('Good')
        post = lambda kind: self.client.post('/admin-portal/process-transaction/', {
            'transaction_type': kind, 'patron_id': patron.patron_id, 'book_ids': [book.book_id]},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertTrue(post('Borrow').get('success'))
        self.client.post('/admin-portal/inventory/condition/', {
            'inventory_id': InventoryRecord.objects.get().inventory_id, 'condition': 'Damaged',
            'reason': 'Reported torn by the borrower', 'admin_password': 'DeskTest123'})
        book.refresh_from_db()
        self.assertEqual(book.status, 'Borrowed')
        self.assertTrue(post('Return').get('success'))
        book.refresh_from_db()
        self.assertEqual(book.status, 'Under Repair')


class AuditFollowUpTests(AuditTrailTests):
    """A count records what the system knows by itself, what the reader marks, and follows it all up."""

    def setUp(self):
        super().setUp()
        from .models import Shelf, ShelfLevel
        other = ShelfLevel.objects.create(shelf=Shelf.objects.get(), level_number=2)
        self.torn = Book.objects.create(title='Torn', author='E', shelf_level=self.level)
        self.repairing = Book.objects.create(title='Repairing', author='F', shelf_level=self.level,
                                             condition='Damaged')
        self.stray = Book.objects.create(title='Stray', author='G', shelf_level=other, status='Missing')
        for b in (self.here, self.gone, self.torn, self.stray):
            InventoryRecord.objects.create(book=b, source='Existing', qr_label='lbl-%d' % b.book_id,
                                           status='Missing' if b is self.stray else 'In Stock')

    def _file(self):
        return self.client.post('/admin-portal/inventory/audit/file/', {
            'level': self.level.shelf_level_id,
            'found_ids': [self.here.book_id], 'missing_ids': [self.gone.book_id],
            'damaged_ids': [self.torn.book_id], 'onshelf_ids': [self.out.book_id],
            'stray_ids': [self.stray.book_id]}).json()

    def test_every_kind_of_finding_is_recorded_and_acted_on(self):
        from .models import StockAudit, StockAuditLine
        r = self._file()
        self.assertTrue(r['success'], r)
        lines = {l.title: l for l in StockAuditLine.objects.all()}
        self.assertEqual(lines['Repairing'].result, 'Under repair')
        self.assertEqual(lines['Torn'].result, 'Damaged')
        self.assertEqual(lines['Out'].result, 'Here on loan')
        self.assertIn('check it in', lines['Out'].detail)
        self.assertEqual(lines['Stray'].result, 'Stray')
        self.assertEqual(lines['Here'].condition_then, 'Good')
        self.assertTrue(lines['Here'].location_then)

        self.torn.refresh_from_db()
        self.assertEqual((self.torn.condition, self.torn.status), ('Damaged', 'Under Repair'))
        self.stray.refresh_from_db()
        self.assertEqual(self.stray.status, 'For Reshelving')
        self.out.refresh_from_db()
        self.assertEqual(self.out.status, 'Borrowed')      # the desk closes the loan, not the count
        audit = StockAudit.objects.get()
        self.assertEqual((audit.damaged_count, audit.loan_mismatch_count, audit.unexpected_count), (1, 1, 1))

        alerts = self.client.get('/portal/alerts/').json()['items']
        self.assertEqual(next(i for i in alerts if i['key'] == 'loan_on_shelf')['count'], 1)

    def test_the_trail_follows_each_finding_to_now(self):
        r = self._file()
        # The loan is closed at the desk, the missing book is written off, the stray goes home.
        loan = Transaction.objects.get(book=self.out)
        loan.return_date = timezone.localdate()
        loan.save()
        self.out.status = 'Available'
        self.out.save()
        self.gone.status = 'Lost'
        self.gone.save()
        self.stray.status = 'Available'
        self.stray.save()

        d = self.client.get('/admin-portal/inventory/audit/%d/' % r['audit_id']).json()
        lines = {l['title']: l for l in d['lines']}
        self.assertEqual((lines['Out']['outcome'], lines['Out']['open']), ('Checked in', False))
        self.assertTrue(any('Returned by Ana' in e['text'] for e in lines['Out']['trail']))
        self.assertEqual(lines['Gone']['outcome'], 'Written off')
        self.assertEqual(lines['Stray']['outcome'], 'Back in place')
        self.assertEqual((lines['Torn']['outcome'], lines['Torn']['open']), ('Still under repair', True))
        self.assertIn('1 still open', d['summary'])

    def test_the_stock_audit_report(self):
        from .reports import build_report
        self._file()
        today = timezone.localdate()
        report = build_report('stock_audits', today, today)
        self.assertEqual(report['title'], 'Stock Audit Report')
        summary = dict(report['summary'])
        self.assertEqual((summary['Counts Filed'], summary['Not Found'], summary['Found Damaged'],
                          summary['On the Shelf While on Loan'], summary['Found on the Wrong Shelf']),
                         (1, 1, 1, 1, 1))
        books = [row[3] for row in report['rows']]
        self.assertIn('Torn', books)
        self.assertIn('Stray', books)
        page = self.client.get('/admin-portal/reports/', {'type': 'stock_audits', 'generate': 1,
                                                          'start': today.isoformat(), 'end': today.isoformat()})
        self.assertContains(page, 'Stock Audit Report')


class NewPatronNoticeTests(TestCase):
    """The people who approve registrations hear about a new one: sidebar, sign-in, dashboard, email."""

    def setUp(self):
        self.admin = _staff(role='Admin', modules='')
        self.reviewer = User.objects.create(fullname='Desk Patrons', email='patrons@example.invalid',
                                            password_hash=hash_password('DeskTest123'), role='Staff',
                                            account_status='Active', modules='patrons')
        self.other = User.objects.create(fullname='Desk Books', email='books@example.invalid',
                                         password_hash=hash_password('DeskTest123'), role='Staff',
                                         account_status='Active', modules='books')

    def _applicant(self, **extra):
        data = dict(first_name='Ana', last_name='Cruz', email='ana@example.invalid', account_status='Pending',
                    registration_channel='Online', otp_verified=False, otp_code='123456',
                    otp_expires_at=timezone.now() + timedelta(minutes=5), patron_type='Student')
        data.update(extra)
        return Patron.objects.create(**data)

    def test_confirming_the_email_tells_every_reviewer(self):
        from django.core import mail
        self._applicant()
        Client().post('/patron/register/', {'action': 'verify_otp', 'email': 'ana@example.invalid', 'otp': '123456'})
        to = sorted(m.to[0] for m in mail.outbox)
        self.assertEqual(to, sorted([self.admin.email, 'patrons@example.invalid']))
        message = next(m for m in mail.outbox if m.to == ['patrons@example.invalid'])
        self.assertIn('New registration: Ana Cruz', message.subject)
        self.assertIn('/library-staff/patrons/', message.body)

    def test_a_resubmission_is_emailed_too(self):
        from django.core import mail
        patron = self._applicant(otp_verified=True, fix_token='tok', fix_note='The ID has expired',
                                 fix_requested_at=timezone.now())
        Client().post('/patron/registration/fix/tok/', {'first_name': 'Ana', 'last_name': 'Cruz',
                                                       'patron_type': 'Student'})
        patron.refresh_from_db()
        self.assertIsNotNone(patron.resubmitted_at)
        self.assertTrue(any('Registration resubmitted' in m.subject and 'The ID has expired' in m.body
                            for m in mail.outbox))

    def test_the_sidebar_and_dashboard_show_what_is_waiting(self):
        self._applicant(otp_verified=True)
        staff = _client(admin_id=self.reviewer.admin_id, admin_role='Staff')
        page = staff.get('/library-staff/dashboard/')
        self.assertEqual(page.context['registrations_waiting'], 1)
        self.assertContains(page, 'data-alert-badge="registrations"')
        self.assertContains(page, 'waiting for review')
        books_only = _client(admin_id=self.other.admin_id, admin_role='Staff')
        self.assertEqual(books_only.get('/library-staff/dashboard/').context['registrations_waiting'], 0)

    def test_the_sign_in_notice_comes_once_per_sign_in(self):
        self._applicant(otp_verified=True)
        admin = _client(admin_id=self.admin.admin_id, admin_role='Admin')
        first = {i['key']: i for i in admin.get('/portal/alerts/').json()['items']}
        self.assertIn('1 registration is waiting for review', first['registrations']['notice'])
        again = {i['key']: i for i in admin.get('/portal/alerts/').json()['items']}
        self.assertNotIn('notice', again['registrations'])


def _sheet(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile('sheet.xlsx', buf.getvalue())


class StricterCataloguingTests(TestCase):
    """Every way in follows one set of cataloguing rules, and Add Book asks for as little as it can."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.noli = Book.objects.create(title='Noli Me Tangere', author='Rizal, Jose', genre='Fiction',
                                        ISBN='978-0143039693', publication_year=1887,
                                        shelf_level=self.level, call_number='FIC R49n 1887')
        self.staff = _staff(role='Admin', modules='books,inventory,donations')
        self.client = _client(admin_id=self.staff.admin_id, admin_role='Admin')
        self.ajax = {'HTTP_X_REQUESTED_WITH': 'XMLHttpRequest'}

    def _add(self, **data):
        return self.client.post('/admin-portal/add-book/', dict(COMPLETE_BOOK, **data), **self.ajax).json()

    def test_the_rules(self):
        from .catalogue import (clean_author, clean_copies, clean_genre, clean_isbn, clean_price,
                                clean_title, clean_year)
        for bad in ('N/A', 'none', '---', '???', 'Untitled', 'Spine not visible'):
            self.assertIsNotNone(clean_title(bad)[1], bad)
        self.assertEqual(clean_title('  BECOMING A  MASTER STUDENT ')[0], 'Becoming a Master Student')
        self.assertEqual(clean_author('anonymous')[0], 'Unknown')
        self.assertEqual(clean_author('Various')[0], 'Various Authors')
        self.assertIsNotNone(clean_author('n/a')[1])
        self.assertIsNotNone(clean_isbn('9780000000000', 'X')[1])          # wrong check digit
        self.assertEqual(clean_isbn('9780000000002', 'X'), ('9780000000002', None))
        self.assertEqual(clean_isbn('ACC-0042', 'X'), ('ACC-0042', None))  # an accession number
        self.assertIsNotNone(clean_isbn('see back cover', 'X')[1])
        # Another title's ISBN, written in its ten-digit form, is refused; its own copies are not.
        self.assertIn('already belongs', clean_isbn('0143039695', 'El Filibusterismo')[1])
        self.assertIsNone(clean_isbn('0143039695', 'Noli Me Tangere')[1])
        self.assertIsNotNone(clean_year('1200')[1])
        self.assertIsNotNone(clean_year(str(timezone.localdate().year + 2))[1])
        self.assertIsNotNone(clean_price('-5')[1])
        self.assertEqual(str(clean_price('₱1,250')[0]), '1250.00')
        self.assertIsNotNone(clean_copies('0')[1])
        self.assertIsNotNone(clean_copies('51')[1])
        self.assertEqual(clean_genre('FICTION'), ('Fiction', None))
        self.assertIsNotNone(clean_genre('Space Opera')[1])
        self.assertEqual(clean_genre('Space Opera', allow_new=True), ('Space Opera', None))
        self.assertIsNotNone(clean_genre('')[1])

    def test_add_book_lists_every_problem_and_saves_nothing(self):
        r = self._add(title='N/A', author='Rizal', ISBN='9780000000000', genre='Ficton')
        self.assertFalse(r['success'])
        self.assertEqual({p['field'] for p in r['problems']}, {'title', 'ISBN', 'genre'})
        self.assertEqual(Book.objects.count(), 1)

    def test_copies_are_made_at_once_and_share_a_call_number(self):
        r = self._add(title='El Filibusterismo', author='Rizal, Jose', genre='fiction',
                      publication_year='1891', copies='3', shelf_level=self.level.shelf_level_id,
                      condition='Minor Damage')
        self.assertTrue(r['success'], r)
        copies = Book.objects.filter(title='El Filibusterismo')
        self.assertEqual(copies.count(), 3)
        self.assertEqual({b.genre for b in copies}, {'Fiction'})          # the list's spelling
        self.assertEqual(len({b.call_number for b in copies}), 1)
        self.assertEqual(len({b.qr_code for b in copies}), 3)
        self.assertEqual({b.condition for b in copies}, {'Minor Damage'})
        self.assertEqual(InventoryRecord.objects.filter(book__in=copies).count(), 3)

    def test_another_copy_takes_the_titles_call_number(self):
        r = self._add(title='Noli Me Tangere', author='Rizal, Jose', genre='Fiction')
        self.assertTrue(r['success'], r)
        self.assertEqual(set(Book.objects.filter(title='Noli Me Tangere')
                             .values_list('call_number', flat=True)), {'FIC R49n 1887'})

    def test_a_new_genre_only_on_purpose(self):
        self.assertFalse(self._add(title='Dune', author='Herbert, Frank', genre='Space Opera')['success'])
        self.assertTrue(self._add(title='Dune', author='Herbert, Frank', genre='Space Opera',
                                  new_genre='1')['success'])
        self.assertEqual(Book.objects.get(title='Dune').genre, 'Space Opera')

    def test_the_shelf_can_wait(self):
        self.assertTrue(self._add(title='Florante at Laura', author='Balagtas', genre='Fiction')['success'])
        self.assertIsNone(Book.objects.get(title='Florante at Laura').shelf_level)

    def test_editing_follows_the_rules_but_its_own_isbn_is_no_clash(self):
        url = '/admin-portal/edit-book/%d/' % self.noli.book_id
        ok = self.client.post(url, {'title': 'Noli Me Tangere', 'author': 'Rizal, Jose', 'genre': 'Fiction',
                                    'ISBN': '978-0143039693'}, **self.ajax).json()
        self.assertTrue(ok['success'], ok)
        bad = self.client.post(url, {'title': 'Noli Me Tangere', 'author': 'Rizal, Jose',
                                     'genre': 'Novela'}, **self.ajax).json()
        self.assertEqual(bad['field'], 'genre')

    def test_the_assist_finds_copies_shelf_and_call_number(self):
        r = self.client.get('/admin-portal/catalogue-assist/', {'title': 'noli me tangere',
                                                                'author': 'Rizal, Jose'}).json()
        self.assertEqual(r['match']['copies'], 1)
        self.assertEqual(r['shelf']['id'], self.level.shelf_level_id)
        self.assertEqual(r['call_number'], 'FIC R49n 1887')
        # A new title is offered the shelf most of its genre is on, and a call number of its own.
        r = self.client.get('/admin-portal/catalogue-assist/', {
            'title': 'Ibong Adarna', 'author': 'Santos, Ana', 'genre': 'fiction', 'publication_year': '1990'}).json()
        self.assertIsNone(r['match'])
        self.assertEqual(r['shelf']['id'], self.level.shelf_level_id)
        self.assertTrue(r['call_number'].startswith('FIC S'))
        self.assertTrue(r['genre_known'])

    def test_the_isbn_lookup_matches_either_length(self):
        r = self.client.get('/admin-portal/isbn-lookup/', {'isbn': '0143039695'}).json()
        self.assertEqual(r['source'], 'catalogue')

    def test_an_illustrator_is_not_who_a_book_files_under(self):
        self.assertEqual(Book._surname_of('Meredith Ann Pierce (Ill. Jamichael Henterly)'), 'Pierce')
        self.assertEqual(Book._surname_of('Jane Austen and John Smith'), 'Austen')

    def test_the_import_refuses_rule_breaking_rows_and_adds_new_genres_on_purpose(self):
        rows = [['Title', 'Author', 'ISBN', 'Genre', 'Publication Year', 'Price'],
                ['Good Book', 'Cruz, Ana', 'none', 'fiction', 2001, 100],
                ['Space Book', 'Cruz, Ana', 'none', 'Space Opera', 2002, 100],
                ['Bad Isbn', 'Cruz, Ana', '9780000000000', 'Fiction', 2004, 100],
                ['No Genre', 'Cruz, Ana', 'none', '', 2005, 100],
                ['Far Future', 'Cruz, Ana', 'none', 'Fiction', 2999, 100]]
        r = self.client.post('/admin-portal/import-books/', {'excel_file': _sheet(rows)}).json()
        self.assertEqual(r['imported'], 1, r)
        self.assertEqual(r['new_genres'], ['Space Opera'])
        self.assertTrue(any(p.startswith('Row 4') and 'check digit' in p for p in r['problems']), r['problems'])
        self.assertIn('Row 5: no genre', r['problems'])
        self.assertTrue(any(p.startswith('Row 6') and '2999' in p for p in r['problems']), r['problems'])
        self.assertEqual(Book.objects.get(title='Good Book').genre, 'Fiction')
        self.assertFalse(Book.objects.filter(title__in=['Space Book', 'Bad Isbn', 'No Genre', 'Far Future']).exists())
        self.client.post('/admin-portal/import-books/', {'excel_file': _sheet([rows[0], rows[2]]),
                                                         'new_genres': '1'})
        self.assertEqual(Book.objects.get(title='Space Book').genre, 'Space Opera')

    def test_the_donation_import_asks_before_adding_new_genres(self):
        rows = [['Donor', 'Title', 'Author', 'Genre'],
                ['Ana Cruz', 'Ibong Adarna', 'Unknown', 'Epic Poem'],
                ['Ana Cruz', 'El Filibusterismo', 'Rizal, Jose', 'Fiction']]
        rows = _with_genre(rows)
        r = self.client.post('/admin-portal/import-donations/', {'excel_file': _sheet(rows)}).json()
        self.assertFalse(r['success'])
        self.assertEqual(r['new_genres'], ['Epic Poem'])
        self.assertFalse(Book.objects.filter(title__in=['Ibong Adarna', 'El Filibusterismo']).exists())
        r = self.client.post('/admin-portal/import-donations/', {'excel_file': _sheet(rows), 'new_genres': '1'}).json()
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.get(title='Ibong Adarna').genre, 'Epic Poem')

    def test_receiving_a_new_title_follows_the_rules(self):
        def receive(**item):
            line = {'title': 'Ibong Adarna', 'author': 'Unknown', 'quantity': 1, 'condition': 'Good',
                    'isbn_none': True, 'year_unknown': True, 'price': '120'}
            line.update(item)
            self.client.post('/admin-portal/inventory/receive/', {'source': 'Purchase', 'supplier': 'Test Supplier', 'items': json.dumps([line])})
            return Book.objects.filter(title='Ibong Adarna').first()
        self.assertIsNone(receive(genre='Epic Poem'))
        self.assertIsNone(receive(genre='Fiction', author='n/a'))
        self.assertEqual(receive(genre='Epic Poem', new_genre=True).genre, 'Epic Poem')
        # A title already catalogued needs nothing more.
        before = Book.objects.filter(title='Noli Me Tangere').count()
        self.client.post('/admin-portal/inventory/receive/', {'source': 'Purchase', 'supplier': 'Test Supplier', 'items': json.dumps(
            [{'book_id': self.noli.book_id, 'quantity': 1, 'condition': 'Good'}])})
        self.assertEqual(Book.objects.filter(title='Noli Me Tangere').count(), before + 1)


class GenreTidyTests(TestCase):
    """584 one-off descriptions became 45 genres; the old wording lives on as the subject."""

    def setUp(self):
        self.staff = _staff(role='Admin', modules='books,inventory,donations')
        self.client = _client(admin_id=self.staff.admin_id, admin_role='Admin')
        self.ajax = {'HTTP_X_REQUESTED_WITH': 'XMLHttpRequest'}

    def test_old_wording_lands_in_the_standard_list(self):
        from .genres import STANDARD_GENRES, standard_genre
        cases = {
            'Legal Thriller': 'Thriller & Suspense',
            "Children's Picture Book / Dinosaurs & Pop-Up": "Children's Picture Books",
            'General Reference / Encyclopedia': 'Encyclopedias',
            'Values Education Textbook': 'Values Education',
            'Social Studies / Textbook (HEKASI)': 'Social Studies',
            'Academic Literature / Poetry': 'Filipino Language & Literature',
            'Non-Fiction / History': 'History',
            'Politics / Non-Fiction': 'Social Sciences',
            'Political Science & Government': 'Social Sciences',
            'Business Statistics': 'Statistics',
            'Auditing & Information Systems': 'Accounting',
            'fiction': 'Fiction',
        }
        for old, new in cases.items():
            self.assertEqual(standard_genre(old), new, old)
        self.assertEqual(len(STANDARD_GENRES), 45)
        self.assertIsNone(standard_genre('Underwater Basketweaving'))

    def test_the_tidy_keeps_the_old_wording_as_the_subject(self):
        import importlib
        from django.apps import apps
        from .models import SystemLog
        legal = Book.objects.create(title='The Firm', author='Grisham, John', genre='Legal Thriller')
        plain = Book.objects.create(title='Emma', author='Austen, Jane', genre='Fiction')
        odd = Book.objects.create(title='Odd', author='Cruz, Ana', genre='Underwater Basketweaving')
        importlib.import_module('library.migrations.0089_tidy_genres').tidy_genres(apps, None)
        for book in (legal, plain, odd):
            book.refresh_from_db()
        self.assertEqual((legal.genre, legal.subject), ('Thriller & Suspense', 'Legal Thriller'))
        self.assertEqual((plain.genre, plain.subject), ('Fiction', None))
        self.assertEqual((odd.genre, odd.subject), ('Underwater Basketweaving', None))
        self.assertTrue(SystemLog.objects.filter(detail__startswith='Tidied genres').exists())

    def test_every_search_box_finds_the_subject(self):
        Book.objects.create(title='The Firm', author='Grisham, John', genre='Thriller & Suspense',
                            subject='Legal Thriller')
        page = self.client.get('/admin-portal/management/', {'q': 'legal thrill'})
        self.assertEqual([b.title for b in page.context['books']], ['The Firm'])

    def test_the_import_matches_wording_and_keeps_it_as_the_subject(self):
        Book.objects.create(title='The Firm', author='Grisham, John', genre='Thriller & Suspense')
        rows = [['Title', 'Author', 'Genre', 'Subject', 'ISBN', 'Publication Year', 'Price'],
                ['The Client', 'Grisham, John', 'Legal Thriller', '', 'none', 1993, 300],
                ['The Spy', 'Le Carre, John', '', 'Spy Thriller', 'none', 1963, 300],
                ['The Basket', 'Cruz, Ana', 'Underwater Basketweaving', '', 'none', 2001, 300]]
        r = self.client.post('/admin-portal/import-books/', {'excel_file': _sheet(rows)}).json()
        client = Book.objects.get(title='The Client')
        self.assertEqual((client.genre, client.subject), ('Thriller & Suspense', 'Legal Thriller'))
        spy = Book.objects.get(title='The Spy')
        self.assertEqual((spy.genre, spy.subject), ('Thriller & Suspense', 'Spy Thriller'))
        self.assertIn('kept as the subject', r['message'])
        # Wording no rule fits is still a new genre, added only on purpose.
        self.assertEqual(r['new_genres'], ['Underwater Basketweaving'])
        self.assertFalse(Book.objects.filter(title='The Basket').exists())

    def test_the_donation_import_matches_too(self):
        Book.objects.create(title='The Firm', author='Grisham, John', genre='Thriller & Suspense')
        rows = [['Donor', 'Title', 'Author', 'Genre'], ['Ana Cruz', 'The Client', 'Grisham, John', 'Legal Thriller']]
        r = self.client.post('/admin-portal/import-donations/', {'excel_file': _sheet(_with_genre(rows))}).json()
        self.assertTrue(r['success'], r)
        client = Book.objects.get(title='The Client')
        self.assertEqual((client.genre, client.subject), ('Thriller & Suspense', 'Legal Thriller'))

    def test_another_copy_takes_its_titles_subject_and_editing_changes_it(self):
        first = Book.objects.create(title='The Firm', author='Grisham, John', genre='Thriller & Suspense',
                                    subject='Legal Thriller')
        r = self.client.post('/admin-portal/add-book/', dict(COMPLETE_BOOK, title='The Firm', author='Grisham, John',
                                                              genre='Thriller & Suspense'), **self.ajax).json()
        self.assertTrue(r['success'], r)
        self.assertEqual(set(Book.objects.filter(title='The Firm').values_list('subject', flat=True)),
                         {'Legal Thriller'})
        self.client.post('/admin-portal/edit-book/%d/' % first.book_id, {
            'title': 'The Firm', 'author': 'Grisham, John', 'genre': 'Thriller & Suspense',
            'subject': '  Courtroom   Drama '}, **self.ajax)
        first.refresh_from_db()
        self.assertEqual(first.subject, 'Courtroom Drama')

    def test_the_edit_form_fills_in_an_apostrophe_as_typed(self):
        book = Book.objects.create(title="Charlotte's Web", author='White, E. B.', genre="Children's Fiction")
        page = self.client.get('/admin-portal/edit-book/%d/' % book.book_id).content.decode()
        self.assertIn("Children\\u0027s Fiction", page)
        self.assertNotIn("value = 'Children&#x27;s", page)

    def test_editing_keeps_a_magazine_a_magazine(self):
        book = Book.objects.create(title='Reader', author='Unknown', genre='Fiction', material_type='Magazine')
        page = self.client.get('/admin-portal/edit-book/%d/' % book.book_id).content.decode()
        self.assertIn('<option value="Magazine">', page)
        self.assertIn("getElementById('editMaterial').value = 'Magazine'", page)

    def test_the_picker_shows_what_each_genre_covers(self):
        from .views import _genre_options
        Book.objects.create(title='The Firm', author='Grisham, John', genre='Thriller & Suspense',
                            subject='Legal Thriller')
        option = next(o for o in _genre_options() if o['label'] == 'Thriller & Suspense')
        self.assertIn('Legal Thriller', option['sub'])
        self.assertIn('Legal Thriller', option['search'])


class ShelvingLabelsAuditTests(TestCase):
    """Books not yet ready for the floor are counted one way everywhere; a shelf can be audited whole."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=1)
        self.fiction = Book.objects.create(title='Noli Me Tangere', author='Rizal, Jose', genre='Fiction',
                                           shelf_level=self.level, shelf_slot=1, call_number='FIC R49n',
                                           qr_code='q-noli')
        self.loose = Book.objects.create(title='El Filibusterismo', author='Rizal, Jose', genre='Fiction',
                                         qr_code='q-fili')
        self.donated = Book.objects.create(title='Florante at Laura', author='Balagtas', genre='Fiction',
                                           status='Donated', qr_code='q-flor')
        # Away or gone: not waiting for anyone to shelve them.
        Book.objects.create(title='Out', author='Cruz, Ana', genre='Fiction', status='Borrowed', qr_code='q-out')
        Book.objects.create(title='Gone', author='Cruz, Ana', genre='Fiction', status='Lost', qr_code='q-gone')
        self.admin = _staff(role='Admin', modules='books,shelf,inventory')
        self.client = _client(admin_id=self.admin.admin_id, admin_role='Admin')

    def test_one_count_everywhere(self):
        from .shelving import need_label, need_shelf
        from .views import _inventory_stats
        self.assertEqual(set(need_shelf().values_list('title', flat=True)), {'El Filibusterismo', 'Florante at Laura'})
        self.assertEqual(list(need_label().values_list('title', flat=True)), ['Noli Me Tangere'])
        self.assertEqual(self.client.get('/admin-portal/management/').context['unshelved_count'], 2)
        self.assertEqual(self.client.get('/admin-portal/shelf-manager/').context['unshelved_count'], 2)
        self.assertEqual(_inventory_stats()['unshelved_count'], 2)
        listed = self.client.get('/admin-portal/board-books/', {'level': 'none'}).json()['books']
        self.assertEqual(len(listed), 2)
        shelving = {i['key']: i for i in self.client.get('/portal/alerts/').json()['items']}['shelving']
        self.assertEqual((shelving['count'], shelving['quiet']), (2, True))
        page = self.client.get('/admin-portal/dashboard/')
        self.assertContains(page, 'books waiting for a shelf')

    def test_unlabelled_books_are_listed_one_by_one_and_print_alone(self):
        second = Book.objects.create(title='Ibong Adarna', author='Anon', genre='Fiction', qr_code='q-ib',
                                     shelf_level=self.level, shelf_slot=5)
        page = self.client.get('/admin-portal/shelving/')
        board = page.context['label_boards'][0]
        self.assertEqual([b.title for b in board['books']], ['Noli Me Tangere', 'Ibong Adarna'])
        self.assertContains(page, 'data-id="%d"' % second.pk)
        self.assertContains(page, 'Print board')
        sheet = self.client.get('/admin-portal/book-qr-labels/', {'ids': str(second.pk)})
        self.assertEqual(sheet['Content-Type'], 'application/pdf')
        from .shelving import need_label
        self.assertEqual(list(need_label().values_list('title', flat=True)), ['Noli Me Tangere'])

    def test_a_waiting_book_is_shelved_on_its_suggested_board_in_one_click(self):
        page = self.client.get('/admin-portal/shelving/')
        rows = {r['book'].title: r for r in page.context['waiting']}
        self.assertEqual(rows['El Filibusterismo']['level'], self.level)
        self.assertIn('Fiction', rows['El Filibusterismo']['why'])
        r = self.client.post('/admin-portal/shelving/place/', {
            'placements': '%d:%d,%d:%d' % (self.loose.pk, self.level.pk, self.donated.pk, self.level.pk)}).json()
        self.assertTrue(r['success'], r)
        self.loose.refresh_from_db()
        self.donated.refresh_from_db()
        self.assertEqual(self.loose.shelf_level, self.level)
        self.assertEqual(sorted([self.loose.shelf_slot, self.donated.shelf_slot]), [2, 3])
        # A donated book on a shelf is released for borrowing.
        self.assertEqual(self.donated.status, 'Available')
        self.assertEqual(sorted(r['ids']), sorted([self.loose.pk, self.donated.pk]))

    def test_add_book_hands_back_what_it_can_label(self):
        shelved = self.client.post('/admin-portal/add-book/', dict(
            COMPLETE_BOOK, title='Ibong Adarna', author='Unknown', genre='Fiction',
            shelf_level=self.level.pk, copies='2'), HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertTrue(shelved['shelved'])
        self.assertEqual(len(shelved['ids']), 2)
        loose = self.client.post('/admin-portal/add-book/', dict(
            COMPLETE_BOOK, title='Mga Ibong Mandaragit', author='Hernandez, Amado', genre='Fiction'),
            HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertFalse(loose['shelved'])

    def test_the_shelf_audit_lists_what_the_records_get_wrong(self):
        Book.objects.create(title='N/A', author='Cruz, Ana', genre='Fiction', shelf_level=self.level,
                            shelf_slot=1, call_number='739', qr_code='q-na')
        Book.objects.create(title='Lost One', author='Cruz, Ana', genre='Fiction', shelf_level=self.level,
                            shelf_slot=4, status='Lost', call_number='FIC C78l', qr_code='q-lost')
        page = self.client.get('/admin-portal/shelf-audit/%d/' % self.shelf.pk)
        kinds = {g['kind']: [p['book'].title for p in g['items']] for g in page.context['problem_groups']}
        self.assertIn('N/A', kinds['rules'])
        self.assertIn('N/A', kinds['call_number'])
        self.assertEqual(sorted(kinds['slot']), ['N/A', 'Noli Me Tangere'])
        self.assertEqual(kinds['gone'], ['Lost One'])
        self.assertEqual(page.context['boards'][0]['books'], 3)

    def test_counting_a_board_records_its_labels(self):
        other = Book.objects.create(title='Doctrina Christiana', author='Unknown', genre='Fiction',
                                    shelf_level=self.level, shelf_slot=2, qr_code='q-doc')
        r = self.client.post('/admin-portal/inventory/audit/file/', {
            'level': self.level.pk, 'found_ids': [self.fiction.pk], 'bulk_ids': [other.pk],
            'nolabel_ids': [other.pk]}).json()
        self.assertTrue(r['success'], r)
        self.fiction.refresh_from_db()
        other.refresh_from_db()
        self.assertIsNotNone(self.fiction.label_printed_at)
        self.assertIsNone(other.label_printed_at)
        self.assertEqual((r['labelled'], r['nolabel'], r['nolabel_ids']), (1, 1, [other.pk]))

    def test_staff_who_shelve_can_audit_but_others_cannot(self):
        shelver = _staff(modules='shelf')
        staff = _client(admin_id=shelver.admin_id, admin_role='Staff')
        self.assertEqual(staff.get('/admin-portal/shelf-audit/%d/' % self.shelf.pk).status_code, 200)
        self.assertTrue(staff.get('/admin-portal/inventory/audit/sheet/', {'level': self.level.pk}).json()['success'])
        self.assertEqual(staff.get('/admin-portal/shelving/').status_code, 200)
        desk = User.objects.create(fullname='Desk Only', email='desk-only@example.invalid', role='Staff',
                                   account_status='Active', modules='transactions')
        other = _client(admin_id=desk.admin_id, admin_role='Staff')
        self.assertEqual(other.get('/admin-portal/shelf-audit/%d/' % self.shelf.pk).status_code, 302)

    def test_the_report_lists_both(self):
        from .reports import build_report
        today = timezone.localdate()
        report = build_report('shelving', today, today)
        needs = sorted(row[0] for row in report['rows'])
        self.assertEqual(needs, ['No label', 'Waiting for a shelf', 'Waiting for a shelf'])
        self.assertIn(('Waiting for a Shelf', 2), report['summary'])


class DonorNamePartsTests(TestCase):
    """A person who donates is named as first name, middle name (optional) and surname."""

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')
        self.book = Book.objects.create(title='Noli', author='Rizal', genre='Fiction', qr_code='q-noli')

    def _give(self, **donor):
        data = {'source': 'Donation', 'donated_date': '2026-10-01',
                'items': json.dumps([{'book_id': self.book.pk, 'quantity': 1, 'condition': 'Good'}])}
        data.update(donor)
        return self.client.post('/admin-portal/inventory/receive/', data, follow=True)

    def test_a_person_is_saved_with_the_name_split(self):
        from .models import Donor
        self._give(donor_type='Individual', donor_first_name=' juan ', donor_middle_name='Perez',
                   donor_last_name='Dela Cruz')
        donor = Donor.objects.get()
        self.assertEqual((donor.first_name, donor.middle_name, donor.last_name), ('juan', 'Perez', 'Dela Cruz'))
        self.assertEqual(donor.name, 'juan Perez Dela Cruz')
        self.assertEqual(InventoryRecord.objects.get().donor_name, 'juan Perez Dela Cruz')

    def test_the_surname_is_required_but_the_middle_name_is_not(self):
        from .models import Donor
        page = self._give(donor_type='Individual', donor_first_name='Juan')
        self.assertContains(page, 'first name and surname')
        self.assertFalse(Donor.objects.exists())
        self._give(donor_type='Individual', donor_first_name='Juan', donor_last_name='Cruz')
        self.assertEqual(Donor.objects.get().middle_name, None)

    def test_an_organisation_keeps_one_name(self):
        from .models import Donor
        self._give(donor_type='School', donor_name='Pulo National High School')
        donor = Donor.objects.get()
        self.assertEqual((donor.name, donor.first_name, donor.donor_type),
                         ('Pulo National High School', None, 'School'))

    def test_a_past_donor_is_matched_by_their_name_parts(self):
        from .models import Donor
        Donor.objects.create(name='Dela Cruz, Juan', first_name='Juan', last_name='Dela Cruz')
        self._give(donor_type='Individual', donor_first_name='JUAN', donor_last_name='dela cruz')
        self.assertEqual(Donor.objects.count(), 1)
        self.assertEqual(InventoryRecord.objects.get().donor_name, 'Dela Cruz, Juan')

    def test_the_donors_page_registers_and_edits_by_parts(self):
        from .models import Donor
        self.client.post('/admin-portal/donors/', {'action': 'add', 'donor_type': 'Individual',
                                                   'first_name': 'Ana', 'last_name': 'Reyes', 'confirm_new': '1'})
        donor = Donor.objects.get()
        self.assertEqual(donor.name, 'Ana Reyes')
        self.client.post('/admin-portal/donors/', {'donor_id': donor.pk, 'donor_type': 'Individual',
                                                   'first_name': 'Ana', 'middle_name': 'Santos',
                                                   'last_name': 'Reyes'})
        donor.refresh_from_db()
        self.assertEqual((donor.name, donor.middle_name), ('Ana Santos Reyes', 'Santos'))

    def test_the_sheet_takes_name_columns(self):
        from .models import Donor
        rows = [['first_name', 'middle_name', 'last_name', 'donor_name', 'title', 'author', 'genre'],
                ['Juan', '', 'Cruz', '', 'Noli', 'Rizal', 'Fiction'],
                ['', '', '', 'Rotary Club', 'Noli', 'Rizal', 'Fiction'],
                ['Pedro', '', '', '', 'Noli', 'Rizal', 'Fiction']]
        r = self.client.post('/admin-portal/import-donations/', {'excel_file': _sheet(rows)}).json()
        self.assertTrue(r['success'], r)
        self.assertIn('Row 4: a person needs both a first name and a surname', r['message'])
        self.assertEqual(Donor.objects.get(first_name='Juan').last_name, 'Cruz')
        self.assertEqual(set(InventoryRecord.objects.values_list('donor_name', flat=True)),
                         {'Juan Cruz', 'Rotary Club'})


class CompleteBookRecordTests(TestCase):
    """A new book has its ISBN and year, or says on purpose it has none, and always its price."""

    def setUp(self):
        Book.objects.create(title='Noli Me Tangere', author='Rizal, Jose', genre='Fiction', qr_code='q-noli')
        self.client = _client(admin_id=_staff(role='Admin', modules='books,inventory,donations').admin_id,
                              admin_role='Admin')
        self.ajax = {'HTTP_X_REQUESTED_WITH': 'XMLHttpRequest'}

    def _add(self, **data):
        base = {'title': 'Florante at Laura', 'author': 'Balagtas', 'genre': 'Fiction'}
        return self.client.post('/admin-portal/add-book/', dict(base, **data), **self.ajax).json()

    def test_add_book_refuses_blanks_left_by_accident(self):
        r = self._add()
        self.assertFalse(r['success'])
        self.assertEqual({p['field'] for p in r['problems']}, {'ISBN', 'publication_year', 'price'})
        self.assertFalse(Book.objects.filter(title='Florante at Laura').exists())

    def test_none_and_unknown_said_on_purpose_are_kept(self):
        r = self._add(no_isbn='1', year_unknown='1', price='150')
        self.assertTrue(r['success'], r)
        book = Book.objects.get(title='Florante at Laura')
        self.assertEqual((book.ISBN, book.isbn_none, book.publication_year, book.year_unknown),
                         (None, True, None, True))

    def test_an_edit_asks_for_what_an_old_record_lacks_and_remembers_the_answer(self):
        old = Book.objects.get(title='Noli Me Tangere')
        url = '/admin-portal/edit-book/%d/' % old.pk
        form = {'title': old.title, 'author': old.author, 'genre': 'Fiction', 'ISBN': '', 'publication_year': '',
                'price': ''}
        r = self.client.post(url, form, **self.ajax).json()
        self.assertEqual({p['field'] for p in r['problems']}, {'ISBN', 'publication_year', 'price'})
        r = self.client.post(url, dict(form, no_isbn='1', publication_year='1887', price='250'), **self.ajax).json()
        self.assertTrue(r['success'], r)
        old.refresh_from_db()
        self.assertTrue(old.isbn_none)
        self.assertIn('no_isbn', '\n'.join(
            line for line in self.client.get(url).content.decode().splitlines() if 'editNoIsbn' in line))

    def test_the_import_refuses_blanks_but_takes_none_and_unknown(self):
        rows = [['Title', 'Author', 'Genre', 'ISBN', 'Publication Year', 'Price'],
                ['Blank Everything', 'Cruz, Ana', 'Fiction', '', '', ''],
                ['Said None', 'Cruz, Ana', 'Fiction', 'N/A', 'unknown', 120],
                ['Complete', 'Cruz, Ana', 'Fiction', '9780000000002', 1999, 300]]
        r = self.client.post('/admin-portal/import-books/', {'excel_file': _sheet(rows)}).json()
        self.assertEqual(r['imported'], 2, r)
        problem = r['problems'][0]
        self.assertTrue(problem.startswith('Row 2'))
        for words in ('no ISBN', 'no publication year', 'no price'):
            self.assertIn(words, problem)
        said = Book.objects.get(title='Said None')
        self.assertEqual((said.isbn_none, said.year_unknown, str(said.price)), (True, True, '120.00'))

    def test_receiving_a_new_title_needs_its_price_but_a_catalogued_one_does_not(self):
        def receive(line):
            self.client.post('/admin-portal/inventory/receive/', {'source': 'Purchase', 'supplier': 'Test Supplier', 'items': json.dumps([line])})
        receive({'title': 'Ibong Adarna', 'author': 'Unknown', 'genre': 'Fiction', 'quantity': 1,
                 'condition': 'Good', 'isbn_none': True, 'year_unknown': True})
        self.assertFalse(Book.objects.filter(title='Ibong Adarna').exists())
        receive({'title': 'Ibong Adarna', 'author': 'Unknown', 'genre': 'Fiction', 'quantity': 2,
                 'condition': 'Good', 'isbn_none': True, 'year_unknown': True, 'price': '90'})
        self.assertEqual(set(Book.objects.filter(title='Ibong Adarna').values_list('price', flat=True)),
                         {Decimal('90.00')})
        noli = Book.objects.get(title='Noli Me Tangere')
        receive({'book_id': noli.pk, 'quantity': 1, 'condition': 'Good'})
        self.assertEqual(Book.objects.filter(title='Noli Me Tangere').count(), 2)

    def test_the_donation_sheet_needs_them_only_for_a_new_title(self):
        rows = [['Donor', 'Title', 'Author', 'Genre', 'ISBN', 'Publication Year', 'Price'],
                ['Ana Cruz', 'Noli Me Tangere', 'Rizal, Jose', '', '', '', ''],
                ['Ana Cruz', 'Doctrina Christiana', 'Unknown', 'Fiction', 'none', 'unknown', ''],
                ['Ana Cruz', 'Urbana at Feliza', 'Pilapil', 'Fiction', 'none', 1864, 80]]
        r = self.client.post('/admin-portal/import-donations/', {'excel_file': _sheet(rows)}).json()
        self.assertTrue(r['success'], r)
        self.assertIn('Row 3', r['message'])
        self.assertIn('no price', r['message'])
        self.assertEqual(Book.objects.filter(title='Noli Me Tangere').count(), 2)
        self.assertFalse(Book.objects.filter(title='Doctrina Christiana').exists())
        self.assertEqual(Book.objects.get(title='Urbana at Feliza').price, Decimal('80.00'))


class DeskAccountIssueTests(TestCase):
    """A membership on hold cannot come in by any route, can always leave, and staff can let it in with a reason."""

    def setUp(self):
        from .models import PatronLog
        self.PatronLog = PatronLog
        self.held = Patron.objects.create(first_name='Ana', last_name='Cruz', fullname='Ana Cruz',
                                          email='ana.held@example.invalid', contact_number='09171234567',
                                          account_status='Suspended', qr_code='card-ana')
        self.staff_user = _staff(role='Admin', modules='logs')
        self.staff = _client(admin_id=self.staff_user.admin_id, admin_role='Admin')
        self.kiosk = _client(admin_id=self.staff_user.admin_id, admin_role='Admin', desk_mode=True)

    def test_typing_the_name_no_longer_makes_a_new_visitor(self):
        before = Patron.objects.count()
        for url, data in (('/desk/sign/', {'name': 'Ana Cruz', 'contact': ''}),
                          ('/desk/visitor/', {'name': 'Ana Cruz', 'contact': '', 'purpose': 'Study'})):
            r = self.kiosk.post(url, data).json()
            self.assertFalse(r['success'], url)
            self.assertIn('on hold', r['error'])
        self.assertEqual(Patron.objects.count(), before)
        self.assertFalse(self.PatronLog.objects.exists())

    def test_the_kiosk_refuses_entry_by_card_but_lets_them_leave(self):
        self.assertFalse(self.kiosk.post('/desk/scan/', {'code': 'card-ana'}).json()['success'])
        self.PatronLog.objects.create(patron=self.held, entry_time=timezone.now() - timedelta(hours=2))
        r = self.kiosk.post('/desk/scan/', {'code': 'card-ana'}).json()
        self.assertEqual((r['success'], r['action']), (True, 'exit'))
        self.PatronLog.objects.create(patron=self.held, entry_time=timezone.now() - timedelta(hours=1))
        who = self.kiosk.post('/desk/identify/', {'value': 'ana.held@example.invalid'}).json()
        self.assertTrue(who['exit_only'])
        self.assertTrue(self.kiosk.post('/desk/visit/', {'direction': 'exit'}).json()['success'])

    def test_staff_can_let_them_in_only_with_a_reason_which_is_kept(self):
        from .models import SystemLog
        r = self.staff.post('/desk/scan/', {'code': 'card-ana'}).json()
        self.assertTrue(r['needs_override'])
        self.assertFalse(self.PatronLog.objects.exists())
        # The kiosk never offers it, reason or not.
        r = self.kiosk.post('/desk/scan/', {'code': 'card-ana', 'override_reason': 'Returning books'}).json()
        self.assertFalse(r['success'])
        self.assertNotIn('needs_override', r)
        r = self.staff.post('/desk/sign/', {'name': 'Ana Cruz', 'contact': '',
                                            'override_reason': 'Returning two books'}).json()
        self.assertTrue(r['success'], r)
        note = self.PatronLog.objects.get().entry_note
        self.assertIn('Returning two books', note)
        self.assertIn('suspended', note)
        self.assertTrue(SystemLog.objects.filter(detail__contains='Returning two books').exists())

    def test_overdue_books_and_charges_are_said_at_the_door_without_stopping_anyone(self):
        from .models import Transaction
        member = Patron.objects.create(first_name='Ben', last_name='Reyes', fullname='Ben Reyes',
                                       account_status='Active', qr_code='card-ben')
        book = Book.objects.create(title='Noli', author='Rizal', status='Borrowed', qr_code='q-noli')
        Transaction.objects.create(book=book, patron=member, transaction_type='Borrow',
                                   due_date=timezone.localdate() - timedelta(days=3))
        r = self.kiosk.post('/desk/scan/', {'code': 'card-ben'}).json()
        self.assertEqual(r['action'], 'entry')
        self.assertIn('1 overdue book', r['notice'])


class PlaceBesideTitleTests(TestCase):
    """A book put on a board stands beside the copies of its title there, not at the far end."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        for slot, title in enumerate(['Alamat', 'Noli Me Tangere', 'Zebra'], start=1):
            Book.objects.create(title=title, author='Rizal, Jose', genre='Fiction', shelf_level=self.level,
                                shelf_slot=slot, qr_code='q-%d' % slot)
        self.client = _client(admin_id=_staff(role='Admin', modules='books,shelf').admin_id, admin_role='Admin')

    def _board(self):
        from django.db.models import F
        return [(b.title, b.copy_number) for b in Book.objects.filter(shelf_level=self.level)
                .order_by(F('shelf_slot').asc(nulls_last=True))]

    def _loose_copy(self, title='Noli Me Tangere'):
        return Book.objects.create(title=title, author='Rizal, Jose', genre='Fiction', qr_code='loose-%s' % title)

    def test_add_book_puts_new_copies_beside_their_title(self):
        r = self.client.post('/admin-portal/add-book/', dict(COMPLETE_BOOK, title='Noli Me Tangere',
                                                              author='Rizal, Jose', genre='Fiction', copies='2',
                                                              shelf_level=self.level.pk),
                             HTTP_X_REQUESTED_WITH='XMLHttpRequest').json()
        self.assertTrue(r['success'], r)
        self.assertEqual(self._board(), [('Alamat', 1), ('Noli Me Tangere', 1), ('Noli Me Tangere', 2),
                                         ('Noli Me Tangere', 3), ('Zebra', 1)])
        slots = list(Book.objects.filter(shelf_level=self.level).values_list('shelf_slot', flat=True))
        self.assertEqual(sorted(slots), [1, 2, 3, 4, 5])

    def test_shelf_managers_add_books_and_move_books_do_the_same(self):
        copy = self._loose_copy()
        self.client.post('/admin-portal/assign-books-to-level/', {'shelf_level_id': self.level.pk,
                                                                  'book_ids': str(copy.pk)})
        self.assertEqual(self._board()[2], ('Noli Me Tangere', 2))
        other = self._loose_copy('Alamat')
        self.client.post('/admin-portal/move-books/', {'ids': str(other.pk), 'level': self.level.pk})
        self.assertEqual(self._board()[:2], [('Alamat', 1), ('Alamat', 2)])
        new_title = self._loose_copy('Bagong Aklat')
        self.client.post('/admin-portal/move-books/', {'ids': str(new_title.pk), 'level': self.level.pk})
        self.assertEqual(self._board()[-1], ('Bagong Aklat', 1))

    def test_the_shelving_page_and_an_edit_do_the_same(self):
        copy = self._loose_copy()
        self.client.post('/admin-portal/shelving/place/', {'placements': '%d:%d' % (copy.pk, self.level.pk)})
        self.assertEqual(self._board()[2], ('Noli Me Tangere', 2))
        moved = self._loose_copy('Zebra')
        self.client.post('/admin-portal/edit-book/%d/' % moved.pk, {
            'title': 'Zebra', 'author': 'Rizal, Jose', 'genre': 'Fiction', 'shelf_level': self.level.pk},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(self._board()[-2:], [('Zebra', 1), ('Zebra', 2)])

    def test_an_imported_copy_joins_its_title_on_the_board(self):
        rows = [['Title', 'Author', 'Genre', 'Location', 'ISBN', 'Publication Year', 'Price'],
                ['Noli Me Tangere', 'Rizal, Jose', 'Fiction', 'Shelf A Level 1', 'none', 1887, 300],
                ['Bagong Aklat', 'Cruz, Ana', 'Fiction', 'Shelf A Level 1', 'none', 2020, 300]]
        self.client.post('/admin-portal/import-books/', {'excel_file': _sheet(rows), 'copy_rows': '2'})
        self.assertEqual(self._board(), [('Alamat', 1), ('Noli Me Tangere', 1), ('Noli Me Tangere', 2),
                                         ('Zebra', 1), ('Bagong Aklat', 1)])


class HoldReasonTests(TestCase):
    """An account is only put on hold with a reason, which the patron and the desk both see."""

    def setUp(self):
        self.user = _staff(role='Admin', modules='patrons,logs')
        self.client = _client(admin_id=self.user.admin_id, admin_role='Admin', admin_fullname='Desk')
        self.patron = Patron.objects.create(first_name='Ana', last_name='Cruz', fullname='Ana Cruz',
                                            email='ana.hold@example.invalid', patron_type='Student',
                                            account_status='Active', qr_code='card-hold',
                                            password_hash=hash_password('Patron123!'))

    def _edit(self, status, reason=''):
        return self.client.post('/admin-portal/edit-patron/%d/' % self.patron.pk, {
            'first_name': 'Ana', 'last_name': 'Cruz', 'patron_type': 'Student',
            'email': 'ana.hold@example.invalid', 'account_status': status, 'hold_reason': reason})

    def test_suspending_without_a_reason_is_refused(self):
        page = self._edit('Suspended')
        self.assertContains(page, 'Give the reason for putting Ana Cruz on hold')
        self.patron.refresh_from_db()
        self.assertEqual(self.patron.account_status, 'Active')

    def test_the_reason_is_kept_logged_and_cleared_when_lifted(self):
        from .models import SystemLog
        self._edit('Suspended', '3 books unreturned since August')
        self.patron.refresh_from_db()
        self.assertEqual((self.patron.account_status, self.patron.hold_reason, self.patron.hold_set_by),
                         ('Suspended', '3 books unreturned since August', 'Desk'))
        self.assertTrue(SystemLog.objects.filter(detail='Put "Ana Cruz" on hold: 3 books unreturned since August').exists())
        self._edit('Active')
        self.patron.refresh_from_db()
        self.assertIsNone(self.patron.hold_reason)
        self.assertTrue(SystemLog.objects.filter(detail='Lifted the hold on "Ana Cruz"').exists())

    def test_the_patron_and_the_desk_are_told_why(self):
        self._edit('Suspended', '3 books unreturned since August')
        page = Client().post('/patron/login/', {'email': 'ana.hold@example.invalid', 'password': 'Patron123!'})
        self.assertContains(page, 'Your account is on hold')
        self.assertContains(page, 'Reason: 3 books unreturned since August.')
        r = self.client.post('/desk/scan/', {'code': 'card-hold'}).json()
        self.assertTrue(r['needs_override'])
        self.assertIn('On hold because: 3 books unreturned since August.', r['error'])
        self.assertEqual((r['hold_reason'], r['hold_by']), ('3 books unreturned since August', 'Desk'))
        self.assertTrue(r['hold_on'])
        kiosk = _client(admin_id=self.user.admin_id, admin_role='Admin', desk_mode=True)
        self.assertNotIn('unreturned', kiosk.post('/desk/scan/', {'code': 'card-hold'}).json()['error'])


class LetInLabelTests(TestCase):
    """The log says whether a visit was let in past a hold or a closed account, and why."""

    def test_the_note_is_read_back(self):
        from .models import PatronLog
        p = Patron.objects.create(first_name='A', last_name='B', fullname='A B', account_status='Active')
        held = PatronLog.objects.create(patron=p, entry_note='Let in by Ben although the account is suspended: Returning books')
        shut = PatronLog.objects.create(patron=p, entry_note='Let in by Ben although the account is inactive: Reading only')
        self.assertEqual(held.let_in, {'kind': 'on hold', 'by': 'Ben', 'reason': 'Returning books'})
        self.assertEqual(shut.let_in['kind'], 'closed')
        self.assertIsNone(PatronLog.objects.create(patron=p).let_in)


class DoorRequestTests(TestCase):
    """A member on hold asks at the kiosk; the librarian lets them in, with a reason, or keeps them out."""

    def setUp(self):
        from .models import DoorRequest, PatronLog
        self.DoorRequest, self.PatronLog = DoorRequest, PatronLog
        self.held = Patron.objects.create(first_name='Ana', last_name='Cruz', fullname='Ana Cruz',
                                          email='ana.door@example.invalid', account_status='Suspended',
                                          qr_code='card-door', hold_reason='2 books unreturned')
        self.user = _staff(role='Admin', modules='logs')
        self.staff = _client(admin_id=self.user.admin_id, admin_role='Admin', admin_fullname='Desk')
        self.kiosk = _client(admin_id=self.user.admin_id, admin_role='Admin', desk_mode=True)

    def _ask(self):
        r = self.kiosk.post('/desk/identify/', {'value': 'ana.door@example.invalid'}).json()
        self.assertTrue(r['can_ask'], r)
        return self.kiosk.post('/desk/ask/').json()

    def test_the_kiosk_offers_to_ask_and_the_librarian_is_told(self):
        r = self._ask()
        self.assertTrue(r['success'], r)
        self.assertIn('wait at the counter', r['message'])
        self.assertEqual(self.DoorRequest.waiting().count(), 1)
        # Asking twice is one request, not two.
        self._ask()
        self.assertEqual(self.DoorRequest.objects.count(), 1)
        page = self.staff.get('/admin-portal/log-management/')
        self.assertContains(page, 'Asking to come in (1)')
        self.assertContains(page, 'On hold: 2 books unreturned')
        door = {i['key']: i for i in self.staff.get('/portal/alerts/').json()['items']}['door']
        self.assertEqual(door['count'], 1)
        self.assertIn('Ana Cruz', door['text'])
        # The visitor form offers the same.
        r = self.kiosk.post('/desk/visitor/', {'name': 'Ana Cruz', 'contact': 'ana.door@example.invalid',
                                               'purpose': 'Study'}).json()
        self.assertTrue(r.get('can_ask'), r)

    def test_letting_in_needs_a_reason_and_makes_the_visit(self):
        ask = self.DoorRequest.objects.create(patron=self.held, purpose='Study')
        r = self.staff.post('/desk/door/', {'request_id': ask.pk}).json()
        self.assertFalse(r['success'])
        self.assertFalse(self.PatronLog.objects.filter(patron=self.held).exists())
        r = self.staff.post('/desk/door/', {'request_id': ask.pk, 'override_reason': 'Returning the books'}).json()
        self.assertTrue(r['success'], r)
        visit = self.PatronLog.objects.get(patron=self.held)
        self.assertIsNone(visit.exit_time)
        self.assertEqual(visit.let_in['reason'], 'Returning the books')
        ask.refresh_from_db()
        self.assertEqual((ask.status, ask.visit, ask.decided_by), ('Let in', visit, self.user))
        again = self.staff.post('/desk/door/', {'request_id': ask.pk, 'override_reason': 'x'}).json()
        self.assertIn('already let in', again['error'])

    def test_keeping_out_records_it_and_signs_no_one_in(self):
        ask = self.DoorRequest.objects.create(patron=self.held)
        r = self.staff.post('/desk/door/', {'request_id': ask.pk, 'decision': 'keep_out'}).json()
        self.assertTrue(r['success'], r)
        ask.refresh_from_db()
        self.assertEqual(ask.status, 'Kept out')
        self.assertFalse(self.PatronLog.objects.filter(patron=self.held).exists())
        self.assertFalse(self.DoorRequest.waiting().exists())

    def test_a_closed_account_cannot_ask(self):
        Patron.objects.create(first_name='Bea', last_name='Lim', fullname='Bea Lim', email='bea.door@example.invalid',
                              account_status='Inactive', qr_code='card-shut')
        r = self.kiosk.post('/desk/identify/', {'value': 'bea.door@example.invalid'}).json()
        self.assertNotIn('can_ask', r)
        self.assertTrue(self.kiosk.post('/desk/ask/').json()['expired'])
        self.assertFalse(self.DoorRequest.objects.exists())

    def test_an_old_request_lapses(self):
        ask = self.DoorRequest.objects.create(patron=self.held)
        self.DoorRequest.objects.filter(pk=ask.pk).update(asked_at=timezone.now() - timedelta(minutes=31))
        self.assertFalse(self.DoorRequest.waiting().exists())
        r = self.staff.post('/desk/door/', {'request_id': ask.pk, 'override_reason': 'x'}).json()
        self.assertFalse(r['success'])
        self.assertIn('tap in again', r['error'])

    def test_the_kiosk_cannot_let_anyone_in(self):
        ask = self.DoorRequest.objects.create(patron=self.held)
        r = self.kiosk.post('/desk/door/', {'request_id': ask.pk, 'override_reason': 'x'})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(self.PatronLog.objects.filter(patron=self.held).exists())


class PatronPanelTests(TestCase):
    """The side panel: who a patron is, what they hold and owe, and why they are on hold."""

    def setUp(self):
        from .models import PatronLog, SystemLog
        self.PatronLog, self.SystemLog = PatronLog, SystemLog
        rule = BorrowingRule.current()
        rule.fine_per_day = Decimal('5')
        rule.save()
        self.admin = _staff(role='Admin', modules='logs,patrons,transactions')
        self.client = _client(admin_id=self.admin.admin_id, admin_role='Admin', admin_fullname='Desk')
        self.p = Patron.objects.create(first_name='Ana', last_name='Cruz', fullname='Ana Cruz',
                                       email='ana.panel@example.invalid', contact_number='09171234567',
                                       patron_type='Student', account_status='Suspended', qr_code='card-panel',
                                       hold_reason='2 books unreturned', hold_set_by='Desk',
                                       hold_set_at=timezone.now())
        today = timezone.localdate()
        late = Book.objects.create(title='Late Book', author='A', genre='Fiction', status='Overdue', qr_code='q-late')
        lost = Book.objects.create(title='Lost Book', author='A', genre='Fiction', status='Lost', qr_code='q-lost')
        Transaction.objects.create(patron=self.p, book=late, transaction_type='Borrow', due_date=today - timedelta(days=3))
        Transaction.objects.create(patron=self.p, book=lost, transaction_type='Borrow', due_date=today - timedelta(days=9),
                                   marked_lost=True, fine_amount=Decimal('350'))
        self.PatronLog.objects.create(patron=self.p, purpose_of_visit='Study', exit_time=timezone.now(),
                                      entry_note='Let in by Desk although the account is suspended: Returning books')

    def test_the_panel_shows_the_hold_the_books_and_what_is_owed(self):
        page = self.client.get('/admin-portal/patron-panel/%d/' % self.p.pk)
        self.assertEqual(page.status_code, 200)
        for text in ('Why it is on hold', '2 books unreturned', 'Set by Desk', 'Late Book', '3 days late',
                     '15.00 in fines so far', 'Lost Book', '350.00 unpaid charge', 'Owes in all', '365.00',
                     'ana.panel@example.invalid', 'Let in &middot; on hold', 'Lift the hold', 'Change reason'):
            self.assertContains(page, text)

    def test_staff_see_it_but_only_those_with_patrons_change_the_hold(self):
        desk = _staff(role='Staff', modules='logs')
        staff = _client(admin_id=desk.admin_id, admin_role='Staff')
        page = staff.get('/admin-portal/patron-panel/%d/' % self.p.pk)
        self.assertContains(page, '2 books unreturned')
        self.assertNotContains(page, 'Lift the hold')
        r = staff.post('/admin-portal/patron-panel/%d/hold/' % self.p.pk, {'action': 'lift'})
        self.p.refresh_from_db()
        self.assertEqual(self.p.account_status, 'Suspended')
        self.assertNotEqual(r.status_code, 200)

    def test_lifting_and_changing_the_hold(self):
        url = '/admin-portal/patron-panel/%d/hold/' % self.p.pk
        self.assertFalse(self.client.post(url, {'action': 'change', 'hold_reason': ' '}).json()['success'])
        r = self.client.post(url, {'action': 'change', 'hold_reason': 'Lost-book charge unpaid'}).json()
        self.assertTrue(r['success'], r)
        self.p.refresh_from_db()
        self.assertEqual(self.p.hold_reason, 'Lost-book charge unpaid')
        self.assertTrue(self.SystemLog.objects.filter(detail='Changed why "Ana Cruz" is on hold: Lost-book charge unpaid').exists())
        r = self.client.post(url, {'action': 'lift'}).json()
        self.assertTrue(r['success'], r)
        self.p.refresh_from_db()
        self.assertEqual((self.p.account_status, self.p.hold_reason, self.p.hold_set_at), ('Active', None, None))
        self.assertTrue(self.SystemLog.objects.filter(detail='Lifted the hold on "Ana Cruz"').exists())
        self.assertFalse(self.client.post(url, {'action': 'lift'}).json()['success'])

    def test_the_pages_link_to_it(self):
        from .models import DoorRequest
        DoorRequest.objects.create(patron=self.p)
        link = 'data-patron-panel="%d"' % self.p.pk
        self.assertContains(self.client.get('/admin-portal/log-management/'), link)
        self.assertContains(self.client.get('/admin-portal/manage-patron/'), link)
        self.assertContains(self.client.get('/admin-portal/transaction/'), 'patron-panel.js')
        r = self.client.post('/desk/scan/', {'code': 'card-panel'}).json()
        self.assertEqual(r['patron_id'], self.p.pk)


class CopyNumbersOnShelvesTests(TestCase):
    """Two copies of one title on one board are told apart in Shelf Manager and Inventory."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.one = Book.objects.create(title='Noli Me Tangere', author='Rizal', genre='Fiction', qr_code='q-n1',
                                       shelf_level=self.level, shelf_slot=1)
        self.two = Book.objects.create(title='Noli Me Tangere', author='Rizal', genre='Fiction', qr_code='q-n2',
                                       shelf_level=self.level, shelf_slot=2)
        for b in (self.one, self.two):
            InventoryRecord.objects.create(book=b, source='Existing', qr_label=b.qr_code)
        self.admin = _staff(role='Admin', modules='books,shelf,inventory')
        self.client = _client(admin_id=self.admin.admin_id, admin_role='Admin')

    def test_the_board_and_the_count_sheet_carry_the_copy_number(self):
        board = self.client.get('/admin-portal/board-books/', {'level': self.level.pk}).json()['books']
        self.assertEqual([(b['title'], b['copy_number']) for b in board],
                         [('Noli Me Tangere', 1), ('Noli Me Tangere', 2)])
        sheet = self.client.get('/admin-portal/inventory/audit/sheet/', {'level': self.level.pk}).json()
        self.assertEqual(sorted(r['copy_number'] for r in sheet['books']), [1, 2])

    def test_inventory_names_the_copy_not_the_record(self):
        page = self.client.get('/admin-portal/inventory/')
        self.assertContains(page, '&middot; Copy 1')
        self.assertContains(page, '&middot; Copy 2')
        self.assertContains(page, 'Position 2')
        record = InventoryRecord.objects.get(book=self.two)
        self.assertNotContains(page, 'Copy #%d' % record.inventory_id)

    def test_the_printed_floor_plan_has_a_copy_column(self):
        from .models import FloorPlan
        from .views import _floorplan_sheet
        sheet = _floorplan_sheet(FloorPlan.objects.get())
        books = [b for sh in sheet['shelves'] for b in sh['books']]
        self.assertEqual(sorted(b['copy'] for b in books), ['Copy 1', 'Copy 2'])


class DonorToggleTests(TestCase):
    """A donation's donor is either picked from those on file or typed in full as someone new."""

    def setUp(self):
        from .models import Donation, Donor
        self.Donation, self.Donor = Donation, Donor
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')
        self.rosario = Donor.objects.create(name='Rosario Mendoza Villanueva', first_name='Rosario',
                                            middle_name='Mendoza', last_name='Villanueva',
                                            donor_type='Individual', contact_number='09175550142')

    def _give(self, **donor):
        data = {'source': 'Donation', 'donated_date': '2026-10-01',
                'items': json.dumps(_with_genre_items([{'title': 'Ibong Adarna', 'author': 'Anon'}]))}
        data.update(donor)
        return self.client.post('/admin-portal/inventory/receive/', data)

    def test_a_saved_donor_is_used_as_is(self):
        # Boxes left over from the other side of the toggle change nothing.
        self._give(donor_id=str(self.rosario.pk), donor_first_name='Someone', donor_last_name='Else',
                   donor_contact_number='09990000000')
        gift = self.Donation.objects.get()
        self.assertEqual((gift.donor, gift.donor_name), (self.rosario, 'Rosario Mendoza Villanueva'))
        self.rosario.refresh_from_db()
        self.assertEqual(self.rosario.contact_number, '09175550142')
        self.assertEqual(self.Donor.objects.count(), 1)

    def test_a_donor_no_longer_on_file_is_refused(self):
        self._give(donor_id='99999')
        self.assertFalse(self.Donation.objects.exists())

    def test_a_new_donor_is_typed_in_full(self):
        self._give(donor_type='Individual', donor_first_name='Carlo', donor_middle_name='Bautista',
                   donor_last_name='Dimaculangan', donor_contact_number='09565550288')
        carlo = self.Donor.objects.get(first_name='Carlo')
        self.assertEqual((carlo.name, carlo.contact_number), ('Carlo Bautista Dimaculangan', '09565550288'))

    def test_both_receiving_pages_offer_the_toggle(self):
        from .views import _donor_options
        self._give(donor_id=str(self.rosario.pk))
        saved = {d['name']: d for d in _donor_options()}['Rosario Mendoza Villanueva']
        self.assertEqual((saved['id'], saved['contact'], saved['gifts'], saved['last']),
                         (self.rosario.pk, '09175550142', 1, 'Oct 01, 2026'))
        staff = _client(admin_id=_staff(role='Staff', modules='inventory').admin_id, admin_role='Staff')
        for page in (self.client.get('/admin-portal/inventory/'), staff.get('/library-staff/receiving/')):
            self.assertContains(page, 'data-donor-mode="new"')
            self.assertContains(page, 'Existing donor <span class="opacity-70">(1)</span>', html=False)
            self.assertContains(page, 'name="donor_id"')



class DirectPrintTests(TestCase):
    """New copies go straight to the print settings, from receiving as from cataloguing."""

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')

    def test_received_copies_open_for_printing_once(self):
        self.client.post('/admin-portal/inventory/receive/', {
            'source': 'Purchase', 'supplier': 'Shop',
            'items': json.dumps(_with_genre_items([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 2}]))})
        new = sorted(Book.objects.filter(title='Ibong Adarna').values_list('book_id', flat=True))
        self.assertEqual(len(new), 2)
        page = self.client.get('/admin-portal/inventory/')
        self.assertEqual(sorted(page.context['print_next_ids']), new)
        self.assertContains(page, 'printNextData')
        self.assertFalse(self.client.get('/admin-portal/inventory/').context['print_next_ids'])

    def test_receiving_staff_may_print_and_get_the_same(self):
        staff = _client(admin_id=_staff(role='Staff', modules='inventory').admin_id, admin_role='Staff')
        staff.post('/admin-portal/inventory/receive/', {
            'source': 'Purchase', 'supplier': 'Shop',
            'items': json.dumps(_with_genre_items([{'title': 'Florante', 'author': 'Balagtas'}]))})
        book = Book.objects.get(title='Florante')
        page = staff.get('/library-staff/receiving/')
        self.assertEqual(page.context['print_next_ids'], [book.book_id])
        self.assertContains(page, 'label-print.js')
        sheet = staff.get('/admin-portal/book-qr-labels/', {'ids': str(book.book_id)})
        self.assertEqual(sheet['Content-Type'], 'application/pdf')

    def test_shelving_offers_only_books_still_without_a_label(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel
        plan = FloorPlan.objects.create(name='G', floor_number=1, is_active=True)
        shelf = Shelf.objects.create(room=Room.objects.create(floor_plan=plan, name='R', map_x=0, map_y=0),
                                     name='Shelf A', map_x=0, map_y=0)
        level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        done = Book.objects.create(title='Done', author='A', genre='Fiction', qr_code='q-done',
                                   label_printed_at=timezone.now())
        fresh = Book.objects.create(title='Fresh', author='A', genre='Fiction', qr_code='q-fresh')
        r = self.client.post('/admin-portal/shelving/place/', {
            'placements': '%d:%d,%d:%d' % (done.pk, level.pk, fresh.pk, level.pk)}).json()
        self.assertEqual(r['ids'], [fresh.pk])
        r = self.client.post('/admin-portal/shelving/place/', {'placements': '%d:%d' % (done.pk, level.pk)}).json()
        self.assertEqual(r['ids'], [])
        self.assertIn('label already', r['message'])


class InventoryValidationTests(TestCase):
    """The inventory's rules, checked on the server as well as in the browser."""

    def setUp(self):
        self.client = _client(admin_id=_staff(role='Admin').admin_id, admin_role='Admin')

    def _receive(self, **data):
        base = {'source': 'Purchase', 'supplier': 'National Book Store',
                'items': json.dumps(_with_genre_items([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 2}]))}
        base.update(data)
        return self.client.post('/admin-portal/inventory/receive/', base)

    def test_reasons_need_real_words(self):
        from .inventory_rules import reason_problem
        for bad in ('', 'x', 'asdasd', 'aaaaaaa', '12345678', 'Damaged'):
            self.assertTrue(reason_problem(bad), bad)
        self.assertIsNone(reason_problem('Water damage found during shelving'))

    def test_a_shipment_needs_its_supplier(self):
        self._receive(supplier=' ')
        self.assertFalse(InventoryRecord.objects.exists())
        self._receive()
        self.assertEqual(InventoryRecord.objects.count(), 2)

    def test_a_donation_cannot_be_dated_in_the_future(self):
        tomorrow = (timezone.localdate() + timedelta(days=1)).isoformat()
        self._receive(source='Donation', donor_type='Individual', donor_first_name='Ana', donor_last_name='Cruz',
                      donated_date=tomorrow)
        self.assertFalse(InventoryRecord.objects.exists())

    def test_a_price_is_at_least_one_peso(self):
        self._receive(items=json.dumps(_with_genre_items([{'title': 'Florante', 'author': 'Balagtas', 'price': '0'}])))
        self.assertFalse(InventoryRecord.objects.exists())

    def test_a_delivery_received_twice_needs_confirming(self):
        self._receive()
        again = self.client.post('/admin-portal/inventory/receive/check/', {
            'source': 'Purchase', 'supplier': 'national book store',
            'items': json.dumps([{'title': 'Ibong Adarna', 'author': 'Anon', 'quantity': 1}])}).json()
        self.assertEqual(len(again['duplicates']), 1)
        self.assertIn('2 copies already received', again['duplicates'][0])
        self._receive()
        self.assertEqual(InventoryRecord.objects.count(), 2, 'not received twice without confirming')
        self._receive(confirm_duplicate='1')
        self.assertEqual(InventoryRecord.objects.count(), 4)

    def test_a_condition_change_with_a_throwaway_reason_is_refused(self):
        self._receive()
        record = InventoryRecord.objects.first()
        self.client.post('/admin-portal/inventory/condition/', {
            'inventory_id': record.pk, 'condition': 'Damaged', 'reason': 'x', 'admin_password': 'DeskTest123'})
        record.refresh_from_db()
        self.assertEqual(record.condition, 'Good')

    def test_the_forms_check_as_you_type(self):
        page = self.client.get('/admin-portal/inventory/')
        for marker in ('inventory-checks.js', 'data-check="reason"', 'data-check="supplier"', 'data-check="donated"',
                       'data-check="price"', 'data-checked-form', "url 'receive_check'" if False else 'receive/check/'):
            self.assertContains(page, marker)


class AuditTrailStateTests(TestCase):
    """A count's follow-up agrees with the book as it is now, whatever happened to it since."""

    def setUp(self):
        from .models import FloorPlan, Room, Shelf, ShelfLevel, StockAudit
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf_a = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.a1 = ShelfLevel.objects.create(shelf=self.shelf_a, level_number=1)
        self.a2 = ShelfLevel.objects.create(shelf=self.shelf_a, level_number=2)
        self.audit = StockAudit.objects.create(shelf=self.shelf_a, shelf_name='Shelf A',
                                               audited_at=timezone.now() - timedelta(days=5))

    def _line(self, result, **book):
        from .models import StockAuditLine
        b = Book.objects.create(title=book.pop('title', result), author='A', shelf_level=self.a1, **book)
        line = StockAuditLine.objects.create(audit=self.audit, book=b, title=b.title, result=result,
                                             status_then=b.status, condition_then=b.condition or 'Good',
                                             location_then=b.location_short)
        return b, line

    def test_a_book_gone_since_settles_any_finding_truthfully(self):
        from .audittrail import outcome
        damaged, d_line = self._line('Damaged', condition='Damaged')
        stray, s_line = self._line('Stray')
        damaged.archive('Desk', 'Deaccessioned: torn beyond use')
        stray.status = 'Lost'
        stray.save()
        self.assertEqual(outcome(d_line), ('Removed from the records', False))
        self.assertEqual(outcome(s_line), ('Reported lost', False))

    def test_a_loan_reported_lost_is_not_called_checked_in(self):
        from .audittrail import events_since, outcome
        book, line = self._line('Here on loan', status='Borrowed')
        reader = Patron.objects.create(first_name='Ana', last_name='Cruz', fullname='Ana Cruz')
        tx = Transaction.objects.create(book=book, patron=reader, transaction_type='Borrow',
                                        due_date=timezone.localdate(), marked_lost=True, fine_amount=Decimal('350'))
        from .models import SystemLog
        SystemLog.objects.create(action='Process', entity_type='Transaction', entity_id=str(tx.transaction_id),
                                 detail='Marked lost')
        book.status = 'Lost'
        book.save()
        self.assertEqual(outcome(line), ('Reported lost', False))
        self.assertTrue(any('Reported lost by Ana Cruz' in text and '350.00' in text for _d, text in events_since(line)))

    def test_a_move_after_the_count_is_explained(self):
        from .audittrail import events_since
        book, line = self._line('Found')
        book.shelf_level = self.a2
        book.save()
        trail = events_since(line)
        self.assertEqual(trail[-1][0], None)
        self.assertIn('moved since the count', trail[-1][1])
