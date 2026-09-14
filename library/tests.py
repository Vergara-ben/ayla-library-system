"""Smoke tests."""

from datetime import date, datetime, timedelta
import io
import json
import os
from io import StringIO
from unittest import mock
from uuid import uuid4

from django.test import TestCase, Client
from django.utils import timezone

from .auth_utils import hash_password, password_length_error
from .models import (
    BLEBeacon, Book, BorrowingRule, Door, FloorPlan, LoginAttempt, Obstacle, Patron, Room,
    InventoryRecord, Obstacle, PatronLog, Shelf, ShelfLevel, Stairway, SystemLog,
    Transaction, User,
    Waypoint,
    WaypointConnection,
)


class _FakeRequest:
    """Just enough request for a middleware or an audit helper to work on."""

    def __init__(self, path='/probe/', method='GET', session=None):
        self.path = path
        self.method = method
        self.session = {} if session is None else session


def _admin(email='smoke-admin@example.invalid', modules=''):
    return User.objects.create(
        fullname='Smoke Admin', email=email,
        password_hash=hash_password('SmokeTest123'),
        role='Admin', account_status='Active', modules=modules,
    )


def _signed_in(user):
    c = Client()
    s = c.session
    s['admin_id'] = user.admin_id
    s['admin_fullname'] = user.fullname
    s['admin_role'] = user.role
    s.save()
    return c


class PagesRenderTests(TestCase):
    """Every admin page answers 200 for someone allowed to see it."""

    def setUp(self):
        self.user = _admin(modules='transactions,books,logs,patrons,shelf,'
                                   'donations,inventory')
        self.client = _signed_in(self.user)

    def test_admin_pages_render(self):
        pages = [
            '/admin-portal/dashboard/',
            '/admin-portal/management/',
            '/admin-portal/transaction/',
            '/admin-portal/log-management/',
            '/admin-portal/manage-patron/',
            '/admin-portal/shelf-manager/',
            '/admin-portal/floorplan-management/',
            '/admin-portal/indoor-map/',
            '/admin-portal/reports/',
            '/admin-portal/users/',
            '/admin-portal/activity-logs/',
            '/admin-portal/announcement-management/',
            '/admin-portal/donation-management/',
            '/admin-portal/inventory/',
            '/admin-portal/borrowing-rules/',
        ]
        for url in pages:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_every_report_builds_and_exports(self):
        from .reports import REPORT_TYPES, build_report, render_report_pdf, render_report_excel
        start = timezone.localdate() - timedelta(days=30)
        end = timezone.localdate()
        for key, _label in REPORT_TYPES:
            with self.subTest(report=key):
                report = build_report(key, start, end)
                self.assertIn('columns', report)
                # Both exporters have broken on a new report key before.
                self.assertTrue(render_report_pdf(report).getvalue())
                self.assertTrue(render_report_excel(report).getvalue())


class AuthGuardTests(TestCase):
    """The guards that a misplaced decorator silently removes."""

    def setUp(self):
        self.user = _admin(modules='patrons')

    def test_admin_pages_reject_anonymous(self):
        anon = Client()
        for url in ['/admin-portal/dashboard/', '/admin-portal/shelf-manager/',
                    '/admin-portal/announcement-management/',
                    '/admin-portal/reports/']:
            with self.subTest(url=url):
                self.assertEqual(anon.get(url).status_code, 302)

    def test_credential_documents_are_not_public(self):
        """Scanned IDs were once served to anyone who had the URL."""
        anon = Client()
        self.assertEqual(anon.get('/media/credentials/anything.png').status_code, 404)
        self.assertEqual(anon.get('/patron-id/anything.png').status_code, 302)

    def test_credential_route_refuses_path_traversal(self):
        client = _signed_in(self.user)
        for probe in ['../../db.sqlite3', '../../../manage.py', 'nope.png']:
            with self.subTest(probe=probe):
                self.assertEqual(client.get('/patron-id/' + probe).status_code, 404)


class LoginSecurityTests(TestCase):
    """Throttling, session rotation, and the absence of an enumeration oracle."""

    def setUp(self):
        self.user = _admin()

    def test_lockout_after_repeated_failures(self):
        client = Client()
        for _ in range(LoginAttempt.MAX_FAILURES):
            client.post('/admin-portal/login/',
                        {'email': self.user.email, 'password': 'wrong'})
        response = client.post('/admin-portal/login/',
                               {'email': self.user.email, 'password': 'wrong'})
        self.assertContains(response, 'Too many failed')

    def test_correct_password_clears_the_run(self):
        client = Client()
        client.post('/admin-portal/login/', {'email': self.user.email, 'password': 'wrong'})
        client.post('/admin-portal/login/',
                    {'email': self.user.email, 'password': 'SmokeTest123'})
        self.assertFalse(
            LoginAttempt.objects.filter(scope='admin',
                                        identifier=self.user.email.lower()).exists())

    def test_session_id_rotates_on_login(self):
        client = Client()
        client.get('/admin-portal/login/')
        session = client.session
        session['planted'] = 'attacker'
        session.save()
        before = session.session_key

        client.post('/admin-portal/login/',
                    {'email': self.user.email, 'password': 'SmokeTest123'})

        self.assertNotEqual(before, client.session.session_key)
        self.assertNotIn('planted', client.session)
        self.assertEqual(client.session.get('admin_id'), self.user.admin_id)

    def test_unknown_and_known_addresses_answer_identically(self):
        client = Client()
        unknown = client.post('/admin-portal/login/',
                              {'email': 'nobody@example.invalid', 'password': 'x'})
        known = Client().post('/admin-portal/login/',
                              {'email': self.user.email, 'password': 'x'})
        self.assertContains(unknown, 'Invalid email or password')
        self.assertContains(known, 'Invalid email or password')
        self.assertNotContains(known, 'suspended')


class PasswordPolicyTests(TestCase):
    def test_weak_passwords_are_refused(self):
        for weak in ['Ab1', 'aaaaaaaaaa', '12345678', 'password1']:
            with self.subTest(password=weak):
                self.assertIsNotNone(password_length_error(weak))

    def test_a_reasonable_password_passes(self):
        self.assertIsNone(password_length_error('Kalachuchi77'))

    def test_the_current_password_cannot_be_reused(self):
        current = hash_password('Kalachuchi77')
        self.assertIsNotNone(
            password_length_error('Kalachuchi77', current_hash=current))


class BorrowingTests(TestCase):
    """The basket either commits whole or not at all, and a copy goes out once."""

    def setUp(self):
        self.admin = _admin(modules='transactions')
        self.client = _signed_in(self.admin)
        self.patron = Patron.objects.create(
            fullname='Smoke Patron', email='smoke-patron@example.invalid',
            password_hash=hash_password('SmokeTest123'),
            patron_type='Student', account_status='Active',
        )
        self.rule = BorrowingRule.current()
        self.rule.max_books_per_patron = 5
        self.rule.save()
        self.free = [Book.objects.create(title=f'Smoke Book {i}', author='Tester',
                                         status='Available') for i in range(2)]
        self.taken = Book.objects.create(title='Smoke Taken', author='Tester',
                                         status='Borrowed')

    def _post(self, kind, books):
        return self.client.post('/admin-portal/process-transaction/', {
            'transaction_type': kind,
            'patron_id': self.patron.patron_id,
            'book_ids': [b.book_id for b in books],
        }).json()

    def test_borrow_and_return_round_trip(self):
        self.assertTrue(self._post('Borrow', self.free)['success'])
        for book in self.free:
            book.refresh_from_db()
            self.assertEqual(book.status, 'Borrowed')

        self.assertTrue(self._post('Return', self.free)['success'])
        for book in self.free:
            book.refresh_from_db()
            self.assertEqual(book.status, 'Available')

    def test_a_failing_book_rolls_the_whole_basket_back(self):
        before = Transaction.objects.count()
        result = self._post('Borrow', [self.free[0], self.taken, self.free[1]])

        self.assertFalse(result['success'])
        for book in self.free:
            book.refresh_from_db()
            self.assertEqual(book.status, 'Available',
                             'a partial basket was committed')
        self.assertEqual(Transaction.objects.count(), before,
                         'transaction rows survived a rolled-back basket')

    def test_a_borrowed_copy_cannot_be_borrowed_again(self):
        self.assertFalse(self._post('Borrow', [self.taken])['success'])
        self.assertEqual(
            Transaction.objects.filter(book=self.taken,
                                       return_date__isnull=True).count(), 0)


class DeskModeTests(TestCase):
    """Arming the desk confines the browser to the log page."""

    def setUp(self):
        from .desk import DESK_SESSION_KEY
        self.user = _admin(modules='logs,patrons')
        self.client = _signed_in(self.user)
        session = self.client.session
        session[DESK_SESSION_KEY] = True
        session.save()

    def test_other_portal_pages_redirect_to_the_log(self):
        response = self.client.get('/admin-portal/manage-patron/')
        self.assertEqual(response.status_code, 302)
        self.assertIn('log-management', response['Location'])


class NoIdleSignOutTests(TestCase):
    """A signed-in session is not closed for sitting quiet."""

    def test_the_idle_middleware_is_not_installed(self):
        from django.conf import settings
        self.assertNotIn('library.middleware.IdleSessionTimeoutMiddleware',
                         settings.MIDDLEWARE)

    def test_a_session_quiet_for_hours_stays_signed_in(self):
        client = _signed_in(_admin())
        session = client.session
        # Leftover timestamp from older sessions.
        session['_last_seen'] = timezone.now().timestamp() - 6 * 3600
        session.save()
        self.assertEqual(client.get('/admin-portal/dashboard/').status_code, 200)
        self.assertIn('admin_id', client.session)

    def test_the_session_age_is_not_a_hidden_one_hour_sign_out(self):
        """Nothing re-saves a session during use, so a short age is a timer."""
        from django.conf import settings
        self.assertGreaterEqual(settings.SESSION_COOKIE_AGE, 7 * 24 * 3600)


class ImportGuardTests(TestCase):
    def test_only_real_workbooks_are_parsed(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .views import check_import_upload

        self.assertIsNotNone(check_import_upload(None))
        self.assertIsNotNone(check_import_upload(
            SimpleUploadedFile('x.csv', b'PK\x03\x04')))
        self.assertIsNotNone(check_import_upload(
            SimpleUploadedFile('x.xlsx', b'a,b,c')))


class ContentSecurityPolicyTests(TestCase):
    def test_the_header_is_sent(self):
        response = Client().get('/admin-portal/login/')
        policy = response.headers.get('Content-Security-Policy', '')
        self.assertIn("frame-ancestors 'none'", policy)
        self.assertIn("connect-src 'self'", policy)


class DailyMaintenanceTests(TestCase):
    """The two daily jobs run from one scheduled-task slot."""

    def test_both_jobs_run(self):
        from io import StringIO
        from django.core.management import call_command

        out = StringIO()
        call_command('daily_maintenance', dry_run=True, stdout=out)
        printed = out.getvalue()
        self.assertIn('Overdue loans', printed)
        self.assertIn('Open visits', printed)

    def test_one_job_failing_does_not_stop_the_other(self):
        from io import StringIO
        from unittest.mock import patch
        from django.core.management import call_command

        real = call_command

        def flaky(name, *args, **kwargs):
            if name == 'send_overdue_notifications':
                raise RuntimeError('simulated mail failure')
            return real(name, *args, **kwargs)

        out, err = StringIO(), StringIO()
        with patch('library.management.commands.daily_maintenance.call_command',
                   side_effect=flaky):
            with self.assertRaises(SystemExit):
                call_command('daily_maintenance', stdout=out, stderr=err)

        # The second job still ran and the failure was reported.
        self.assertIn('Open visits', out.getvalue())
        self.assertIn('simulated mail failure', err.getvalue())


class EmailTransportTests(TestCase):
    """Mail can leave over SMTP or over an HTTPS API without the code changing."""

    def test_send_goes_through_whatever_backend_is_configured(self):
        from django.core import mail
        from .emails import send_email

        with self.settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend'):
            mail.outbox = []
            self.assertTrue(send_email('Subject', 'Body', 'someone@example.invalid'))
            self.assertEqual(len(mail.outbox), 1)
            self.assertEqual(mail.outbox[0].subject, 'Subject')

    def test_a_send_failure_never_raises_into_the_caller(self):
        """A mail failure must not stop a loan."""
        from .emails import send_email

        with self.settings(EMAIL_BACKEND='django.core.mail.backends.smtp.EmailBackend',
                           EMAIL_HOST='127.0.0.1', EMAIL_PORT=1, EMAIL_TIMEOUT=1):
            self.assertFalse(send_email('Subject', 'Body', 'someone@example.invalid'))


class CrossFloorRoutingTests(TestCase):
    """Routing to a shelf that is not on the floor the patron is standing on."""

    def setUp(self):
        self.f1 = FloorPlan.objects.create(
            name='Ground', floor_number=1, is_active=True, pixels_per_meter=10)
        self.f2 = FloorPlan.objects.create(
            name='Upper', floor_number=2, is_active=True, pixels_per_meter=10)

        def corridor(plan, y):
            wps = [Waypoint.objects.create(floor_plan=plan, map_x=x, map_y=y,
                                           label='%s-%d' % (plan.name, x))
                   for x in (100, 300, 500)]
            for a, b in zip(wps, wps[1:]):
                WaypointConnection.objects.create(
                    waypoint_from=a, waypoint_to=b,
                    distance=abs(b.map_x - a.map_x))
            return wps

        self.lower = corridor(self.f1, 200)
        self.upper = corridor(self.f2, 200)

        # The book, upstairs at the far end of the corridor.
        room = Room.objects.create(floor_plan=self.f2, name='Upstairs Reading',
                                   map_x=500, map_y=200)
        self.shelf = Shelf.objects.create(room=room, name='Shelf U3',
                                          map_x=500, map_y=210)

        self.near = Stairway.objects.create(
            floor_plan=self.f1, kind='Stairs', name='North Stairs',
            geometry=[[80, 240], [120, 240], [120, 300], [80, 300]],
            map_x=100, map_y=270, direction='up', connects_to=self.f2)
        self.landing = Stairway.objects.create(
            floor_plan=self.f2, kind='Stairs', name='North Stairs',
            geometry=[[80, 240], [120, 240], [120, 300], [80, 300]],
            map_x=100, map_y=270, direction='down', connects_to=self.f1)
        # Real stairs, never linked -- the wrong answer that has to be avoided.
        self.far = Stairway.objects.create(
            floor_plan=self.f1, kind='Stairs', name='South Stairs',
            geometry=[[480, 240], [520, 240], [520, 300], [480, 300]],
            map_x=500, map_y=270, direction='up')

        self.client = Client()

    def _route(self, **params):
        params.setdefault('start_x', 100)
        params.setdefault('start_y', 200)
        r = self.client.get('/patron/navigation-route/', params)
        self.assertEqual(r.status_code, 200)
        return r.json()

    def test_same_floor_route_is_untouched(self):
        """The floor the stairs are on still routes as it always did."""
        room = Room.objects.create(floor_plan=self.f1, name='Ground Reading',
                                   map_x=500, map_y=200)
        shelf = Shelf.objects.create(room=room, name='Shelf G3',
                                     map_x=500, map_y=210)
        data = self._route(target_shelf_id=shelf.shelf_id,
                           floor=self.f1.floor_plan_id)

        self.assertTrue(data['success'])
        self.assertNotIn('cross_floor', data)
        # Three waypoints, no stairway spliced in, and the plain corridor length.
        self.assertEqual([p['waypoint_id'] for p in data['route']],
                         [w.waypoint_id for w in self.lower])
        self.assertEqual(data['distance'], 400)
        self.assertEqual(len(data['legs']), 1)

    def test_route_crosses_the_floor_and_reaches_the_shelf(self):
        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)

        self.assertTrue(data['success'])
        self.assertTrue(data['cross_floor'])
        self.assertFalse(data['stairs_only'])
        self.assertEqual(data['leg'], 'full')
        self.assertEqual(data['floors_crossed'], 2)

        # One continuous path: down the graph, up the stairs, on to the shelf.
        floors = [l['floor_plan_id'] for l in data['legs']]
        self.assertEqual(floors, [self.f1.floor_plan_id, self.f2.floor_plan_id])
        self.assertEqual(data['legs'][-1]['points'][-1]['waypoint_id'],
                         self.upper[-1].waypoint_id)
        self.assertEqual(data['target_shelf']['shelf_id'], self.shelf.shelf_id)

    def test_it_uses_the_staircase_that_is_actually_linked(self):
        """Starting beside the unlinked far stairs must not tempt it."""
        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id,
                           start_x=500, start_y=200)

        self.assertTrue(data['success'])
        self.assertEqual(data['via_stairway']['stairway_id'], self.near.stairway_id)
        self.assertTrue(data['via_stairway']['linked'])

    def test_route_is_the_leg_for_the_floor_being_shown(self):
        """`route` follows the map, so the same request answers each floor."""
        below = self._route(target_shelf_id=self.shelf.shelf_id,
                            from_floor=self.f1.floor_plan_id)
        above = self._route(target_shelf_id=self.shelf.shelf_id,
                            from_floor=self.f2.floor_plan_id,
                            start_x=100, start_y=200)

        self.assertTrue(all(p['floor_plan_id'] == self.f1.floor_plan_id
                            for p in below['route']))
        self.assertTrue(all(p['floor_plan_id'] == self.f2.floor_plan_id
                            for p in above['route']))
        # Standing on the target floor is not a crossing at all.
        self.assertNotIn('cross_floor', above)

    def test_total_distance_includes_the_climb(self):
        """The walk is longer than the two floors laid end to end."""
        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)
        flat = sum(l['distance'] for l in data['legs'])
        self.assertGreater(data['distance'], flat)
        self.assertEqual(data['leg_distance'], data['legs'][0]['distance'])

    def test_instruction_names_the_stairs_and_the_floor(self):
        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)
        self.assertEqual(
            data['instruction'],
            'Take the stairs at North Stairs up to the %s, then follow the map '
            'from there.' % self.f2.floor_label)

    def test_unlinked_stairs_still_walk_the_patron_to_them(self):
        """With nothing joining the floors, guidance beats an error."""
        self.near.connects_to = None
        self.near.save(update_fields=['connects_to'])
        self.landing.connects_to = None
        self.landing.save(update_fields=['connects_to'])

        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)

        self.assertTrue(data['success'])
        self.assertTrue(data['stairs_only'])
        self.assertEqual(data['leg'], 'to_stairs')
        self.assertEqual(data['via_stairway']['stairway_id'], self.near.stairway_id)
        self.assertFalse(data['via_stairway']['linked'])
        self.assertEqual(data['target_floor']['floor_plan_id'], self.f2.floor_plan_id)

    def test_linking_only_the_upward_flight_is_enough(self):
        """The landing forgot to name the floor below."""
        self.landing.connects_to = None
        self.landing.save(update_fields=['connects_to'])

        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)

        self.assertTrue(data['success'])
        self.assertFalse(data['stairs_only'])
        self.assertEqual(data['floors_crossed'], 2)
        self.assertEqual(data['via_stairway']['stairway_id'], self.near.stairway_id)

    def test_a_distant_stairway_is_not_the_same_staircase(self):
        """Half a floor away is a different stairwell, not the other end."""
        self.landing.connects_to = None
        self.landing.map_x = 900          # nowhere near the flight below
        self.landing.save(update_fields=['connects_to', 'map_x'])

        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)

        # No pairing, so no through route -- and the honest fallback instead.
        self.assertTrue(data['success'])
        self.assertTrue(data['stairs_only'])
        self.assertEqual(data['leg'], 'to_stairs')

    def test_no_stairs_at_all_says_so_plainly(self):
        Stairway.objects.all().delete()
        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)

        self.assertFalse(data['success'])
        self.assertTrue(data['cross_floor'])
        self.assertIn('stairway', data['error'].lower())
        self.assertEqual(data['target_floor']['label'], self.f2.floor_label)

    def test_a_floor_with_stairs_but_no_corridor_explains_itself(self):
        """Half-drawn upstairs: stairs placed, waypoints not."""
        Waypoint.objects.filter(floor_plan=self.f2).delete()

        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id)
        self.assertFalse(data['success'])
        self.assertIn('waypoint', data['error'].lower())

        # And from the other side: standing on the floor that has none.
        blank = self._route(target_shelf_id=self.shelf.shelf_id,
                            from_floor=self.f2.floor_plan_id)
        self.assertFalse(blank['success'])
        self.assertIn('waypoint', blank['error'].lower())

    def test_a_lift_is_preferred_over_stairs_when_it_is_nearer(self):
        """Stairs and lift are one graph, so the shorter walk simply wins."""
        Stairway.objects.create(
            floor_plan=self.f1, kind='Elevator', name='West Lift',
            geometry=[[480, 240], [520, 240], [520, 280], [480, 280]],
            map_x=500, map_y=260, direction='both', connects_to=self.f2)
        Stairway.objects.create(
            floor_plan=self.f2, kind='Elevator', name='West Lift',
            geometry=[[480, 240], [520, 240], [520, 280], [480, 280]],
            map_x=500, map_y=260, direction='both', connects_to=self.f1)

        data = self._route(target_shelf_id=self.shelf.shelf_id,
                           from_floor=self.f1.floor_plan_id,
                           start_x=500, start_y=200)

        self.assertTrue(data['success'])
        self.assertEqual(data['via_stairway']['kind'], 'Elevator')
        self.assertTrue(data['instruction'].startswith('Take the lift at West Lift'))


class QRLabelSheetTests(TestCase):
    """The printable QR label sheet."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        self.copies = [
            Book.objects.create(title='Noli Me Tangere', author='Jose Rizal',
                                qr_code=str(uuid4()), status='Available')
            for _ in range(3)
        ]
        self.single = Book.objects.create(title='Florante at Laura',
                                          author='Francisco Balagtas',
                                          qr_code=str(uuid4()), status='Available')

    def _sheet(self, ids, **params):
        params['ids'] = ','.join(str(i) for i in ids)
        return self.client.get('/admin-portal/book-qr-labels/', params)

    def test_sheet_is_a_pdf(self):
        r = self._sheet([b.book_id for b in self.copies])
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertIn('attachment', r['Content-Disposition'])
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_every_copy_gets_its_own_code(self):
        """Three copies of one title are three different stickers."""
        codes = {b.qr_code for b in self.copies}
        self.assertEqual(len(codes), 3)

    def test_copy_numbers_count_the_whole_catalogue(self):
        """Printing two of three copies still reads 2 of 3, not 1 of 2."""
        from .labels import copy_numbers
        subset = self.copies[1:]
        numbers = copy_numbers(subset)
        self.assertEqual(numbers[subset[0].book_id], (2, 3))
        self.assertEqual(numbers[subset[1].book_id], (3, 3))

    def test_a_single_copy_is_not_numbered(self):
        from .labels import copy_numbers
        self.assertNotIn(self.single.book_id, copy_numbers([self.single]))

    def test_a_book_with_no_qr_gets_one_minted(self):
        """Otherwise it is silently left off the sheet and nobody notices."""
        bare = Book.objects.create(title='Ibong Adarna', author='Anonymous',
                                   qr_code=None, status='Available')
        r = self._sheet([bare.book_id])
        self.assertEqual(r.status_code, 200)
        bare.refresh_from_db()
        self.assertTrue(bare.qr_code)

    def test_requested_size_is_honoured_and_clamped(self):
        from .labels import build_label_sheet, MAX_QR_MM, MIN_QR_MM
        _pdf, big = build_label_sheet(self.copies, qr_mm=9999)
        _pdf, small = build_label_sheet(self.copies, qr_mm=0.1)
        self.assertEqual(big['qr_mm'], MAX_QR_MM)
        self.assertEqual(small['qr_mm'], MIN_QR_MM)

    def test_a_smaller_label_fits_more_on_the_page(self):
        from .labels import build_label_sheet
        _pdf, small = build_label_sheet(self.copies, qr_mm=15)
        _pdf, large = build_label_sheet(self.copies, qr_mm=40)
        self.assertGreater(small['per_page'], large['per_page'])
        self.assertGreater(small['per_page'], 50)      # the grid is worth having

    def test_skip_pushes_the_first_label_down_the_sheet(self):
        from .labels import build_label_sheet
        _pdf, plain = build_label_sheet(self.copies, qr_mm=25)
        _pdf, offset = build_label_sheet(self.copies, qr_mm=25,
                                         skip=plain['per_page'])
        self.assertEqual(plain['pages'], 1)
        self.assertEqual(offset['pages'], 2)

    def test_no_selection_is_refused(self):
        r = self.client.get('/admin-portal/book-qr-labels/')
        self.assertFalse(r.json()['success'])

    def test_too_many_is_refused(self):
        from .views import MAX_LABELS_PER_SHEET
        r = self._sheet(range(1, MAX_LABELS_PER_SHEET + 5))
        self.assertFalse(r.json()['success'])

    def test_it_is_gated(self):
        stranger = Client()
        r = stranger.get('/admin-portal/book-qr-labels/',
                         {'ids': str(self.single.book_id)})
        self.assertNotEqual(r.status_code, 200)

    def test_import_reports_what_it_created(self):
        """So the label sheet can target exactly that delivery."""
        import openpyxl
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Title', 'Author', 'Quantity'])
        ws.append(['Mga Ibong Mandaragit', 'Amado V. Hernandez', 2])
        buf = BytesIO()
        wb.save(buf)

        upload = SimpleUploadedFile(
            'delivery.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r = self.client.post('/admin-portal/import-books/', {'excel_file': upload})
        data = r.json()
        self.assertTrue(data['success'], data)
        self.assertEqual(len(data['created_ids']), 2)
        self.assertFalse(data['created_truncated'])

        sheet = self._sheet(data['created_ids'])
        self.assertTrue(sheet.content.startswith(b'%PDF'))


class EveryAdminEndpointIsGuardedTests(TestCase):
    """No /admin-portal/ view is missing its sign-in guard."""

    # Reached before sign-in by design.
    PUBLIC = {
        'admin_login', 'admin_signin', 'admin_forgot_password', 'admin_logout',
        'staff_login', 'staff_forgot_password',
        'desk_sign', 'desk_sign_out', 'desk_scan', 'desk_unlock', 'desk_arm',
        'desk_identify', 'desk_visit', 'desk_visitor',
    }
    GUARDS = {
        'admin_login_required', 'admin_only_required', 'admin_module_required',
        'granted_module_required', 'admin_or_module_required',
        'patron_login_required',
    }

    def _admin_routes(self):
        """(route, name, view callable) for every admin-portal URL."""
        from library import urls as library_urls

        found = []
        for pattern in library_urls.urlpatterns:
            route = str(getattr(pattern, 'pattern', ''))
            name = getattr(pattern, 'name', None)
            if not route.startswith('admin-portal/') or name in self.PUBLIC:
                continue
            found.append((route, name, getattr(pattern, 'callback', None)))
        return found

    def test_every_admin_view_carries_a_guard(self):
        """Checked on the decorated object, so a stranded decorator shows up."""
        import ast
        from pathlib import Path

        here = Path(__file__).resolve().parent
        decorated = {}
        for module in ('views.py', 'desk.py', 'chat.py'):
            tree = ast.parse((here / module).read_text(encoding='utf-8'))
            for node in tree.body:
                if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    continue
                # Count decorators on inner functions too.
                names = set()
                for inner in ast.walk(node):
                    if not isinstance(inner, (ast.FunctionDef, ast.AsyncFunctionDef)):
                        continue
                    for dec in inner.decorator_list:
                        target = dec.func if isinstance(dec, ast.Call) else dec
                        names.add(getattr(target, 'id', None)
                                  or getattr(target, 'attr', None))
                decorated[node.name] = names

        routes = self._admin_routes()
        self.assertGreater(len(routes), 30, 'the sweep found almost no URLs')

        unguarded = [
            '%s -> %s()' % (route, view.__name__)
            for route, _name, view in routes
            if view is not None
            and view.__name__ in decorated
            and not (decorated[view.__name__] & self.GUARDS)
        ]
        self.assertEqual(unguarded, [], 'no sign-in guard: %s' % unguarded)

    def test_a_stranger_gets_no_admin_page(self):
        anonymous = Client()
        open_doors = []
        for route, name, _view in self._admin_routes():
            if '<' in route:
                continue          # needs an id; a 404 would prove nothing
            url = '/' + route
            for response in (anonymous.get(url), anonymous.post(url, {})):
                if response.status_code == 200:
                    open_doors.append('%s (%s)' % (url, name))
                    break
        self.assertEqual(open_doors, [],
                         'reachable without signing in: %s' % open_doors)


class LabelPickerTests(TestCase):
    """The picker's feed, and printing in shelf order."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)

        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=10, map_y=10)
        self.shelf_a = Shelf.objects.create(room=room, name='Shelf A', map_x=1, map_y=1)
        self.shelf_b = Shelf.objects.create(room=room, name='Shelf B', map_x=2, map_y=2)
        self.a1 = ShelfLevel.objects.create(shelf=self.shelf_a, level_number=1,
                                            category='Fiction')
        self.a2 = ShelfLevel.objects.create(shelf=self.shelf_a, level_number=2,
                                            category='Filipiniana')
        self.b1 = ShelfLevel.objects.create(shelf=self.shelf_b, level_number=1,
                                            category='Reference')

        def book(title, level):
            return Book.objects.create(title=title, author='A. Writer',
                                       qr_code=str(uuid4()), status='Available',
                                       shelf_level=level)

        # Deliberately created out of shelf order.
        self.unplaced = book('Zebra Book', None)
        self.b1_book = book('Beta', self.b1)
        self.a2_book = book('Alpha', self.a2)
        self.a1_two = book('Second', self.a1)
        self.a1_one = book('First', self.a1)

    def test_the_feed_reports_where_each_book_lives(self):
        data = self.client.get('/admin-portal/book-label-picker/').json()
        self.assertTrue(data['success'])
        self.assertEqual(len(data['books']), 5)
        self.assertFalse(data['truncated'])
        self.assertEqual(data['max_per_sheet'], 500)

        by_id = {b['book_id']: b for b in data['books']}
        placed = by_id[self.a2_book.book_id]
        self.assertEqual(placed['location'], 'Shelf A · L2')
        self.assertEqual(placed['shelf_name'], 'Shelf A')
        self.assertEqual(placed['level_number'], 2)
        self.assertEqual(placed['level_category'], 'Filipiniana')

        # An unplaced book says so rather than inventing a location.
        self.assertEqual(by_id[self.unplaced.book_id]['location'], '')
        self.assertIsNone(by_id[self.unplaced.book_id]['level_id'])

    def test_the_feed_is_gated(self):
        r = Client().get('/admin-portal/book-label-picker/')
        self.assertNotEqual(r.status_code, 200)

    def test_printing_order_walks_the_shelves(self):
        """Shelf, then level, then title -- with the unplaced ones last."""
        from .labels import sort_for_printing

        ordered = sort_for_printing([
            self.unplaced, self.b1_book, self.a2_book, self.a1_two, self.a1_one])
        self.assertEqual([b.title for b in ordered],
                         ['First', 'Second', 'Alpha', 'Beta', 'Zebra Book'])

    def test_the_sheet_is_returned_in_that_order(self):
        ids = ','.join(str(b.book_id) for b in
                       [self.unplaced, self.b1_book, self.a1_one])
        r = self.client.get('/admin-portal/book-qr-labels/', {'ids': ids})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_sort_by_id_is_still_available(self):
        from .labels import sort_for_printing
        by_shelf = sort_for_printing([self.unplaced, self.a1_one])
        self.assertEqual(by_shelf[0].title, 'First')     # shelf order
        ids = ','.join(str(b.book_id) for b in [self.unplaced, self.a1_one])
        r = self.client.get('/admin-portal/book-qr-labels/',
                            {'ids': ids, 'sort': 'id'})
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_groups_do_not_get_their_own_page(self):
        """Packed tight: five books across three shelf levels is still one page."""
        from .labels import build_label_sheet, sort_for_printing

        books = sort_for_printing(list(Book.objects.select_related(
            'shelf_level', 'shelf_level__shelf').all()))
        _pdf, layout = build_label_sheet(books, qr_mm=25,
                                         show={'title': True, 'location': True})
        self.assertEqual(layout['pages'], 1)
        self.assertGreater(layout['per_page'], 5)

    def test_location_reads_as_shelf_and_level(self):
        from .labels import location_of
        self.assertEqual(location_of(self.a1_one), 'Shelf A · L1')
        self.assertEqual(location_of(self.unplaced), '')


class ShelfShapeTests(TestCase):
    """Shelves that are not rectangles, tops, and tables that hold books."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True)
        self.room = Room.objects.create(floor_plan=self.plan, name='Main',
                                        map_x=100, map_y=100)

    def test_a_rectangular_shelf_still_derives_its_corners(self):
        """Everything drawn before geometry existed must not move."""
        shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                     map_x=100, map_y=100,
                                     width=40, depth=10, rotation=0)
        self.assertEqual(shelf.footprint(),
                         [[80.0, 95.0], [120.0, 95.0], [120.0, 105.0], [80.0, 105.0]])

    def test_rotation_still_turns_the_rectangle(self):
        shelf = Shelf.objects.create(room=self.room, name='Turned',
                                     map_x=0, map_y=0, width=20, depth=10,
                                     rotation=90)
        xs = [p[0] for p in shelf.footprint()]
        ys = [p[1] for p in shelf.footprint()]
        # A quarter turn swaps the run and the depth.
        self.assertAlmostEqual(max(xs) - min(xs), 10.0, places=4)
        self.assertAlmostEqual(max(ys) - min(ys), 20.0, places=4)

    def test_a_traced_shelf_keeps_the_shape_it_was_drawn_as(self):
        """A corner unit is a polygon, not a rectangle with a rotation."""
        chevron = [[0, 0], [60, 0], [60, 12], [12, 12], [12, 60], [0, 60]]
        shelf = Shelf.objects.create(room=self.room, name='Corner',
                                     map_x=20, map_y=20, geometry=chevron)
        self.assertEqual(shelf.footprint(), [[float(x), float(y)] for x, y in chevron])
        self.assertEqual(len(shelf.footprint()), 6)

    def test_an_unplaced_shelf_has_no_footprint(self):
        shelf = Shelf.objects.create(room=self.room, name='Unplaced')
        self.assertEqual(shelf.footprint(), [])

    def test_drawing_a_shelf_saves_the_outline_and_centres_on_it(self):
        chevron = [[0, 0], [60, 0], [60, 12], [12, 12], [12, 60], [0, 60]]
        r = self.client.post('/admin-portal/add-shelf/', {
            'room_id': self.room.room_id, 'name': 'Corner Unit',
            'kind': 'Shelf', 'geometry': json.dumps(chevron),
        })
        data = r.json()
        self.assertTrue(data['success'], data)
        shelf = Shelf.objects.get(shelf_id=data['shelf_id'])
        self.assertEqual(len(shelf.geometry), 6)
        # Positioned by its outline.
        self.assertIsNotNone(shelf.map_x)
        self.assertGreater(shelf.map_x, 0)
        self.assertGreater(shelf.map_y, 0)

    def test_the_top_of_a_case_is_not_a_numbered_level(self):
        shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                     map_x=10, map_y=10)
        ShelfLevel.objects.create(shelf=shelf, level_number=1)
        r = self.client.post('/admin-portal/add-shelf-level/',
                             {'shelf_id': shelf.shelf_id, 'is_top': '1'})
        self.assertTrue(r.json()['success'], r.json())
        top = ShelfLevel.objects.get(shelf=shelf, is_top=True)
        self.assertEqual(top.label, 'Top')
        self.assertEqual(top.short_label, 'Top')
        # It sorts above every real shelf, which is where it physically is.
        self.assertGreater(top.level_number, 1)

    def test_a_case_gets_only_one_top(self):
        shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                     map_x=10, map_y=10)
        self.client.post('/admin-portal/add-shelf-level/',
                         {'shelf_id': shelf.shelf_id, 'is_top': '1'})
        again = self.client.post('/admin-portal/add-shelf-level/',
                                 {'shelf_id': shelf.shelf_id, 'is_top': '1'})
        self.assertFalse(again.json()['success'])
        self.assertEqual(ShelfLevel.objects.filter(shelf=shelf, is_top=True).count(), 1)

    def test_a_label_says_top_rather_than_a_level_number(self):
        """L5 would send someone looking inside for a book sitting on top."""
        from .labels import location_of

        shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                     map_x=10, map_y=10)
        inside = ShelfLevel.objects.create(shelf=shelf, level_number=2)
        on_top = ShelfLevel.objects.create(shelf=shelf, level_number=5, is_top=True)

        a = Book.objects.create(title='Inside', author='X', qr_code=str(uuid4()),
                                shelf_level=inside)
        b = Book.objects.create(title='On top', author='X', qr_code=str(uuid4()),
                                shelf_level=on_top)
        self.assertEqual(location_of(a), 'Shelf A · L2')
        self.assertEqual(location_of(b), 'Shelf A · Top')

    def test_furniture_that_holds_books_is_saved_as_a_shelf(self):
        """A book points at a ShelfLevel, so a table holding one must be a Shelf."""
        top = [[10, 10], [90, 10], [90, 50], [10, 50]]
        r = self.client.post('/admin-portal/add-obstacle/', {
            'floor_plan_id': self.plan.floor_plan_id, 'kind': 'Table',
            'name': 'New Arrivals', 'geometry': json.dumps(top),
            'holds_books': '1',
        })
        data = r.json()
        self.assertTrue(data['success'], data)
        self.assertTrue(data['as_shelf'])
        self.assertEqual(Obstacle.objects.count(), 0)

        shelf = Shelf.objects.get(shelf_id=data['shelf_id'])
        self.assertEqual(shelf.kind, 'Table')
        self.assertEqual(shelf.name, 'New Arrivals')
        self.assertEqual(len(shelf.geometry), 4)
        # One surface, and it is the top of it.
        level = ShelfLevel.objects.get(shelf=shelf)
        self.assertTrue(level.is_top)
        self.assertEqual(shelf.label, 'New Arrivals (table)')

    def test_furniture_that_holds_nothing_stays_an_obstacle(self):
        top = [[10, 10], [90, 10], [90, 50], [10, 50]]
        r = self.client.post('/admin-portal/add-obstacle/', {
            'floor_plan_id': self.plan.floor_plan_id, 'kind': 'Table',
            'name': 'Reading Table', 'geometry': json.dumps(top),
        })
        self.assertTrue(r.json()['success'])
        self.assertEqual(Obstacle.objects.count(), 1)
        self.assertEqual(Shelf.objects.filter(kind='Table').count(), 0)

    def test_a_book_table_needs_a_room_to_live_in(self):
        """It is catalogued inside one, so it cannot be saved without one."""
        Room.objects.all().delete()
        r = self.client.post('/admin-portal/add-obstacle/', {
            'floor_plan_id': self.plan.floor_plan_id, 'kind': 'Table',
            'geometry': json.dumps([[1, 1], [9, 1], [9, 5], [1, 5]]),
            'holds_books': '1',
        })
        data = r.json()
        self.assertFalse(data['success'])
        self.assertIn('room', data['error'].lower())
        self.assertEqual(Shelf.objects.count(), 0)

    def test_the_map_payload_carries_a_resolved_outline(self):
        """So no map has to know the difference between the two shapes."""
        Shelf.objects.create(room=self.room, name='Rect', map_x=50, map_y=50)
        Shelf.objects.create(room=self.room, name='Traced', map_x=50, map_y=50,
                             geometry=[[0, 0], [10, 0], [10, 10]])
        data = self.client.get('/admin-portal/map-data/').json()
        shelves = {s['name']: s for s in data.get('shelves', [])}
        self.assertGreaterEqual(len(shelves['Rect']['footprint']), 4)
        self.assertEqual(len(shelves['Traced']['footprint']), 3)
        self.assertIsNone(shelves['Rect']['geometry'])


class ShelfPositionTests(TestCase):
    """Underneath, columns, and units that are off the floor."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True)
        self.room = Room.objects.create(floor_plan=self.plan, name='Main',
                                        map_x=100, map_y=100)
        self.shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                          map_x=10, map_y=10)

    def _level(self, **params):
        params.setdefault('shelf_id', self.shelf.shelf_id)
        return self.client.post('/admin-portal/add-shelf-level/', params).json()

    def test_underneath_is_not_a_numbered_shelf(self):
        r = self._level(is_under='1')
        self.assertTrue(r['success'], r)
        under = ShelfLevel.objects.get(shelf=self.shelf, is_under=True)
        self.assertEqual(under.label, 'Underneath')
        self.assertEqual(under.short_label, 'Under')
        # Below everything, so it sorts first.
        self.assertEqual(under.level_number, 0)

    def test_a_level_cannot_be_both_top_and_underneath(self):
        r = self._level(is_top='1', is_under='1')
        self.assertFalse(r['success'])
        self.assertEqual(ShelfLevel.objects.count(), 0)

    def test_a_column_addresses_a_bay(self):
        r = self._level(level_number='2', column_number='3')
        self.assertTrue(r['success'], r)
        lv = ShelfLevel.objects.get(shelf_level_id=r['shelf_level_id'])
        self.assertEqual(lv.label, 'Level 2, Column 3')
        self.assertEqual(lv.short_label, 'L2C3')

    def test_a_level_without_a_column_is_unchanged(self):
        r = self._level(level_number='2')
        lv = ShelfLevel.objects.get(shelf_level_id=r['shelf_level_id'])
        self.assertIsNone(lv.column_number)
        self.assertEqual(lv.label, 'Level 2')
        self.assertEqual(lv.short_label, 'L2')

    def test_a_bad_column_is_refused_rather_than_guessed(self):
        self.assertFalse(self._level(level_number='1', column_number='abc')['success'])

    def test_labels_walk_a_case_bottom_to_top_and_left_to_right(self):
        from .labels import sort_for_printing

        under = ShelfLevel.objects.create(shelf=self.shelf, level_number=0, is_under=True)
        l1c2 = ShelfLevel.objects.create(shelf=self.shelf, level_number=1, column_number=2)
        l1c1 = ShelfLevel.objects.create(shelf=self.shelf, level_number=1, column_number=1)
        top = ShelfLevel.objects.create(shelf=self.shelf, level_number=9, is_top=True)

        def book(name, level):
            return Book.objects.create(title=name, author='X',
                                       qr_code=str(uuid4()), shelf_level=level)
        made = [book('D top', top), book('C l1c2', l1c2),
                book('A under', under), book('B l1c1', l1c1)]
        self.assertEqual([b.title for b in sort_for_printing(made)],
                         ['A under', 'B l1c1', 'C l1c2', 'D top'])

    def test_an_elevated_unit_says_so_and_is_flagged(self):
        wall = Shelf.objects.create(room=self.room, name='Corner Ledge',
                                    map_x=5, map_y=5, mount='Ceiling',
                                    mount_height_m=2.4)
        floor = Shelf.objects.create(room=self.room, name='Stack 1',
                                     map_x=6, map_y=6)
        self.assertTrue(wall.is_elevated)
        self.assertFalse(floor.is_elevated)
        self.assertEqual(wall.label, 'Corner Ledge (ceiling-mounted)')
        self.assertEqual(floor.label, 'Stack 1')

    def test_drawing_an_elevated_shelf_keeps_its_mount(self):
        r = self.client.post('/admin-portal/add-shelf/', {
            'room_id': self.room.room_id, 'name': 'Ceiling Corner',
            'mount': 'Ceiling', 'mount_height_m': '2.4',
            'geometry': json.dumps([[0, 0], [40, 0], [40, 20]]),
        }).json()
        self.assertTrue(r['success'], r)
        shelf = Shelf.objects.get(shelf_id=r['shelf_id'])
        self.assertEqual(shelf.mount, 'Ceiling')
        self.assertEqual(shelf.mount_height_m, 2.4)

    def test_an_impossible_mounting_height_is_refused(self):
        r = self.client.post('/admin-portal/add-shelf/', {
            'room_id': self.room.room_id, 'name': 'Too High',
            'mount': 'Wall', 'mount_height_m': '95',
        }).json()
        self.assertFalse(r['success'])
        self.assertIn('metres', r['error'])

    def test_the_map_payload_flags_an_elevated_unit(self):
        """So a map can draw it hollow instead of as floor furniture."""
        Shelf.objects.create(room=self.room, name='Hung', map_x=5, map_y=5,
                             mount='Ceiling')
        data = self.client.get('/admin-portal/map-data/').json()
        by_name = {s['name']: s for s in data.get('shelves', [])}
        self.assertTrue(by_name['Hung']['elevated'])
        self.assertFalse(by_name['Shelf A']['elevated'])


class AuditSweepTests(TestCase):
    """Scanning a room without saying which shelf you are standing at."""

    def setUp(self):
        self.user = _admin(modules='inventory')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='G', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=1, map_y=1)
        self.a = Shelf.objects.create(room=room, name='Shelf A', map_x=1, map_y=1)
        self.b = Shelf.objects.create(room=room, name='Shelf B', map_x=2, map_y=2)
        self.la = ShelfLevel.objects.create(shelf=self.a, level_number=1)
        self.lb = ShelfLevel.objects.create(shelf=self.b, level_number=1)

        self.labels = {}
        for shelf_key, level, n in (('a', self.la, 3), ('b', self.lb, 2)):
            for i in range(n):
                book = Book.objects.create(title=f'{shelf_key}{i}', author='X',
                                           qr_code=str(uuid4()), shelf_level=level)
                rec = InventoryRecord.objects.create(book=book, status='In Stock',
                                                     condition='Good',
                                                     qr_label=f'{shelf_key.upper()}-{i}')
                self.labels.setdefault(shelf_key, []).append(rec.qr_label)

    def _sweep(self, scanned):
        return self.client.post('/admin-portal/inventory/audit/progress/',
                                {'scanned': scanned}).json()

    def test_scans_sort_themselves_into_shelves(self):
        """No shelf is chosen up front; each label says where it belongs."""
        data = self._sweep([self.labels['a'][0], self.labels['b'][0],
                            self.labels['a'][1]])
        self.assertTrue(data['success'])
        groups = {g['shelf_name']: g for g in data['groups']}
        self.assertEqual(groups['Shelf A']['found_count'], 2)
        self.assertEqual(groups['Shelf A']['expected_count'], 3)
        self.assertEqual(groups['Shelf B']['found_count'], 1)
        self.assertEqual(groups['Shelf B']['expected_count'], 2)

    def test_progress_says_when_a_shelf_is_done(self):
        data = self._sweep(self.labels['a'])
        shelf_a = data['groups'][0]
        self.assertEqual(shelf_a['found_count'], 3)
        self.assertEqual(shelf_a['percent'], 100)
        self.assertTrue(shelf_a['complete'])

    def test_a_half_counted_shelf_is_not_complete(self):
        data = self._sweep(self.labels['a'][:1])
        self.assertFalse(data['groups'][0]['complete'])
        self.assertEqual(data['groups'][0]['percent'], 33)

    def test_duplicate_scans_count_once(self):
        """Waving the scanner twice at one book must not invent a copy."""
        one = self.labels['a'][0]
        data = self._sweep([one, one, one])
        self.assertEqual(data['total_scanned'], 1)
        self.assertEqual(data['groups'][0]['found_count'], 1)

    def test_an_unknown_label_is_reported_not_dropped(self):
        data = self._sweep([self.labels['a'][0], 'NOT-A-REAL-LABEL'])
        self.assertEqual(data['unknown'], ['NOT-A-REAL-LABEL'])
        self.assertEqual(len(data['groups']), 1)

    def test_a_copy_with_no_shelf_is_kept_separate(self):
        loose = Book.objects.create(title='Loose', author='X', qr_code=str(uuid4()))
        InventoryRecord.objects.create(book=loose, status='In Stock',
                                       condition='Good', qr_label='LOOSE-1')
        data = self._sweep(['LOOSE-1'])
        self.assertEqual(len(data['unshelved']), 1)
        self.assertEqual(data['groups'], [])

    def test_it_is_gated(self):
        r = Client().post('/admin-portal/inventory/audit/progress/', {'scanned': ['x']})
        self.assertNotEqual(r.status_code, 200)


class BulkReshelvingTests(TestCase):
    """A trolley of returns goes back in one action, not thirty."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        self.books = [Book.objects.create(title=f'B{i}', author='X',
                                          qr_code=str(uuid4()),
                                          status='For Reshelving')
                      for i in range(4)]

    def _post(self, ids):
        return self.client.post('/admin-portal/reshelving/',
                                {'book_id': [str(i) for i in ids]}).json()

    def test_many_at_once(self):
        ids = [b.book_id for b in self.books[:3]]
        data = self._post(ids)
        self.assertTrue(data['success'])
        self.assertEqual(data['shelved_count'], 3)
        self.assertEqual(data['remaining'], 1)
        self.assertEqual(Book.objects.filter(status='Available').count(), 3)

    def test_one_still_works_unchanged(self):
        data = self._post([self.books[0].book_id])
        self.assertTrue(data['success'])
        self.assertEqual(data['shelved_count'], 1)
        self.assertEqual(data['title'], 'B0')

    def test_a_book_someone_else_already_shelved_is_skipped_not_fatal(self):
        """Two people working the trolley must not error each other out."""
        self.books[0].status = 'Available'
        self.books[0].save(update_fields=['status'])
        data = self._post([self.books[0].book_id, self.books[1].book_id])
        self.assertTrue(data['success'])
        self.assertEqual(data['shelved_count'], 1)
        self.assertEqual(data['skipped_count'], 1)

    def test_nothing_selected_is_refused(self):
        self.assertFalse(self.client.post('/admin-portal/reshelving/', {}).json()['success'])

    def test_all_already_shelved_reports_it(self):
        Book.objects.update(status='Available')
        data = self._post([b.book_id for b in self.books])
        self.assertFalse(data['success'])
        self.assertIn('already', data['error'].lower())


class DoorPlacementTests(TestCase):
    """Where a door lands, and how wide it is once it is there."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True)
        # A wide, short room.
        self.room = Room.objects.create(
            floor_plan=self.plan, name='Hallway', map_x=0, map_y=0,
            geometry=[[100, 100], [500, 100], [500, 200], [100, 200]])

    def _add(self, x, y, **extra):
        params = {'room_id': self.room.room_id, 'map_x': x, 'map_y': y}
        params.update(extra)
        return self.client.post('/admin-portal/add-door/', params).json()

    def _edit(self, door_id, **params):
        params['door_id'] = door_id
        return self.client.post('/admin-portal/edit-door/', params).json()

    def test_the_door_lands_on_the_wall_nearest_the_click(self):
        """Each click belongs to one wall, and it is the one clicked."""
        for label, (cx, cy), expect in (
                ('top', (300, 110), (300, 100)),
                ('bottom', (300, 190), (300, 200)),
                ('left', (110, 150), (100, 150)),
                ('right', (490, 150), (500, 150))):
            r = self._add(cx, cy)
            self.assertTrue(r['success'], (label, r))
            self.assertAlmostEqual(r['door']['x'], expect[0], places=3, msg=label)
            self.assertAlmostEqual(r['door']['y'], expect[1], places=3, msg=label)

    def test_a_new_door_takes_the_default_width(self):
        r = self._add(300, 110)
        self.assertEqual(r['door']['width'], 28)

    def test_width_can_be_changed_after_placement(self):
        """The gap this closes: before edit_door existed, 28 was for ever."""
        door_id = self._add(300, 110)['door']['door_id']
        r = self._edit(door_id, width=90)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['door']['width'], 90)
        self.assertEqual(Door.objects.get(pk=door_id).width, 90)

    def test_an_impossible_width_is_refused(self):
        door_id = self._add(300, 110)['door']['door_id']
        for bad in (0, 5, 401, 'wide'):
            r = self._edit(door_id, width=bad)
            self.assertFalse(r['success'], bad)
        self.assertEqual(Door.objects.get(pk=door_id).width, 28)

    def test_resizing_from_one_end_leaves_the_other_end_alone(self):
        """Dragging an end grip sends a new width AND a new centre together."""
        door_id = self._add(300, 110)['door']['door_id']
        door = Door.objects.get(pk=door_id)
        self.assertAlmostEqual(door.rotation, 0.0, places=3)   # along the top wall
        left_end = door.map_x - door.width / 2                 # the end held still

        new_width = 100
        r = self._edit(door_id, width=new_width,
                       map_x=left_end + new_width / 2, map_y=door.map_y)
        self.assertTrue(r['success'], r)

        door.refresh_from_db()
        self.assertEqual(door.width, new_width)
        self.assertAlmostEqual(door.map_x - door.width / 2, left_end, places=3)
        # And it is still on the wall it was cut into.
        self.assertAlmostEqual(door.map_y, 100.0, places=3)

    def test_a_moved_door_is_re_snapped_to_the_wall(self):
        """A centre sent from a drag is never trusted as final."""
        door_id = self._add(300, 110)['door']['door_id']
        r = self._edit(door_id, width=40, map_x=320, map_y=137)
        self.assertTrue(r['success'], r)
        door = Door.objects.get(pk=door_id)
        self.assertAlmostEqual(door.map_y, 100.0, places=3)    # pulled back onto the wall
        self.assertAlmostEqual(door.map_x, 320.0, places=3)

    def test_editing_nothing_is_refused_rather_than_silently_accepted(self):
        door_id = self._add(300, 110)['door']['door_id']
        r = self._edit(door_id)
        self.assertFalse(r['success'])


class WallCrossingTests(TestCase):
    """A connection asserts someone can walk straight between two points."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, pixels_per_meter=100)
        # Two rooms sharing the wall x = 300.
        self.left = Room.objects.create(
            floor_plan=self.plan, name='Left', map_x=200, map_y=200,
            geometry=[[100, 100], [300, 100], [300, 300], [100, 300]])
        self.right = Room.objects.create(
            floor_plan=self.plan, name='Right', map_x=400, map_y=200,
            geometry=[[300, 100], [500, 100], [500, 300], [300, 300]])

    def _wp(self, x, y):
        return Waypoint.objects.create(floor_plan=self.plan, map_x=x, map_y=y)

    def _connect(self, a, b):
        return self.client.post('/admin-portal/add-waypoint-connection/', {
            'waypoint_from_id': a.waypoint_id, 'waypoint_to_id': b.waypoint_id}).json()

    def test_a_connection_through_a_solid_wall_is_refused(self):
        r = self._connect(self._wp(200, 200), self._wp(400, 200))
        self.assertFalse(r['success'])
        self.assertIn('wall', r['error'].lower())
        self.assertEqual(WaypointConnection.objects.count(), 0)

    def test_the_refusal_names_the_room_whose_wall_it_is(self):
        r = self._connect(self._wp(200, 200), self._wp(400, 200))
        self.assertTrue('Left' in r['error'] or 'Right' in r['error'], r['error'])

    def test_the_same_connection_is_allowed_once_a_door_is_there(self):
        """The whole point: a door is what makes a crossing legal."""
        Door.objects.create(room=self.left, map_x=300, map_y=200,
                            width=60, rotation=90)
        r = self._connect(self._wp(200, 200), self._wp(400, 200))
        self.assertTrue(r['success'], r)
        self.assertEqual(WaypointConnection.objects.count(), 1)

    def test_a_door_elsewhere_on_the_wall_does_not_help(self):
        """A doorway 200 units up the wall is not the one being walked through."""
        Door.objects.create(room=self.left, map_x=300, map_y=110,
                            width=20, rotation=90)
        r = self._connect(self._wp(200, 250), self._wp(400, 250))
        self.assertFalse(r['success'], r)

    def test_two_waypoints_in_one_room_connect_freely(self):
        r = self._connect(self._wp(150, 150), self._wp(250, 250))
        self.assertTrue(r['success'], r)

    def test_a_connection_running_along_a_wall_is_not_a_crossing(self):
        """Sitting on the wall, or tracking it, is not passing through it."""
        r = self._connect(self._wp(300, 120), self._wp(300, 280))
        self.assertTrue(r['success'], r)

    def test_floors_are_joined_by_stairs_not_by_connections(self):
        upstairs = FloorPlan.objects.create(name='First', floor_number=2)
        far = Waypoint.objects.create(floor_plan=upstairs, map_x=200, map_y=200)
        r = self._connect(self._wp(150, 150), far)
        self.assertFalse(r['success'])
        self.assertIn('floor', r['error'].lower())


class DoorResnapTests(TestCase):
    """Reshaping a room used to leave its doors floating in mid-air."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, pixels_per_meter=100)
        self.room = Room.objects.create(
            floor_plan=self.plan, name='Hall', map_x=200, map_y=200,
            geometry=[[100, 100], [300, 100], [300, 300], [100, 300]])
        self.door = Door.objects.create(room=self.room, map_x=200, map_y=100,
                                        width=40, rotation=0)

    def _reshape(self, geometry):
        import json
        return self.client.post('/admin-portal/edit-room/', {
            'room_id': self.room.room_id,
            'geometry': json.dumps(geometry)}).json()

    def test_a_door_follows_the_wall_it_was_cut_into(self):
        # Push the top wall down from y=100 to y=160.
        r = self._reshape([[100, 160], [300, 160], [300, 300], [100, 300]])
        self.assertTrue(r['success'], r)
        self.door.refresh_from_db()
        self.assertAlmostEqual(self.door.map_y, 160.0, places=3)
        self.assertAlmostEqual(self.door.map_x, 200.0, places=3)

    def test_the_move_is_reported_rather_than_silent(self):
        """Nearest-edge is a guess, so it has to be visible to be correctable."""
        r = self._reshape([[100, 160], [300, 160], [300, 300], [100, 300]])
        self.assertEqual(len(r['doors_moved']), 1)
        self.assertEqual(r['doors_moved'][0]['door_id'], self.door.door_id)
        self.assertAlmostEqual(r['doors_moved'][0]['moved_by'], 60.0, places=1)

    def test_a_door_already_on_the_wall_is_left_alone(self):
        r = self._reshape([[100, 100], [300, 100], [300, 400], [100, 400]])
        self.assertEqual(r['doors_moved'], [])
        self.door.refresh_from_db()
        self.assertAlmostEqual(self.door.map_y, 100.0, places=3)

    def test_renaming_a_room_does_not_disturb_its_doors(self):
        r = self.client.post('/admin-portal/edit-room/', {
            'room_id': self.room.room_id, 'name': 'Lobby'}).json()
        self.assertTrue(r['success'], r)
        self.assertEqual(r.get('doors_moved'), [])


class FloorPlanReadinessTests(TestCase):
    """Everything this reports is something that otherwise fails silently."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1)

    def _check(self):
        return self.client.get('/admin-portal/floor-plan-readiness/',
                               {'floor_plan_id': self.plan.floor_plan_id}).json()

    def _texts(self, data):
        return [i['text'] for i in data['issues']]

    def test_a_bare_plan_reports_its_blockers(self):
        data = self._check()
        self.assertFalse(data['ready'])
        joined = ' | '.join(self._texts(data))
        self.assertIn('No scale set', joined)
        self.assertIn('No rooms drawn', joined)
        self.assertIn('No waypoints placed', joined)

    def test_walls_that_almost_touch_are_flagged(self):
        """A tenth of a unit apart looks joined at every zoom level."""
        self.plan.pixels_per_meter = 100
        self.plan.save()
        Room.objects.create(floor_plan=self.plan, name='A', map_x=0, map_y=0,
                            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        Room.objects.create(floor_plan=self.plan, name='B', map_x=0, map_y=0,
                            geometry=[[100.5, 0], [200, 0], [200, 100], [100.5, 100]])
        joined = ' | '.join(self._texts(self._check()))
        self.assertIn('almost touch', joined)

    def test_rooms_that_genuinely_touch_are_not_flagged(self):
        Room.objects.create(floor_plan=self.plan, name='A', map_x=0, map_y=0,
                            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        Room.objects.create(floor_plan=self.plan, name='B', map_x=0, map_y=0,
                            geometry=[[100, 0], [200, 0], [200, 100], [100, 100]])
        joined = ' | '.join(self._texts(self._check()))
        self.assertNotIn('almost touch', joined)

    def test_an_unreachable_room_is_named(self):
        Room.objects.create(floor_plan=self.plan, name='Store', map_x=0, map_y=0,
                            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        Room.objects.create(floor_plan=self.plan, name='Lobby', map_x=0, map_y=0,
                            geometry=[[200, 0], [300, 0], [300, 100], [200, 100]])
        Waypoint.objects.create(floor_plan=self.plan, map_x=250, map_y=50)
        joined = ' | '.join(self._texts(self._check()))
        self.assertIn('No waypoint inside "Store"', joined)
        self.assertNotIn('No waypoint inside "Lobby"', joined)

    def test_a_split_route_network_is_reported(self):
        w = [Waypoint.objects.create(floor_plan=self.plan, map_x=x, map_y=0)
             for x in (0, 10, 100, 110)]
        WaypointConnection.objects.create(waypoint_from=w[0], waypoint_to=w[1], distance=10)
        WaypointConnection.objects.create(waypoint_from=w[2], waypoint_to=w[3], distance=10)
        joined = ' | '.join(self._texts(self._check()))
        self.assertIn('2 separate pieces', joined)

    def test_orphan_waypoints_are_counted(self):
        Waypoint.objects.create(floor_plan=self.plan, map_x=0, map_y=0)
        joined = ' | '.join(self._texts(self._check()))
        self.assertIn('1 waypoint(s) connected to nothing', joined)

    def test_a_finished_plan_reports_ready(self):
        self.plan.pixels_per_meter = 100
        self.plan.is_active = True
        self.plan.save()
        room = Room.objects.create(
            floor_plan=self.plan, name='Main', map_x=50, map_y=50,
            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        Door.objects.create(room=room, map_x=50, map_y=0, width=30, rotation=0)
        for i in range(3):
            BLEBeacon.objects.create(floor_plan=self.plan, beacon_uuid='u%d' % i,
                                     map_x=i * 10, map_y=i * 7, tx_power=-59, height=2.4)
        a = Waypoint.objects.create(floor_plan=self.plan, map_x=30, map_y=30)
        b = Waypoint.objects.create(floor_plan=self.plan, map_x=70, map_y=70)
        WaypointConnection.objects.create(waypoint_from=a, waypoint_to=b, distance=56)
        data = self._check()
        self.assertTrue(data['ready'], data['issues'])
        self.assertEqual(data['blockers'], 0)


class SharedDoorTests(TestCase):
    """A doorway joins two rooms, and that is what makes a plan connected."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, pixels_per_meter=100)
        # Left and Right share the wall x = 300. Far is nowhere near either.
        self.left = Room.objects.create(
            floor_plan=self.plan, name='Left', map_x=200, map_y=200,
            geometry=[[100, 100], [300, 100], [300, 300], [100, 300]])
        self.right = Room.objects.create(
            floor_plan=self.plan, name='Right', map_x=400, map_y=200,
            geometry=[[300, 100], [500, 100], [500, 300], [300, 300]])
        self.far = Room.objects.create(
            floor_plan=self.plan, name='Far', map_x=900, map_y=200,
            geometry=[[800, 100], [1000, 100], [1000, 300], [800, 300]])

    def _add(self, room, x, y):
        return self.client.post('/admin-portal/add-door/', {
            'room_id': room.room_id, 'map_x': x, 'map_y': y}).json()

    def _edit(self, door_id, **params):
        params['door_id'] = door_id
        return self.client.post('/admin-portal/edit-door/', params).json()

    def test_placing_a_door_on_a_shared_wall_suggests_the_other_room(self):
        r = self._add(self.left, 295, 200)
        self.assertTrue(r['success'], r)
        self.assertIsNotNone(r['suggested_room_b'])
        self.assertEqual(r['suggested_room_b']['name'], 'Right')

    def test_the_suggestion_is_only_a_suggestion(self):
        """Guessing outright would record adjacency nobody agreed to."""
        r = self._add(self.left, 295, 200)
        self.assertIsNone(Door.objects.get(pk=r['door']['door_id']).room_b)

    def test_a_door_with_nothing_on_the_far_side_suggests_nothing(self):
        r = self._add(self.left, 105, 200)     # the outer wall, x = 100
        self.assertTrue(r['success'], r)
        self.assertIsNone(r['suggested_room_b'])

    def test_rooms_that_only_nearly_touch_are_still_found(self):
        """Hand-drawn walls miss by a fraction; that must not hide adjacency."""
        gap = Room.objects.create(
            floor_plan=self.plan, name='Gap', map_x=200, map_y=500,
            geometry=[[100, 300.4], [300, 300.4], [300, 500], [100, 500]])
        r = self._add(self.left, 200, 295)     # the wall at y = 300
        self.assertEqual(r['suggested_room_b']['name'], gap.name)

    def test_a_wall_far_away_is_not_offered(self):
        r = self._add(self.far, 805, 200)
        self.assertIsNone(r['suggested_room_b'])

    def test_the_far_room_can_be_set_and_cleared(self):
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        r = self._edit(door_id, room_b=self.right.room_id)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['door']['room_b'], self.right.room_id)
        self.assertEqual(r['door']['room_b_name'], 'Right')

        r = self._edit(door_id, room_b='')
        self.assertTrue(r['success'], r)
        self.assertIsNone(r['door']['room_b'])

    def test_a_door_cannot_join_a_room_to_itself(self):
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        r = self._edit(door_id, room_b=self.left.room_id)
        self.assertFalse(r['success'])
        self.assertIn('itself', r['error'])

    def test_the_two_rooms_must_be_on_the_same_floor(self):
        upstairs = FloorPlan.objects.create(name='First', floor_number=2)
        elsewhere = Room.objects.create(floor_plan=upstairs, name='Up',
                                        map_x=0, map_y=0,
                                        geometry=[[0, 0], [10, 0], [10, 10], [0, 10]])
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        r = self._edit(door_id, room_b=elsewhere.room_id)
        self.assertFalse(r['success'])
        self.assertIn('same floor', r['error'])

    def test_deleting_the_far_room_leaves_the_door_standing(self):
        """SET_NULL, not CASCADE: the doorway is still a hole in a wall."""
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        self._edit(door_id, room_b=self.right.room_id)
        self.right.delete()
        door = Door.objects.filter(pk=door_id).first()
        self.assertIsNotNone(door, 'the door was deleted with the far room')
        self.assertIsNone(door.room_b)

    def test_adjacency_comes_from_the_doors(self):
        from library.views import _room_adjacency
        self.assertEqual(_room_adjacency(self.plan), set())
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        self._edit(door_id, room_b=self.right.room_id)
        self.assertEqual(
            _room_adjacency(self.plan),
            {(min(self.left.room_id, self.right.room_id),
              max(self.left.room_id, self.right.room_id))})

    def test_a_linked_door_excuses_a_crossing_between_its_own_two_rooms(self):
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        self._edit(door_id, room_b=self.right.room_id, width=60)
        a = Waypoint.objects.create(floor_plan=self.plan, map_x=200, map_y=200)
        b = Waypoint.objects.create(floor_plan=self.plan, map_x=400, map_y=200)
        r = self.client.post('/admin-portal/add-waypoint-connection/', {
            'waypoint_from_id': a.waypoint_id, 'waypoint_to_id': b.waypoint_id}).json()
        self.assertTrue(r['success'], r)

    def test_a_door_linked_to_other_rooms_excuses_nothing_here(self):
        """The precision the second room buys: a doorway belongs to one wall."""
        # A door that links unrelated rooms.
        stray = Door.objects.create(room=self.far, room_b=self.left,
                                    map_x=300, map_y=200, width=80, rotation=90)
        self.assertEqual(stray.room_b_id, self.left.room_id)
        a = Waypoint.objects.create(floor_plan=self.plan, map_x=200, map_y=200)
        b = Waypoint.objects.create(floor_plan=self.plan, map_x=400, map_y=200)
        r = self.client.post('/admin-portal/add-waypoint-connection/', {
            'waypoint_from_id': a.waypoint_id, 'waypoint_to_id': b.waypoint_id}).json()
        # It crosses Right's wall, and this door says nothing about Right.
        self.assertFalse(r['success'], r)

    def test_reshaping_away_from_the_far_room_drops_the_link(self):
        """The link was true of the old wall; it has to be re-earned."""
        import json
        door_id = self._add(self.left, 295, 200)['door']['door_id']
        self._edit(door_id, room_b=self.right.room_id)
        # Pull Left's right-hand wall back to x = 200, well away from Right.
        r = self.client.post('/admin-portal/edit-room/', {
            'room_id': self.left.room_id,
            'geometry': json.dumps([[100, 100], [200, 100], [200, 300], [100, 300]])}).json()
        self.assertTrue(r['success'], r)
        door = Door.objects.get(pk=door_id)
        self.assertIsNone(door.room_b)
        self.assertEqual(r['doors_moved'][0]['unlinked_from'], 'Right')

    def test_readiness_names_rooms_no_door_reaches(self):
        # Doors exist but connect nothing.
        self._add(self.left, 105, 200)
        data = self.client.get('/admin-portal/floor-plan-readiness/',
                               {'floor_plan_id': self.plan.floor_plan_id}).json()
        joined = ' | '.join(i['text'] for i in data['issues'])
        self.assertIn('not joined to any other by a door', joined)

    def test_readiness_points_out_doors_that_look_shared(self):
        self._add(self.left, 295, 200)
        data = self.client.get('/admin-portal/floor-plan-readiness/',
                               {'floor_plan_id': self.plan.floor_plan_id}).json()
        joined = ' | '.join(i['text'] for i in data['issues'])
        self.assertIn('look shared but are not linked', joined)


class GeneratedRouteTests(TestCase):
    """Deriving the walkable network instead of clicking it out by hand."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, pixels_per_meter=100)
        # Two rooms sharing the wall x = 300, joined by a door in the middle.
        self.left = Room.objects.create(
            floor_plan=self.plan, name='Left', map_x=200, map_y=200,
            geometry=[[100, 100], [300, 100], [300, 300], [100, 300]])
        self.right = Room.objects.create(
            floor_plan=self.plan, name='Right', map_x=400, map_y=200,
            geometry=[[300, 100], [500, 100], [500, 300], [300, 300]])
        self.door = Door.objects.create(room=self.left, room_b=self.right,
                                        map_x=300, map_y=200, width=70, rotation=90)

    def _generate(self):
        return self.client.post('/admin-portal/generate-waypoints/',
                                {'floor_plan_id': self.plan.floor_plan_id}).json()

    def _links(self):
        wps = Waypoint.objects.filter(floor_plan=self.plan)
        return WaypointConnection.objects.filter(waypoint_from__in=wps,
                                                 waypoint_to__in=wps)

    def test_it_places_and_joins_something(self):
        r = self._generate()
        self.assertTrue(r['success'], r)
        self.assertGreater(r['placed'], 0)
        self.assertGreater(r['connections'], 0)
        self.assertEqual(r['unreachable'], [])

    def test_nothing_it_draws_goes_through_a_wall(self):
        """The property that matters, checked independently of the generator."""
        from library.views import _walls_crossed
        self._generate()
        for c in self._links().select_related('waypoint_from', 'waypoint_to'):
            a, b = c.waypoint_from, c.waypoint_to
            self.assertEqual(
                _walls_crossed(self.plan, a.map_x, a.map_y, b.map_x, b.map_y), [],
                'a generated connection crosses a wall')

    def test_the_two_rooms_end_up_joined(self):
        self._generate()
        wps = {w.waypoint_id: w for w in Waypoint.objects.filter(floor_plan=self.plan)}
        parent = dict((i, i) for i in wps)

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]
                i = parent[i]
            return i

        for c in self._links():
            a, b = find(c.waypoint_from_id), find(c.waypoint_to_id)
            if a != b:
                parent[a] = b
        self.assertEqual(len(set(find(i) for i in wps)), 1,
                         'the generated network is not one connected piece')

    def test_a_second_run_replaces_only_what_it_made(self):
        """The whole reason generated waypoints are marked as such."""
        mine = Waypoint.objects.create(floor_plan=self.plan, map_x=150, map_y=150,
                                       label='mine by hand')
        first = self._generate()
        second = self._generate()
        self.assertEqual(second['replaced'], first['placed'])
        self.assertEqual(second['kept_manual'], 1)
        self.assertTrue(Waypoint.objects.filter(pk=mine.waypoint_id).exists())

    def test_a_room_with_no_door_is_reported_not_silently_skipped(self):
        Room.objects.create(
            floor_plan=self.plan, name='Cupboard', map_x=900, map_y=200,
            geometry=[[800, 100], [1000, 100], [1000, 300], [800, 300]])
        r = self._generate()
        # It still gets its own middle waypoint, but nothing connects to it.
        self.assertTrue(r['success'], r)
        cupboard = [w for w in Waypoint.objects.filter(floor_plan=self.plan)
                    if w.label == 'Cupboard']
        self.assertEqual(len(cupboard), 1)
        self.assertFalse(
            WaypointConnection.objects.filter(waypoint_from=cupboard[0]).exists()
            or WaypointConnection.objects.filter(waypoint_to=cupboard[0]).exists())

    def test_it_refuses_when_there_is_nothing_drawn(self):
        empty = FloorPlan.objects.create(name='Blank', floor_number=9)
        r = self.client.post('/admin-portal/generate-waypoints/',
                             {'floor_plan_id': empty.floor_plan_id}).json()
        self.assertFalse(r['success'])
        self.assertIn('room', r['error'].lower())

    def test_furniture_is_walked_around_not_through(self):
        from library.views import _crosses_obstacle
        Obstacle.objects.create(
            floor_plan=self.plan, kind='Table', name='Long table',
            map_x=200, map_y=200,
            geometry=[[120, 190], [280, 190], [280, 210], [120, 210]])
        self._generate()
        for c in self._links().select_related('waypoint_from', 'waypoint_to'):
            a, b = c.waypoint_from, c.waypoint_to
            if a.label == 'Doorway' and b.label == 'Doorway':
                continue          # the step through the door is exempt
            self.assertFalse(
                _crosses_obstacle(self.plan, a.map_x, a.map_y, b.map_x, b.map_y),
                'a generated connection walks through the table')


class BookLocationTests(TestCase):
    """Where a copy is, to the space on the board it sits on."""

    def setUp(self):
        self.user = _admin(modules='books,shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, pixels_per_meter=100)
        self.room = Room.objects.create(floor_plan=self.plan, name='Main',
                                        map_x=200, map_y=200,
                                        geometry=[[0, 0], [400, 0], [400, 400], [0, 400]])
        self.shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                          map_x=200, map_y=200, width=200, depth=40,
                                          rotation=0)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=3,
                                               category='Fiction')

    def _book(self, **kw):
        kw.setdefault('title', 'A Book')
        kw.setdefault('author', 'Someone')
        kw.setdefault('shelf_level', self.level)
        return Book.objects.create(**kw)

    def test_a_blank_slot_is_not_position_one(self):
        """Not recorded and 'first on the shelf' must stay different."""
        from library.views import parse_shelf_slot
        for blank in ('', '   ', None, 'four', '0', '-2', '10000'):
            self.assertIsNone(parse_shelf_slot(blank), repr(blank))
        self.assertEqual(parse_shelf_slot('1'), 1)
        self.assertEqual(parse_shelf_slot(' 12 '), 12)

    def test_the_location_reads_in_walking_order(self):
        from library.views import book_location_words
        for i in range(1, 4):
            self._book(title='Filler %d' % i, shelf_slot=i, status='Available')
        book = self._book(shelf_slot=4, status='Available')
        words = book_location_words(book)
        self.assertIn('Shelf A', words)
        self.assertIn('Level 3', words)
        # Counted among the copies on the shelf, not read off the stored slot.
        self.assertIn('4th book along', words)
        self.assertLess(words.index('Shelf A'), words.index('Level 3'))

    def test_precision_nobody_entered_is_not_claimed(self):
        from library.views import book_location_words
        book = self._book()
        self.assertNotIn('from the left', book_location_words(book))
        self.assertEqual(book_location_words(self._book(shelf_level=None)), '')

    def test_ordinals_survive_the_teens(self):
        from library.views import _ordinal
        self.assertEqual([_ordinal(n) for n in (1, 2, 3, 11, 12, 13, 21, 22)],
                         ['1st', '2nd', '3rd', '11th', '12th', '13th', '21st', '22nd'])

    def test_the_slot_can_be_set_and_cleared_by_editing(self):
        book = self._book()
        base = {'title': book.title, 'author': book.author,
                'shelf_level': self.level.shelf_level_id}
        r = self.client.post('/admin-portal/edit-book/%d/' % book.book_id,
                             dict(base, shelf_slot='7'),
                             HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        book.refresh_from_db()
        self.assertEqual(book.shelf_slot, 7)

        self.client.post('/admin-portal/edit-book/%d/' % book.book_id,
                         dict(base, shelf_slot=''),
                         HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        book.refresh_from_db()
        self.assertIsNone(book.shelf_slot)

    # Map drawing checks.

    def test_the_route_is_arrows_with_no_line_through_them(self):
        """The line said a route exists; the arrows say which way to walk."""
        html = self.client.get('/patron/map/').content.decode()
        self.assertNotIn('routeLine', html)
        self.assertIn('drawRouteArrows', html)

    def test_the_marker_is_held_still_rather_than_crept_along(self):
        """Creeping toward the mean was still creeping, so it never stopped."""
        html = self.client.get('/patron/map/').content.decode()
        self.assertIn('STILL_HOLD_M', html)
        self.assertIn('stillAnchor', html)
        self.assertNotIn('POS_SMOOTHING_STILL', html)

    def test_there_is_no_facing_arrow(self):
        """It could not be measured honestly, so it is not drawn."""
        html = self.client.get('/patron/map/').content.decode()
        for gone in ('headingMarker', 'headingDeg', 'deviceorientation',
                     'travelBearing'):
            self.assertNotIn(gone, html, gone + ' should have gone with the arrow')
        # Where they are, and which way the route runs, are still shown.
        self.assertIn('patronMarker', html)
        self.assertIn('drawRouteArrows', html)

    def test_the_shelf_strip_shows_the_neighbours_not_the_whole_board(self):
        """One bar per book wanted 1,536 px on a 192-book board."""
        html = self.client.get('/patron/map/').content.decode()
        self.assertIn('SPINES_EACH_SIDE', html)
        self.assertIn('shelfTrack', html)
        self.assertIn('b.more', html)

    def test_the_marker_is_kept_off_the_furniture(self):
        """Trilateration returns a point, not a place."""
        html = self.client.get('/patron/map/').content.decode()
        for handle in ('clampToFloor', 'pushOffSolids', 'isClearFloor',
                       'nearestClearPoint', 'solids('):
            self.assertIn(handle, html, handle + ' is missing')

    def test_the_map_is_given_what_it_needs_to_know_what_is_solid(self):
        """Shelves, furniture and the stairs all block; rooms bound the floor."""
        d = self.client.get('/patron/map-data/').json()
        for key in ('shelves', 'obstacles', 'stairways', 'rooms'):
            self.assertIn(key, d, key + ' is missing from the map payload')

    def test_the_map_is_told_where_on_the_shelf_to_look(self):
        self._book(title='First', shelf_slot=1, status='Available')
        book = self._book(shelf_slot=2, status='Available')
        self._book(title='Another', shelf_slot=5, status='Available')
        r = self.client.get('/patron/map/', {'book_id': book.book_id})
        target = r.context['target']
        self.assertEqual(target['shelf_slot'], 2)
        self.assertEqual(target['level_number'], 3)
        # Position counts only copies on the shelf.
        self.assertEqual(target['position'], 2)
        self.assertEqual(target['position_of'], 3)
        self.assertIn('2nd book along', target['location'])

    def test_a_book_with_no_shelf_tells_the_map_nothing_extra(self):
        book = self._book(shelf_level=None)
        r = self.client.get('/patron/map/', {'book_id': book.book_id})
        self.assertNotIn('shelf_slot', r.context['target'])


class BookHistoryTests(TestCase):
    """Who has had this copy, without leaving the book you are looking at."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        self.book = Book.objects.create(title='A Book', author='Someone')
        self.patron = Patron.objects.create(fullname='Ben Vergara',
                                            email='ben@example.invalid',
                                            account_status='Active')

    def _loan(self, **kw):
        kw.setdefault('patron', self.patron)
        kw.setdefault('book', self.book)
        kw.setdefault('transaction_type', 'Borrow')
        return Transaction.objects.create(**kw)

    def _details(self):
        return self.client.get('/admin-portal/book-details/%d/' % self.book.book_id).json()

    def test_a_never_borrowed_copy_says_so_rather_than_erroring(self):
        d = self._details()
        self.assertTrue(d['success'])
        self.assertEqual(d['history'], [])
        self.assertEqual(d['history_total'], 0)

    def test_each_loan_reports_its_state(self):
        import datetime
        today = datetime.date.today()
        self._loan(due_date=today + datetime.timedelta(days=7))
        self._loan(due_date=today - datetime.timedelta(days=1), overdue_flag=True)
        self._loan(due_date=today, return_date=today)
        states = sorted(t['state'] for t in self._details()['history'])
        self.assertEqual(states, ['Out', 'Overdue', 'Returned'])

    def test_the_borrower_is_named(self):
        self._loan()
        self.assertEqual(self._details()['history'][0]['patron'], 'Ben Vergara')

    def test_a_long_history_is_capped_but_the_total_is_honest(self):
        from library.views import BOOK_HISTORY_LIMIT
        for _ in range(BOOK_HISTORY_LIMIT + 6):
            self._loan()
        d = self._details()
        self.assertEqual(len(d['history']), BOOK_HISTORY_LIMIT)
        self.assertEqual(d['history_total'], BOOK_HISTORY_LIMIT + 6)

    def test_history_is_staff_side_only(self):
        """The patron endpoint must not say who has had a book."""
        self._loan()
        anon = Client()
        r = anon.get('/admin-portal/book-details/%d/' % self.book.book_id)
        self.assertNotEqual(r.status_code, 200)

        patron_view = anon.get('/patron/book-details/%d/' % self.book.book_id)
        body = patron_view.content.decode('utf-8', 'replace')
        self.assertNotIn('Ben Vergara', body)


class LivePositionTests(TestCase):
    """The number shown is worked out from what is on the shelf right now."""

    def setUp(self):
        self.user = _admin(modules='books,shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                        is_active=True, pixels_per_meter=100)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0,
                                   geometry=[[0, 0], [400, 0], [400, 400], [0, 400]])
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=200,
                                          map_y=200, width=200, depth=40, rotation=0)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=3)
        self.titles = ['Algorithms', 'Botany', 'Calculus', 'Data Structures',
                       'Ecology', 'Fractals', 'Geology', 'History', 'Immunology', 'Jazz']
        for i, t in enumerate(self.titles, start=1):
            Book.objects.create(title=t, author='A', shelf_level=self.level,
                                shelf_slot=i, status='Available')

    def _book(self, title):
        return Book.objects.get(title=title, shelf_level=self.level)

    def _focus(self, title):
        from library.views import level_layout
        return level_layout(self.level, self._book(title))['focus']

    def test_with_nothing_borrowed_the_count_matches_the_slot(self):
        f = self._focus('Immunology')
        self.assertEqual(f['position'], 9)
        self.assertEqual(f['space'], 9)

    def test_a_borrowed_neighbour_moves_the_count_but_not_the_space(self):
        """The case that started this: slot 9 with slot 6 out on loan."""
        Book.objects.filter(title='Fractals').update(status='Borrowed')
        f = self._focus('Immunology')
        self.assertEqual(f['position'], 8, 'the count should skip the borrowed copy')
        self.assertEqual(f['space'], 9, 'its place in the recorded order is unchanged')

    def test_every_way_of_being_off_the_shelf_counts(self):
        for status in ('Borrowed', 'Overdue', 'Lost', 'Being Read', 'For Reshelving'):
            Book.objects.filter(shelf_level=self.level).update(status='Available')
            Book.objects.filter(title='Botany').update(status=status)
            self.assertEqual(self._focus('Calculus')['position'], 2,
                             '%s should not be counted as on the shelf' % status)

    def test_the_neighbours_are_the_ones_actually_there(self):
        Book.objects.filter(title='History').update(status='Borrowed')
        f = self._focus('Immunology')
        self.assertEqual(f['before'], 'Geology')
        self.assertEqual(f['after'], 'Jazz')

    def test_a_copy_that_is_out_reports_no_position(self):
        Book.objects.filter(title='Immunology').update(status='Borrowed')
        f = self._focus('Immunology')
        self.assertFalse(f['on_shelf'])
        self.assertIsNone(f['position'])
        self.assertIn('not on the shelf', book_location_words_for(self._book('Immunology')))

    def test_the_map_marker_follows_the_live_count(self):
        book = self._book('Immunology')
        before = self.client.get('/patron/map/', {'book_id': book.book_id}).context['target']
        Book.objects.filter(title='Fractals').update(status='Borrowed')
        after = self.client.get('/patron/map/', {'book_id': book.book_id}).context['target']
        self.assertEqual(before['position'], 9)
        self.assertEqual(after['position'], 8)
        self.assertEqual(after['position_of'], 9)

    def test_the_layout_marks_the_gaps(self):
        from library.views import level_layout
        Book.objects.filter(title='Calculus').update(status='Borrowed')
        rows = level_layout(self.level, self._book('Immunology'))['rows']
        out = [r for r in rows if not r['on_shelf']]
        self.assertEqual([r['title'] for r in out], ['Calculus'])
        self.assertIsNone(out[0]['position'])
        self.assertEqual(sum(1 for r in rows if r['is_focus']), 1)


def book_location_words_for(book):
    from library.views import book_location_words
    return book_location_words(book)


class ReorderBooksTests(TestCase):
    """Dragging a book into place, and dragging one in from elsewhere."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.other = ShelfLevel.objects.create(shelf=shelf, level_number=2)
        self.books = [Book.objects.create(title=t, author='A', shelf_level=self.level,
                                          shelf_slot=i, status='Available')
                      for i, t in enumerate(['One', 'Two', 'Three'], start=1)]

    def _reorder(self, level, ids):
        return self.client.post('/admin-portal/reorder-books-on-level/', {
            'shelf_level_id': level.shelf_level_id,
            'book_ids': ','.join(str(i) for i in ids)}).json()

    def test_the_new_order_is_written_as_one_two_three(self):
        a, b, c = self.books
        r = self._reorder(self.level, [c.book_id, a.book_id, b.book_id])
        self.assertTrue(r['success'], r)
        for book, expected in ((c, 1), (a, 2), (b, 3)):
            book.refresh_from_db()
            self.assertEqual(book.shelf_slot, expected, book.title)

    def test_a_book_dragged_from_another_level_moves_and_takes_its_place(self):
        stray = Book.objects.create(title='Stray', author='A',
                                    shelf_level=self.other, shelf_slot=1)
        a, b, c = self.books
        r = self._reorder(self.level, [a.book_id, stray.book_id, b.book_id, c.book_id])
        self.assertTrue(r['success'], r)
        self.assertEqual(r['moved'], 1)
        stray.refresh_from_db()
        self.assertEqual(stray.shelf_level_id, self.level.shelf_level_id)
        self.assertEqual(stray.shelf_slot, 2)

    def test_renumbering_closes_gaps_left_by_hand_entry(self):
        """Slots are an ordering key, so a tidy-up makes them dense again."""
        for book, slot in zip(self.books, (5, 40, 900)):
            book.shelf_slot = slot
            book.save()
        self._reorder(self.level, [b.book_id for b in self.books])
        slots = list(Book.objects.filter(shelf_level=self.level)
                     .order_by('shelf_slot').values_list('shelf_slot', flat=True))
        self.assertEqual(slots, [1, 2, 3])

    def test_an_order_naming_a_book_that_is_gone_changes_nothing(self):
        """All or nothing: half an ordering is worse than none."""
        a, b, c = self.books
        r = self._reorder(self.level, [c.book_id, 999999, a.book_id])
        self.assertFalse(r['success'])
        a.refresh_from_db()
        self.assertEqual(a.shelf_slot, 1)

    def test_an_empty_order_is_refused(self):
        self.assertFalse(self._reorder(self.level, [])['success'])

    def _board_node(self, data):
        def walk(nodes):
            for n in nodes:
                if n['type'] == 'shelflevel' and n['id'] == self.level.shelf_level_id:
                    return n
                found = walk(n.get('children', []))
                if found:
                    return found
            return None
        return walk(data['tree'])

    def test_a_board_returns_its_books_in_shelf_order(self):
        """Where the tree's book list went."""
        a, b, c = self.books
        self._reorder(self.level, [c.book_id, b.book_id, a.book_id])
        d = self.client.get('/admin-portal/board-books/',
                            {'level': self.level.shelf_level_id}).json()
        self.assertTrue(d['success'], d)
        self.assertEqual([bk['title'] for bk in d['books']], ['Three', 'Two', 'One'])
        self.assertEqual([bk['slot'] for bk in d['books']], [1, 2, 3])

    def test_the_tree_counts_a_boards_books_instead_of_listing_them(self):
        data = self.client.get('/admin-portal/get-shelf-tree/').json()
        board = self._board_node(data)
        self.assertIsNotNone(board)
        self.assertEqual(board['book_count'], 3)
        self.assertEqual(board.get('children'), [],
                         'the tree carries structure, not the catalogue')

    def test_a_board_that_is_gone_is_reported_not_crashed(self):
        gone = self.level.shelf_level_id
        self.level.delete()
        d = self.client.get('/admin-portal/board-books/', {'level': gone}).json()
        self.assertFalse(d['success'])
        self.assertIn('no longer exists', d['error'])



class ShelfColumnTreeTests(TestCase):
    """Shelf > Column > Level, but only on a shelf that has more than one."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        self.room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)

    def _shelf(self, name):
        return Shelf.objects.create(room=self.room, name=name, map_x=0, map_y=0)

    def _tree(self):
        return self.client.get('/admin-portal/get-shelf-tree/').json()['tree']

    def _shelf_node(self, name):
        for fp in self._tree():
            for room in fp['children']:
                for shelf in room['children']:
                    if shelf['name'] == name:
                        return shelf
        return None

    def test_a_single_column_shelf_lists_its_boards_directly(self):
        shelf = self._shelf('Plain')
        for n in (1, 2, 3):
            ShelfLevel.objects.create(shelf=shelf, level_number=n)
        node = self._shelf_node('Plain')
        self.assertEqual([c['type'] for c in node['children']],
                         ['shelflevel'] * 3)
        self.assertEqual([c['name'] for c in node['children']],
                         ['Level 1', 'Level 2', 'Level 3'])

    def test_a_divided_shelf_groups_its_boards_under_columns(self):
        shelf = self._shelf('Divided')
        for column in (1, 2):
            for n in (1, 2):
                ShelfLevel.objects.create(shelf=shelf, level_number=n,
                                          column_number=column)
        node = self._shelf_node('Divided')
        self.assertEqual([c['type'] for c in node['children']],
                         ['shelfcolumn', 'shelfcolumn'])
        self.assertEqual([c['name'] for c in node['children']],
                         ['Column 1', 'Column 2'])
        for column in node['children']:
            self.assertEqual([b['name'] for b in column['children']],
                             ['Level 1', 'Level 2'])

    def test_a_board_under_a_column_does_not_repeat_the_column(self):
        """The heading says it once. Saying it again on every child is noise."""
        shelf = self._shelf('Divided')
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=1)
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=2)
        node = self._shelf_node('Divided')
        for column in node['children']:
            for board in column['children']:
                self.assertNotIn('Column', board['name'])

    def test_a_board_still_names_its_column_when_it_stands_alone(self):
        """A dropdown row or a QR label has no heading above it to inherit."""
        shelf = self._shelf('Divided')
        board = ShelfLevel.objects.create(shelf=shelf, level_number=2, column_number=3)
        board.shelf = shelf
        self.assertEqual(board.board_label, 'Level 2')
        self.assertEqual(board.label, 'Level 2, Column 3')

    def test_a_column_carries_the_books_of_the_boards_in_it(self):
        shelf = self._shelf('Divided')
        one = ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=1)
        two = ShelfLevel.objects.create(shelf=shelf, level_number=2, column_number=1)
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=2)
        for level, count in ((one, 2), (two, 3)):
            for i in range(count):
                Book.objects.create(title='B%d-%d' % (level.pk, i), author='X',
                                    shelf_level=level, status='Available')
        node = self._shelf_node('Divided')
        first, second = node['children']
        self.assertEqual(first['book_count'], 5)
        self.assertEqual(second['book_count'], 0)

    def test_a_top_board_is_not_renumbered_by_the_column_grouping(self):
        """A top is not level N, and dividing the bay does not make it one."""
        shelf = self._shelf('Divided')
        ShelfLevel.objects.create(shelf=shelf, level_number=9, column_number=1,
                                  is_top=True)
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=2)
        node = self._shelf_node('Divided')
        names = [b['name'] for c in node['children'] for b in c['children']]
        self.assertIn('Top', names)

    def test_boards_with_no_column_recorded_count_as_one_column(self):
        """Most of the real data stores NULL, not 1, for an undivided bay."""
        shelf = self._shelf('Null columns')
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=None)
        ShelfLevel.objects.create(shelf=shelf, level_number=2, column_number=None)
        node = self._shelf_node('Null columns')
        self.assertEqual([c['type'] for c in node['children']],
                         ['shelflevel', 'shelflevel'])

    def test_a_null_column_mixed_with_a_numbered_one_still_groups(self):
        """Half-divided is divided: the reader has to be able to tell them apart."""
        shelf = self._shelf('Half')
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=None)
        ShelfLevel.objects.create(shelf=shelf, level_number=1, column_number=2)
        node = self._shelf_node('Half')
        self.assertEqual([c['name'] for c in node['children']],
                         ['Column 1', 'Column 2'])


class ShelfReadAuditTests(TestCase):
    """Reading a board against a list, and filing what was found."""

    def setUp(self):
        self.user = _admin(modules='inventory')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=1)
        self.books = [self._book('One', 1), self._book('Two', 2), self._book('Three', 3)]

    def _book(self, title, slot, status='Available'):
        return Book.objects.create(title=title, author='X', genre='REF',
                                   shelf_level=self.level, shelf_slot=slot,
                                   status=status, qr_code='qr-' + title.lower())

    @property
    def audits(self):
        from library.models import StockAudit
        return StockAudit.objects

    def _sheet(self):
        return self.client.get('/admin-portal/inventory/audit/sheet/',
                               {'level': self.level.shelf_level_id}).json()

    def _file(self, **kw):
        data = {'level': self.level.shelf_level_id}
        data.update(kw)
        return self.client.post('/admin-portal/inventory/audit/file/', data).json()

    # The sheet

    def test_the_sheet_lists_the_board_in_shelf_order(self):
        """The order is half the point: a shelf-read finds books out of place."""
        d = self._sheet()
        self.assertTrue(d['success'], d)
        self.assertEqual([b['title'] for b in d['books']], ['One', 'Two', 'Three'])
        self.assertEqual(d['expected_count'], 3)

    def test_the_sheet_carries_the_code_that_is_printed_on_the_sticker(self):
        """The old audit looked up a copy label no printed sticker carries."""
        d = self._sheet()
        self.assertEqual([b['qr_code'] for b in d['books']],
                         ['qr-one', 'qr-two', 'qr-three'])

    def test_a_borrowed_copy_is_shown_but_not_asked_about(self):
        """The librarian is standing at a gap. Saying why beats hiding it."""
        self.books[1].status = 'Borrowed'
        self.books[1].save(update_fields=['status'])
        d = self._sheet()
        row = next(b for b in d['books'] if b['title'] == 'Two')
        self.assertFalse(row['expected_present'])
        self.assertEqual(d['expected_count'], 2)
        self.assertEqual(d['elsewhere_count'], 1)

    def test_a_written_off_copy_is_not_on_the_sheet_at_all(self):
        self.books[0].status = 'Lost'
        self.books[0].save(update_fields=['status'])
        d = self._sheet()
        self.assertEqual([b['title'] for b in d['books']], ['Two', 'Three'])

    # Filing

    def test_filing_records_what_was_confirmed_by_hand_against_what_was_swept(self):
        """A swept board and a board read spine by spine are not the same claim."""
        d = self._file(found_ids=[self.books[0].book_id],
                       bulk_ids=[self.books[1].book_id, self.books[2].book_id])
        self.assertTrue(d['success'], d)
        self.assertEqual((d['confirmed'], d['individually'], d['in_bulk']), (3, 1, 2))
        audit = self.audits.get(audit_id=d['audit_id'])
        self.assertEqual(audit.found_count, 3)
        self.assertEqual(audit.scanned_count, 1)

    def test_a_copy_not_found_is_marked_missing_and_not_written_off(self):
        """One bad count must not destroy a book that was simply misplaced."""
        gone = self.books[2]
        d = self._file(found_ids=[self.books[0].book_id, self.books[1].book_id],
                       missing_ids=[gone.book_id])
        self.assertEqual(d['flagged'], 1)
        gone.refresh_from_db()
        self.assertEqual(gone.status, 'Missing')
        self.assertEqual(gone.audit_misses, 1)
        self.assertIsNotNone(gone.missing_since)
        self.assertNotEqual(gone.status, 'Lost')

    def test_a_second_miss_counts_but_keeps_the_original_date(self):
        """The gap is measured from when it went, not from the last count."""
        gone = self.books[2]
        self._file(missing_ids=[gone.book_id], found_ids=[self.books[0].book_id])
        gone.refresh_from_db()
        first_seen_missing = gone.missing_since
        self._file(missing_ids=[gone.book_id], found_ids=[self.books[0].book_id])
        gone.refresh_from_db()
        self.assertEqual(gone.audit_misses, 2)
        self.assertEqual(gone.missing_since, first_seen_missing)

    def test_a_missing_copy_that_turns_up_is_recovered(self):
        """The good news in a count, and the reason a miss writes nothing off."""
        gone = self.books[2]
        self._file(missing_ids=[gone.book_id], found_ids=[self.books[0].book_id])
        d = self._file(found_ids=[b.book_id for b in self.books])
        self.assertEqual(d['recovered'], 1)
        gone.refresh_from_db()
        self.assertEqual(gone.status, 'Available')
        self.assertIsNone(gone.missing_since)
        self.assertEqual(gone.audit_misses, 0)

    def test_a_borrowed_copy_cannot_be_marked_missing(self):
        """Not finding it on the shelf is the expected result, not a discrepancy."""
        out = self.books[1]
        out.status = 'Borrowed'
        out.save(update_fields=['status'])
        d = self._file(found_ids=[self.books[0].book_id], missing_ids=[out.book_id])
        self.assertEqual(d['flagged'], 0)
        out.refresh_from_db()
        self.assertEqual(out.status, 'Borrowed')

    def test_a_confirmation_beats_a_sweep_and_a_miss_beats_both(self):
        """The verdicts arrive as three lists and a copy can only be in one."""
        book = self.books[0]
        d = self._file(found_ids=[book.book_id], bulk_ids=[book.book_id],
                       missing_ids=[book.book_id])
        self.assertEqual((d['confirmed'], d['flagged']), (0, 1))
        book.refresh_from_db()
        self.assertEqual(book.status, 'Missing')

    def test_confirming_a_copy_records_when_it_was_last_seen(self):
        self._file(found_ids=[self.books[0].book_id])
        self.books[0].refresh_from_db()
        self.assertIsNotNone(self.books[0].last_seen)

    def test_a_board_nobody_marked_cannot_be_filed(self):
        """Filing an untouched sheet would date a shelf nobody walked."""
        d = self._file()
        self.assertFalse(d['success'])
        self.assertIn('Mark what you found', d['error'])
        self.assertEqual(self.audits.count(), 0)

    def test_a_book_from_another_board_cannot_be_counted_here(self):
        elsewhere = Book.objects.create(title='Stranger', author='X',
                                        status='Available')
        d = self._file(found_ids=[self.books[0].book_id, elsewhere.book_id])
        self.assertEqual(d['confirmed'], 1)
        elsewhere.refresh_from_db()
        self.assertIsNone(elsewhere.last_seen)

    def test_the_sweep_writes_a_different_mark_from_a_tap(self):
        """The bug: "the rest are here" wrote the same mark a tap writes."""
        html = self.client.get('/admin-portal/inventory/?tab=audit').content.decode()
        self.assertIn("sheet.marks[b.book_id] = 'bulk'", html,
                      'the sweep must write its own mark')
        self.assertIn("if (m === 'here') body.append('found_ids'", html)
        self.assertIn("body.append('bulk_ids'", html)

    def test_the_board_picker_reports_the_oldest_confirmation_not_the_newest(self):
        """A board is only as counted as its least-recently-seen book."""
        self._file(found_ids=[self.books[0].book_id])
        html = self.client.get('/admin-portal/inventory/?tab=audit').content.decode()
        self.assertIn('never read', html,
                      'two books here have never been confirmed')


class ShapedStairTests(TestCase):
    """A staircase that turns, drawn once."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True)
        self.upstairs = FloorPlan.objects.create(name='First', floor_number=2)
        # A tall well: 200 across, 400 along, travel running north-south.
        self.well = [[100, 100], [300, 100], [300, 500], [100, 500]]

    def _add(self, **kw):
        data = {'floor_plan_id': self.plan.floor_plan_id,
                'geometry': json.dumps(self.well)}
        data.update(kw)
        return self.client.post('/admin-portal/add-stairway/', data).json()

    @staticmethod
    def _box(geometry):
        xs = [p[0] for p in geometry]
        ys = [p[1] for p in geometry]
        return min(xs), min(ys), max(xs), max(ys)

    # The shape

    def test_a_straight_stair_stores_no_flights(self):
        """The overwhelming majority. Footprint and bearing say it all."""
        r = self._add(shape='straight')
        self.assertTrue(r['success'], r)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        self.assertIsNone(st.flights)
        self.assertEqual(r['stairway']['shape'], 'straight')

    def test_a_half_turn_becomes_two_flights_and_a_landing(self):
        r = self._add(shape='half', bearing=180)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        self.assertEqual([f['kind'] for f in st.flights],
                         ['flight', 'landing', 'flight'])
        self.assertEqual(r['stairway']['shape'], 'half')

    def test_the_two_flights_of_a_half_turn_climb_opposite_ways(self):
        """That is what turning about means, and it is why one bearing failed."""
        r = self._add(shape='half', bearing=180)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        first, _, second = st.flights
        self.assertEqual(first['bearing'], 180)
        self.assertEqual(second['bearing'], 0)

    def test_the_flights_of_a_half_turn_stand_side_by_side(self):
        r = self._add(shape='half', bearing=180)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        first, landing, second = st.flights
        fx0, fy0, fx1, fy1 = self._box(first['geometry'])
        sx0, sy0, sx1, sy1 = self._box(second['geometry'])
        self.assertLess(fx1, sx0, 'the flights must not overlap')
        self.assertEqual((fy0, fy1), (sy0, sy1), 'they run the same length')

    def test_the_landing_sits_at_the_end_the_flights_climb_towards(self):
        """A half-landing is where you turn round, so it is at the far end."""
        r = self._add(shape='half', bearing=180)          # travelling south
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        first, landing, _ = st.flights
        self.assertGreater(self._box(landing['geometry'])[1],
                           self._box(first['geometry'])[1],
                           'travelling south, the landing is at the south end')

        r = self._add(shape='half', bearing=0)            # travelling north
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        first, landing, _ = st.flights
        self.assertLess(self._box(landing['geometry'])[1],
                        self._box(first['geometry'])[1])

    def test_the_landing_spans_the_whole_well(self):
        """You walk off one flight and onto the other across it."""
        r = self._add(shape='half', bearing=180)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        landing = st.flights[1]
        lx0, _, lx1, _ = self._box(landing['geometry'])
        wx0, _, wx1, _ = self._box(self.well)
        self.assertEqual((lx0, lx1), (wx0, wx1))

    def test_a_quarter_turn_leaves_its_landing_sideways(self):
        r = self._add(shape='quarter', bearing=180)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        first, _, second = st.flights
        self.assertEqual((first['bearing'], second['bearing']), (180, 270))

    def test_a_well_too_small_to_divide_stays_one_run(self):
        """Better a straight stair than two flights of three steps."""
        tiny = [[0, 0], [20, 0], [20, 20], [0, 20]]
        r = self.client.post('/admin-portal/add-stairway/', {
            'floor_plan_id': self.plan.floor_plan_id,
            'geometry': json.dumps(tiny), 'shape': 'half'}).json()
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        self.assertIsNone(st.flights)

    # The direction

    def test_the_bearing_is_read_off_the_shape_when_none_is_given(self):
        """It was a number box asking to convert a direction just drawn."""
        r = self._add()                                   # no bearing sent
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        self.assertEqual(st.bearing, 0, 'the well is taller than it is wide')

        wide = [[100, 100], [500, 100], [500, 300], [100, 300]]
        r = self.client.post('/admin-portal/add-stairway/', {
            'floor_plan_id': self.plan.floor_plan_id,
            'geometry': json.dumps(wide)}).json()
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        self.assertEqual(st.bearing, 90, 'a wide well runs left to right')

    def test_an_explicit_bearing_still_wins(self):
        r = self._add(bearing=270)
        st = Stairway.objects.get(stairway_id=r['stairway']['stairway_id'])
        self.assertEqual(st.bearing, 270)

    def test_turning_a_shaped_stair_redraws_what_is_inside_it(self):
        """Arrows that disagree with the footprint they sit on are worse than none."""
        r = self._add(shape='half', bearing=180)
        sid = r['stairway']['stairway_id']
        self.client.post('/admin-portal/edit-stairway/',
                         {'stairway_id': sid, 'bearing': 0})
        st = Stairway.objects.get(stairway_id=sid)
        self.assertEqual([f['bearing'] for f in st.flights], [0, 0, 180])
        self.assertEqual(len(st.flights), 3, 'it is still a half turn')

    # What gets drawn

    def test_each_flight_gets_its_own_treads(self):
        r = self._add(shape='half', bearing=180)
        parts = r['stairway']['parts']
        self.assertEqual([p['kind'] for p in parts], ['flight', 'landing', 'flight'])
        self.assertTrue(parts[0]['treads'])
        self.assertTrue(parts[2]['treads'])

    def test_a_landing_has_no_treads_because_it_is_a_floor(self):
        """Steps across it would say you climb it."""
        r = self._add(shape='half', bearing=180)
        landing = r['stairway']['parts'][1]
        self.assertEqual(landing['treads'], [])

    def test_a_lift_has_no_treads_at_all(self):
        """Drawing steps on one would be wrong, not merely decorative."""
        r = self._add(kind='Elevator', shape='half')
        self.assertEqual(r['stairway']['parts'], [])
        self.assertEqual(r['stairway']['treads'], [])

    def test_a_straight_stair_still_reports_treads_for_the_maps(self):
        """Four maps draw a stair from this. None of them may go blank."""
        r = self._add(shape='straight')
        self.assertTrue(r['stairway']['treads'])
        self.assertEqual(len(r['stairway']['parts']), 1)


class UnshelvedInMoverTests(TestCase):
    """Books on no board, opened where they can be put away."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=1)

    def _book(self, title, level=None, slot=None, status='Available'):
        return Book.objects.create(title=title, author='X', genre='REF',
                                   shelf_level=level, shelf_slot=slot, status=status)

    def _pile(self):
        return self.client.get('/admin-portal/board-books/', {'level': 'none'}).json()

    def test_the_pile_opens_like_any_other_board(self):
        self._book('Homeless one')
        self._book('Homeless two')
        self._book('Shelved', self.level, 1)
        d = self._pile()
        self.assertTrue(d['success'], d)
        self.assertEqual(d['level'], 'none')
        self.assertEqual(d['shelf'], 'Not shelved')
        self.assertEqual([b['title'] for b in d['books']],
                         ['Homeless one', 'Homeless two'])

    def test_a_book_with_no_board_has_no_position(self):
        """A number here would invent a place along a shelf it is not on."""
        self._book('Homeless')
        self.assertIsNone(self._pile()['books'][0]['slot'])

    def test_a_written_off_copy_is_not_waiting_to_be_shelved(self):
        """It is gone, not mislaid, and offering it as work would be a lie."""
        self._book('Lost one', status='Lost')
        self._book('Real one')
        self.assertEqual([b['title'] for b in self._pile()['books']], ['Real one'])

    def test_shelving_from_the_pile_is_the_same_move_as_any_other(self):
        a, b = self._book('One'), self._book('Two')
        r = self.client.post('/admin-portal/move-books/', {
            'ids': '%d,%d' % (a.book_id, b.book_id),
            'level': self.level.shelf_level_id}).json()
        self.assertTrue(r['success'], r)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual((a.shelf_level, b.shelf_level), (self.level, self.level))
        self.assertEqual(sorted([a.shelf_slot, b.shelf_slot]), [1, 2])

    def test_books_shelved_from_the_pile_land_after_the_residents(self):
        self._book('Resident', self.level, 1)
        incoming = self._book('Arriving')
        self.client.post('/admin-portal/move-books/',
                         {'ids': str(incoming.book_id),
                          'level': self.level.shelf_level_id})
        incoming.refresh_from_db()
        self.assertEqual(incoming.shelf_slot, 2)

    def test_the_page_says_how_many_are_waiting(self):
        """Otherwise they sit there indefinitely with the page looking complete."""
        for i in range(3):
            self._book('Homeless %d' % i)
        html = self.client.get('/admin-portal/shelf-manager/').content.decode()
        self.assertIn('3 not on a shelf', html)

    def test_the_button_is_rendered_even_with_none_waiting(self):
        """Taking a book off a shelf must be able to make it appear."""
        html = self.client.get('/admin-portal/shelf-manager/').content.decode()
        self.assertIn('unshelvedBtn', html)
        self.assertIn('hidden', html[html.index('unshelvedBtn') - 80:
                                     html.index('unshelvedBtn') + 120])

    def test_the_tree_carries_the_count_so_it_survives_a_move(self):
        """It is reloaded after every move, and a stale count is a wrong one."""
        self._book('Homeless')
        d = self.client.get('/admin-portal/get-shelf-tree/').json()
        self.assertEqual(d['unshelved'], 1)

    def test_the_count_falls_as_books_are_put_away(self):
        book = self._book('Homeless')
        self.client.post('/admin-portal/move-books/',
                         {'ids': str(book.book_id),
                          'level': self.level.shelf_level_id})
        d = self.client.get('/admin-portal/get-shelf-tree/').json()
        self.assertEqual(d['unshelved'], 0)

    def test_taking_a_book_off_a_shelf_puts_it_in_the_pile(self):
        """The two directions are the same tool, which is the point."""
        book = self._book('Shelved', self.level, 1)
        self.client.post('/admin-portal/move-books/',
                         {'ids': str(book.book_id), 'level': 'none'})
        self.assertEqual([b['title'] for b in self._pile()['books']], ['Shelved'])


class PrintSeveralFloorsTests(TestCase):
    """A building printed in one pass rather than one storey at a time."""

    def setUp(self):
        self.user = _admin(modules='indoor_map')
        self.client = _signed_in(self.user)
        self.ground = self._plan('Ground', 1)
        self.first = self._plan('First', 2)
        self.draft = FloorPlan.objects.create(name='Attic', floor_number=3,
                                              is_active=False)

    def _plan(self, name, number):
        plan = FloorPlan.objects.create(name=name, floor_number=number, is_active=True)
        room = Room.objects.create(floor_plan=plan, name=name + ' room',
                                   map_x=100, map_y=100,
                                   geometry=[[0, 0], [200, 0], [200, 200], [0, 200]])
        Shelf.objects.create(room=room, name='Shelf on ' + name, map_x=50, map_y=50)
        return plan

    def _get(self, query=''):
        return self.client.get('/admin-portal/floor-plan/print/' + query)

    def _pages(self, r):
        return r.content.decode().count('class="floor-page"')

    def test_one_floor_by_default(self):
        """A bare link keeps doing what it always did."""
        r = self._get()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._pages(r), 1)

    def test_a_named_floor_is_the_one_printed(self):
        r = self._get('?floor=%d' % self.first.floor_plan_id)
        self.assertEqual(self._pages(r), 1)
        self.assertContains(r, 'Shelf on First')
        self.assertNotContains(r, 'Shelf on Ground')

    def test_all_floors_come_out_as_one_document(self):
        """The point: one trip to the printer, not one per storey."""
        r = self._get('?floor=all')
        self.assertEqual(self._pages(r), 2)
        self.assertContains(r, 'Shelf on Ground')
        self.assertContains(r, 'Shelf on First')

    def test_two_named_floors_print_together(self):
        r = self._get('?floor=%d&floor=%d'
                      % (self.ground.floor_plan_id, self.first.floor_plan_id))
        self.assertEqual(self._pages(r), 2)

    def test_a_floor_out_of_service_is_never_printed(self):
        """A wall map of a storey nobody may enter is a wrong map."""
        r = self._get('?floor=%d' % self.draft.floor_plan_id)
        self.assertNotContains(r, 'Attic')

    def test_nonsense_falls_back_to_one_floor_rather_than_failing(self):
        for query in ('?floor=999999', '?floor=abc', '?floor='):
            r = self._get(query)
            self.assertEqual(r.status_code, 200, query)
            self.assertEqual(self._pages(r), 1, query)

    def test_each_floor_gets_its_own_page_break(self):
        r = self._get('?floor=all')
        self.assertContains(r, 'page-break-after')

    def test_the_indoor_map_offers_both(self):
        html = self.client.get('/admin-portal/indoor-map/').content.decode()
        self.assertIn('Print this floor', html)
        self.assertIn('Print all floors', html)


class ShelfRoomTests(TestCase):
    """Which room a shelf is filed under, and correcting it."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True)
        self.left = Room.objects.create(
            floor_plan=self.plan, name='Reading Room', map_x=50, map_y=50,
            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        self.right = Room.objects.create(
            floor_plan=self.plan, name='Study Room', map_x=150, map_y=50,
            geometry=[[100, 0], [200, 0], [200, 100], [100, 100]])

    def _shelf(self, name, x, y, room):
        return Shelf.objects.create(room=room, name=name, map_x=x, map_y=y)

    def _mismatches(self):
        from library.views import _shelf_room_mismatches
        return _shelf_room_mismatches()

    # Setting it by hand

    def test_a_shelf_can_be_moved_to_another_room(self):
        shelf = self._shelf('Shelf A', 50, 50, self.left)
        r = self.client.post('/admin-portal/edit-shelf/',
                             {'shelf_id': shelf.shelf_id,
                              'room_id': self.right.room_id}).json()
        self.assertTrue(r['success'], r)
        shelf.refresh_from_db()
        self.assertEqual(shelf.room, self.right)

    def test_the_shelf_does_not_move_on_the_plan_when_its_room_changes(self):
        """Filing is not placing. Only where it is listed changes."""
        shelf = self._shelf('Shelf A', 50, 50, self.left)
        self.client.post('/admin-portal/edit-shelf/',
                         {'shelf_id': shelf.shelf_id, 'room_id': self.right.room_id})
        shelf.refresh_from_db()
        self.assertEqual((shelf.map_x, shelf.map_y), (50, 50))

    def test_a_room_on_another_floor_is_refused(self):
        """A shelf cannot be in a room upstairs."""
        upstairs = FloorPlan.objects.create(name='First', floor_number=2, is_active=True)
        elsewhere = Room.objects.create(floor_plan=upstairs, name='Attic',
                                        map_x=0, map_y=0)
        shelf = self._shelf('Shelf A', 50, 50, self.left)
        r = self.client.post('/admin-portal/edit-shelf/',
                             {'shelf_id': shelf.shelf_id,
                              'room_id': elsewhere.room_id}).json()
        self.assertFalse(r['success'])
        shelf.refresh_from_db()
        self.assertEqual(shelf.room, self.left)

    def test_leaving_the_room_out_changes_nothing(self):
        """Renaming a shelf must not silently re-file it."""
        shelf = self._shelf('Shelf A', 50, 50, self.left)
        self.client.post('/admin-portal/edit-shelf/',
                         {'shelf_id': shelf.shelf_id, 'name': 'Shelf One'})
        shelf.refresh_from_db()
        self.assertEqual((shelf.name, shelf.room), ('Shelf One', self.left))

    # Bulk correction

    def test_a_shelf_drawn_in_another_room_is_reported(self):
        self._shelf('Wanderer', 150, 50, self.left)      # drawn right, filed left
        rows = self._mismatches()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['name'], 'Wanderer')
        self.assertEqual(rows[0]['filed'], 'Reading Room')
        self.assertEqual(rows[0]['drawn'], 'Study Room')

    def test_a_shelf_in_the_room_it_says_is_not_reported(self):
        self._shelf('Settled', 50, 50, self.left)
        self.assertEqual(self._mismatches(), [])

    def test_a_shelf_inside_no_room_is_left_alone(self):
        """A shelf in a corridor nobody drew a polygon round is not a mistake."""
        self._shelf('Corridor', 900, 900, self.left)
        self.assertEqual(self._mismatches(), [])

    def test_correcting_them_files_each_where_it_is_drawn(self):
        a = self._shelf('Wanderer', 150, 50, self.left)
        b = self._shelf('Other way', 50, 50, self.right)
        settled = self._shelf('Settled', 60, 60, self.left)
        r = self.client.post('/admin-portal/fix-shelf-rooms/').json()
        self.assertTrue(r['success'], r)
        self.assertEqual(r['fixed'], 2)
        a.refresh_from_db(); b.refresh_from_db(); settled.refresh_from_db()
        self.assertEqual(a.room, self.right)
        self.assertEqual(b.room, self.left)
        self.assertEqual(settled.room, self.left)

    def test_correcting_twice_finds_nothing_the_second_time(self):
        self._shelf('Wanderer', 150, 50, self.left)
        self.client.post('/admin-portal/fix-shelf-rooms/')
        r = self.client.post('/admin-portal/fix-shelf-rooms/').json()
        self.assertEqual(r['fixed'], 0)

    def test_the_page_says_nothing_when_every_shelf_agrees(self):
        self._shelf('Settled', 50, 50, self.left)
        html = self.client.get('/admin-portal/shelf-manager/').content.decode()
        self.assertNotIn('roomMismatchBar', html)

    def test_the_page_names_the_shelves_that_disagree(self):
        self._shelf('Wanderer', 150, 50, self.left)
        html = self.client.get('/admin-portal/shelf-manager/').content.decode()
        self.assertIn('roomMismatchBar', html)
        self.assertIn('Wanderer', html)
        self.assertIn('Study Room', html)




class AnalyticsLayoutTests(TestCase):
    """The page answers before it explains, and one section at a time."""

    def setUp(self):
        self.user = _admin()
        self.client = _signed_in(self.user)
        self.start = timezone.localdate() - timedelta(days=7)
        self.end = timezone.localdate()

    def _parts(self):
        from library import analytics
        return dict(
            arrivals=analytics.visits_by_hour(self.start, self.end),
            occupancy=analytics.occupancy_by_hour(self.start, self.end),
            unshelved=analytics.unshelved_summary(),
            conditions=analytics.books_by_condition(),
            borrowed_books=analytics.most_borrowed_books(self.start, self.end),
            penalties=analytics.penalties_over_time(self.start, self.end, 'day'),
        )

    def _headline(self):
        from library import analytics
        return analytics.headline(**self._parts())

    # -- the headline ------------------------------------------------------

    def test_it_leads_with_a_figure_from_every_section(self):
        """A strip that only covered the collection would answer a third of the page."""
        labels = [f['label'] for f in self._headline()]
        self.assertIn('Visits', labels)
        self.assertIn('Loans', labels)
        self.assertIn('Copies held', labels)

    def test_every_figure_says_which_period_it_belongs_to(self):
        """Some figures follow the date range and some are current."""
        for f in self._headline():
            self.assertIn(f['scope'], ('period', 'now'), f)

    def test_the_collection_figures_are_not_claimed_to_be_period_scoped(self):
        scopes = dict((f['label'], f['scope']) for f in self._headline())
        self.assertEqual(scopes['Copies held'], 'now')
        self.assertEqual(scopes['Visits'], 'period')

    def test_it_cannot_disagree_with_the_chart_below_it(self):
        """Built from the dicts the builders already returned, not re-queried."""
        parts = self._parts()
        from library import analytics
        tiles = dict((f['label'], f['value']) for f in analytics.headline(**parts))
        self.assertEqual(tiles['Visits'], parts['arrivals']['total'])
        self.assertEqual(tiles['Loans'], parts['borrowed_books']['total'])
        self.assertEqual(tiles['Copies held'], parts['unshelved']['total'])

    def test_a_clean_collection_reads_as_good_not_as_nothing(self):
        """Zero books needing attention is a result, and it is the good one."""
        tone = dict((f['label'], f['tone']) for f in self._headline())
        self.assertEqual(tone['Needs attention'], 'good')

    def test_the_page_shows_the_strip(self):
        html = self.client.get('/admin-portal/analytics/').content.decode()
        self.assertIn('figure-tile', html)
        self.assertIn('Copies held', html)

    # -- the tabs ----------------------------------------------------------

    def test_all_three_sections_are_on_the_page(self):
        """Theme switching runs in the browser, so every page includes it."""
        html = self.client.get('/admin-portal/analytics/').content.decode()
        for panel in ('panel-visits', 'panel-collection', 'panel-borrowing'):
            self.assertIn(panel, html)

    def test_it_opens_on_visits_by_default(self):
        html = self.client.get('/admin-portal/analytics/').content.decode()
        self.assertIn('id="viewField" value="visits"', html)

    def test_a_tab_can_be_linked_to(self):
        html = self.client.get('/admin-portal/analytics/',
                               {'view': 'collection'}).content.decode()
        self.assertIn('id="viewField" value="collection"', html)

    def test_a_view_nobody_recognises_falls_back_rather_than_failing(self):
        """Query strings get edited, truncated and pasted half-copied."""
        r = self.client.get('/admin-portal/analytics/', {'view': 'nonsense'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('id="viewField" value="visits"', r.content.decode())

    def test_the_date_form_carries_the_tab(self):
        """Otherwise every date change drops the reader back on the first tab."""
        html = self.client.get('/admin-portal/analytics/',
                               {'view': 'borrowing'}).content.decode()
        form = html[html.index('id="periodForm"'):html.index('</form>', html.index('id="periodForm"'))]
        self.assertIn('name="view"', form)

    def test_the_range_still_filters_when_a_tab_is_named(self):
        """The tab is presentation. It must not touch what gets counted."""
        a = self.client.get('/admin-portal/analytics/',
                            {'start': '2020-01-01', 'end': '2020-01-31'})
        self.assertEqual(a.status_code, 200)
        self.assertIn('2020-01-01', a.content.decode())

    # -- the layout itself -------------------------------------------------

    def test_few_category_charts_are_drawn_as_shares_not_as_bars(self):
        """Three bars in a box built for twenty-four hours is mostly white space."""
        html = self.client.get('/admin-portal/analytics/').content.decode()
        self.assertIn('fa-chart-pie', html, 'the proportion list is not on the page')

    def test_the_proportion_bars_are_shares_of_the_whole(self):
        """Scaling to the biggest row would end every list in one full bar."""
        tpl = open('templates/admin/_proportions.html', encoding='utf-8').read()
        self.assertIn('width: {{ m.share }}%', tpl)

    def test_the_period_form_says_when_it_does_not_apply(self):
        """A control that vanishes looks broken. One that explains itself does not."""
        html = self.client.get('/admin-portal/analytics/').content.decode()
        self.assertIn('period-note', html)
        self.assertIn('do not apply here', html)

    def test_the_tablist_is_reachable_without_a_mouse(self):
        html = self.client.get('/admin-portal/analytics/').content.decode()
        self.assertIn('role="tablist"', html)
        self.assertIn('role="tabpanel"', html)
        self.assertIn('aria-selected', html)

class RestrictedRoomTests(TestCase):
    """A room patrons may not enter: drawn, named, and not walked into."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, pixels_per_meter=100)
        # Two rooms sharing the wall x = 300, joined by a door.
        self.public = Room.objects.create(
            floor_plan=self.plan, name='Reading Room', map_x=200, map_y=200,
            geometry=[[100, 100], [300, 100], [300, 300], [100, 300]])
        self.store = Room.objects.create(
            floor_plan=self.plan, name='Store', map_x=400, map_y=200,
            geometry=[[300, 100], [500, 100], [500, 300], [300, 300]])
        Door.objects.create(room=self.public, room_b=self.store,
                            map_x=300, map_y=200, width=70, rotation=90)

    def _close_the_store(self):
        self.store.patron_access = False
        self.store.save()

    def _generate(self):
        return self.client.post('/admin-portal/generate-waypoints/',
                                {'floor_plan_id': self.plan.floor_plan_id}).json()

    def _readiness(self):
        return self.client.get('/admin-portal/floor-plan-readiness/',
                               {'floor_plan_id': self.plan.floor_plan_id}).json()

    # -- the field ---------------------------------------------------------

    def test_rooms_are_open_unless_somebody_closes_them(self):
        """A library is a place people are allowed into. That is the default."""
        self.assertTrue(Room.objects.get(pk=self.public.pk).patron_access)

    def test_closing_a_room_does_not_hide_it(self):
        """The two settings are independent, which is the whole point."""
        self._close_the_store()
        self.store.refresh_from_db()
        self.assertFalse(self.store.patron_access)
        self.assertTrue(self.store.is_active)

    # -- routing -----------------------------------------------------------

    def test_no_waypoint_is_placed_in_a_closed_room(self):
        """A* only travels along waypoints, so this is what makes it unroutable."""
        from library.views import _point_in_polygon
        self._close_the_store()
        self.assertTrue(self._generate()['success'])
        inside = [w for w in Waypoint.objects.filter(floor_plan=self.plan)
                  if _point_in_polygon(w.map_x, w.map_y, self.store.geometry)]
        self.assertEqual(inside, [], 'a route was offered into a staff-only room')

    def test_the_open_room_still_gets_waypoints(self):
        """Closing one room must not quietly disable routing everywhere."""
        from library.views import _point_in_polygon
        self._close_the_store()
        self._generate()
        inside = [w for w in Waypoint.objects.filter(floor_plan=self.plan)
                  if _point_in_polygon(w.map_x, w.map_y, self.public.geometry)]
        self.assertTrue(inside, 'the public room lost its waypoints too')

    def test_closing_every_room_is_refused_rather_than_silently_empty(self):
        """Otherwise it reports success and quietly leaves a plan nobody can use."""
        self._close_the_store()
        self.public.patron_access = False
        self.public.save()
        r = self._generate()
        self.assertFalse(r['success'])
        self.assertIn('staff-only', r['error'])

    # -- the patron map ----------------------------------------------------

    def _map_rooms(self):
        r = self.client.get('/patron/map-data/').json()
        return {room['name']: room for room in r['rooms']}

    def test_a_closed_room_is_still_sent_to_the_patron_map(self):
        """Hiding it would leave a hole in the plan, which reads as a broken map."""
        self._close_the_store()
        rooms = self._map_rooms()
        self.assertIn('Store', rooms)
        self.assertFalse(rooms['Store']['patron_access'])

    def test_an_open_room_says_so(self):
        rooms = self._map_rooms()
        self.assertTrue(rooms['Reading Room']['patron_access'])

    def test_an_inactive_room_is_still_withheld(self):
        """The old flag keeps its old meaning. This adds one, it replaces none."""
        self.store.is_active = False
        self.store.save()
        self.assertNotIn('Store', self._map_rooms())

    # -- the map matching that keeps the marker out ------------------------

    def test_the_map_leaves_closed_rooms_out_of_the_walkable_set(self):
        """roomPolygons is what "somewhere a person could be standing" means."""
        html = open('templates/patron/patronmap.html', encoding='utf-8').read()
        start = html.index('function roomPolygons()')
        body = html[start:html.index('\n}', start)]
        self.assertIn('patron_access', body,
                      'the walkable set still includes staff-only rooms')

    def test_closed_rooms_also_block_like_furniture(self):
        """Belt and braces, for a store room drawn inside a bigger room."""
        html = open('templates/patron/patronmap.html', encoding='utf-8').read()
        start = html.index('function solids()')
        body = html[start:html.index('\n}', start)]
        self.assertIn('patron_access', body)

    # -- telling the Administrator ----------------------------------------

    def test_the_readiness_check_states_it_rather_than_warning(self):
        """It is a decision the library made, not a fault in the plan."""
        self._close_the_store()
        r = self._readiness()
        said = [i for i in r['issues'] if 'staff-only' in i['text']]
        self.assertEqual(len(said), 1, r['issues'])
        self.assertEqual(said[0]['level'], 'info')

    def test_an_info_line_is_not_counted_as_a_warning(self):
        """warnings used to be everything-that-is-not-a-blocker."""
        self._close_the_store()
        r = self._readiness()
        self.assertEqual(
            r['warnings'],
            sum(1 for i in r['issues'] if i['level'] == 'warning'))

    def test_it_does_not_complain_that_a_closed_room_has_no_waypoint(self):
        """That is the setting working, reported back as a fault."""
        self._close_the_store()
        self._generate()
        r = self._readiness()
        for i in r['issues']:
            self.assertNotIn('No waypoint inside "Store"', i['text'])

    # -- editing it --------------------------------------------------------

    def test_an_administrator_can_close_and_reopen_a_room(self):
        self.client.post('/admin-portal/edit-room/',
                         {'room_id': self.store.room_id, 'patron_access': '0'})
        self.store.refresh_from_db()
        self.assertFalse(self.store.patron_access)

        self.client.post('/admin-portal/edit-room/',
                         {'room_id': self.store.room_id, 'patron_access': '1'})
        self.store.refresh_from_db()
        self.assertTrue(self.store.patron_access)

    def test_the_edit_dialog_can_actually_be_submitted(self):
        """step="0.1" on the centroid fields made the whole form unsubmittable."""
        html = open('templates/admin/floorplanadmin.html', encoding='utf-8').read()
        block = html[html.index('id="editRoomForm"'):html.index('</form>', html.index('id="editRoomForm"'))]
        for field in ('map_x', 'map_y'):
            line = block[block.index('name="%s"' % field) - 120:block.index('name="%s"' % field)]
            self.assertIn('step="any"', line,
                          '%s cannot hold a centroid, so the form will not submit' % field)

    def test_a_post_that_does_not_mention_access_leaves_it_alone(self):
        """Reshaping a room posts room_id and geometry and nothing else."""
        self._close_the_store()
        self.client.post('/admin-portal/edit-room/',
                         {'room_id': self.store.room_id,
                          'geometry': json.dumps([[300, 100], [520, 100],
                                                  [520, 300], [300, 300]])})
        self.store.refresh_from_db()
        self.assertFalse(self.store.patron_access)

class PatronSchoolTests(TestCase):
    """School belongs to the patron, and a visit inherits it."""

    def setUp(self):
        self.user = _admin(modules='patrons,logs')
        self.client = _signed_in(self.user)

    def _patron(self, name='Ana Cruz', email='ana@example.invalid',
                school=None, patron_type='Student'):
        return Patron.objects.create(
            fullname=name, first_name=name.split()[0], last_name=name.split()[-1],
            email=email, patron_type=patron_type, school=school,
            account_status='Active', password_hash=hash_password('SmokeTest123'))

    # The field itself

    def test_a_patron_may_have_no_school(self):
        """A parent or a resident who walks in has none, and that is not an error."""
        p = self._patron(patron_type='Parent')
        p.full_clean(exclude=['password_hash'])
        self.assertIsNone(p.school)

    def _register(self, email, patron_type, **extra):
        """An online application. The uploaded ID is the whole identity check."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        data = {
            'action': 'register', 'first_name': 'Ana', 'last_name': 'Cruz',
            'email': email, 'password': 'SmokeTest123',
            'confirm_password': 'SmokeTest123', 'patron_type': patron_type,
            'contact_number': '09171234567', 'address': 'Sala, Cabuyao',
            'credential_document': SimpleUploadedFile(
                # A real PNG header.
                'id.png', bytes([137, 80, 78, 71, 13, 10, 26, 10]) + b'0' * 80,
                content_type='image/png'),
        }
        data.update(extra)
        return self.client.post('/patron/register/', data)

    def test_registering_records_the_school(self):
        self._register('ana.new@example.invalid', 'Student',
                       school='San Juan National High School')
        p = Patron.objects.filter(email='ana.new@example.invalid').first()
        self.assertIsNotNone(p, 'registration should have created the patron')
        self.assertEqual(p.school, 'San Juan National High School')

    def test_registering_without_one_leaves_it_empty_not_blank_string(self):
        """Null and empty are two states, and the pickers filter on null."""
        self._register('boy@example.invalid', 'Parent')
        p = Patron.objects.filter(email='boy@example.invalid').first()
        self.assertIsNotNone(p)
        self.assertIsNone(p.school)

    def test_an_administrator_can_change_it(self):
        p = self._patron(school='Old School')
        self.client.post('/admin-portal/edit-patron/%d/' % p.patron_id, {
            'first_name': 'Ana', 'last_name': 'Cruz', 'patron_type': 'Student',
            'email': 'ana@example.invalid', 'contact_number': '', 'address': '',
            'school': 'New School', 'account_status': 'Active'})
        p.refresh_from_db()
        self.assertEqual(p.school, 'New School')

    def test_clearing_it_stores_nothing_rather_than_an_empty_string(self):
        p = self._patron(school='Old School')
        self.client.post('/admin-portal/edit-patron/%d/' % p.patron_id, {
            'first_name': 'Ana', 'last_name': 'Cruz', 'patron_type': 'Parent',
            'email': 'ana@example.invalid', 'contact_number': '', 'address': '',
            'school': '', 'account_status': 'Active'})
        p.refresh_from_db()
        self.assertIsNone(p.school)

    # A visit inherits it

    def test_a_visit_takes_the_school_from_the_patron(self):
        """Nobody should retype it at every entry. That is what drifted."""
        p = self._patron(school='San Juan National High School')
        r = self.client.post('/admin-portal/log-entry/',
                             {'name': p.fullname, 'email': p.email}).json()
        self.assertTrue(r['success'], r)
        log = PatronLog.objects.filter(patron=p).first()
        self.assertEqual(log.school, 'San Juan National High School')

    def test_a_school_typed_for_one_visit_wins(self):
        """Somebody transfers, or is visiting from elsewhere that day."""
        p = self._patron(school='San Juan National High School')
        self.client.post('/admin-portal/log-entry/',
                         {'name': p.fullname, 'email': p.email,
                          'school': 'Rizal High School'})
        log = PatronLog.objects.filter(patron=p).first()
        self.assertEqual(log.school, 'Rizal High School')
        p.refresh_from_db()
        self.assertEqual(p.school, 'San Juan National High School',
                         'one visit must not rewrite the patron record')

    def test_a_patron_with_no_school_records_a_visit_with_none(self):
        p = self._patron(patron_type='Parent')
        self.client.post('/admin-portal/log-entry/',
                         {'name': p.fullname, 'email': p.email})
        log = PatronLog.objects.filter(patron=p).first()
        self.assertIsNone(log.school)

    # The picker

    def test_the_known_school_list_collapses_case(self):
        """"pnc" and "PNC" were two institutions in the library's own figures."""
        from library.views import known_schools
        a = self._patron('Ana Cruz', 'a@example.invalid', school='PNC')
        b = self._patron('Ben Cruz', 'b@example.invalid', school='pnc')
        PatronLog.objects.create(patron=a, school='PNC')
        PatronLog.objects.create(patron=b, school='pnc')
        PatronLog.objects.create(patron=b, school='PNC')
        names = known_schools()
        self.assertEqual(len([n for n in names if n.casefold() == 'pnc']), 1)

    def test_the_list_keeps_the_spelling_used_most(self):
        """The commonest is likeliest to be the one typed carefully."""
        from library.views import known_schools
        p = self._patron(school=None)
        for _ in range(3):
            PatronLog.objects.create(patron=p, school='Rizal High School')
        PatronLog.objects.create(patron=p, school='rizal high school')
        self.assertIn('Rizal High School', known_schools())

    def test_the_list_draws_on_visits_as_well_as_patrons(self):
        """Years of visits carry names no patron record has yet."""
        from library.views import known_schools
        p = self._patron(school=None)
        PatronLog.objects.create(patron=p, school='Only On A Visit')
        self.assertIn('Only On A Visit', known_schools())

    def test_the_forms_offer_the_list(self):
        self._patron(school='San Juan National High School')
        for url in ('/admin-portal/manage-patron/', '/admin-portal/log-management/'):
            html = self.client.get(url).content.decode()
            self.assertIn('knownSchools', html, url)
            self.assertIn('San Juan National High School', html, url)

    def test_patron_type_still_exists(self):
        """Reports group the sample by this field."""
        self.assertIn(('Student', 'Student'), Patron.PATRON_TYPE_CHOICES)
        self.assertEqual(len(Patron.PATRON_TYPE_CHOICES), 4)

class TracedShelfTests(TestCase):
    """A shelf drawn corner by corner keeps that shape through every gesture."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        self.room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        # An L, so a rectangle could never stand in for it.
        self.shape = [[100, 100], [200, 100], [200, 140], [140, 140], [140, 200], [100, 200]]
        self.shelf = Shelf.objects.create(
            room=self.room, name='Corner unit', map_x=150, map_y=150,
            width=46, depth=14, rotation=0, geometry=self.shape)

    def _move(self, x, y):
        return self.client.post('/admin-portal/move-shelf/', {
            'shelf_id': self.shelf.shelf_id, 'map_x': x, 'map_y': y}).json()

    def _rotate(self, deg):
        return self.client.post('/admin-portal/rotate-shelf/', {
            'shelf_id': self.shelf.shelf_id, 'rotation': deg}).json()

    def test_moving_carries_the_outline_with_it(self):
        """The bug: the anchor moved and the drawn shape stayed behind."""
        r = self._move(200, 250)
        self.assertTrue(r['success'], r)
        self.shelf.refresh_from_db()
        self.assertEqual(self.shelf.geometry,
                         [[x + 50, y + 100] for x, y in self.shape])

    def test_moving_does_not_change_the_shape(self):
        import math
        def edges(geom):
            return [round(math.dist(geom[i], geom[(i + 1) % len(geom)]), 6)
                    for i in range(len(geom))]
        before = edges(self.shape)
        self._move(-30, 900)
        self.shelf.refresh_from_db()
        self.assertEqual(edges(self.shelf.geometry), before)
        self.assertEqual(len(self.shelf.geometry), len(self.shape))

    def test_turning_rotates_the_outline_about_the_anchor(self):
        r = self._rotate(90)
        self.assertTrue(r['success'], r)
        self.shelf.refresh_from_db()
        # Expected position after a quarter turn.
        self.assertEqual(self.shelf.geometry[0], [200.0, 100.0])
        self.assertEqual(len(self.shelf.geometry), len(self.shape))

    def test_turning_twice_is_the_same_as_turning_once_by_the_sum(self):
        """Rotation is applied as a delta, so it must not accumulate wrongly."""
        self._rotate(45)
        self._rotate(90)
        once = list(self.shelf.geometry)
        self.shelf.refresh_from_db()
        turned = list(self.shelf.geometry)

        other = Shelf.objects.create(room=self.room, name='Twin', map_x=150,
                                     map_y=150, rotation=0, geometry=self.shape)
        self.client.post('/admin-portal/rotate-shelf/', {
            'shelf_id': other.shelf_id, 'rotation': 90})
        other.refresh_from_db()
        for a, b in zip(turned, other.geometry):
            self.assertAlmostEqual(a[0], b[0], places=1)
            self.assertAlmostEqual(a[1], b[1], places=1)
        del once

    def test_a_sized_shelf_is_untouched_by_any_of_this(self):
        plain = Shelf.objects.create(room=self.room, name='Plain', map_x=10,
                                     map_y=10, width=46, depth=14, rotation=0)
        r = self.client.post('/admin-portal/move-shelf/', {
            'shelf_id': plain.shelf_id, 'map_x': 99, 'map_y': 99}).json()
        self.assertTrue(r['success'], r)
        plain.refresh_from_db()
        self.assertIsNone(plain.geometry)
        self.assertEqual((plain.map_x, plain.map_y), (99, 99))

    def test_placing_an_unplaced_shelf_does_not_shift_its_outline(self):
        """A shelf with no position yet has nothing to be offset from."""
        fresh = Shelf.objects.create(room=self.room, name='Fresh', geometry=self.shape)
        self.assertIsNone(fresh.map_x)
        r = self.client.post('/admin-portal/place-shelf/', {
            'shelf_id': fresh.shelf_id, 'map_x': 500, 'map_y': 500}).json()
        self.assertTrue(r['success'], r)
        fresh.refresh_from_db()
        self.assertEqual(fresh.geometry, self.shape)

    def test_the_move_response_carries_the_new_outline(self):
        """The editor takes the server's answer so the two cannot drift."""
        r = self._move(160, 160)
        self.assertEqual(r['geometry'], [[x + 10, y + 10] for x, y in self.shape])


class ShelfGridTests(TestCase):
    """A bay's boards, made once instead of one dialog at a time."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        self.room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)

    def _add(self, **extra):
        params = {'room_id': self.room.room_id, 'name': 'Bay A',
                  'map_x': 10, 'map_y': 10}
        params.update(extra)
        return self.client.post('/admin-portal/add-shelf/', params).json()

    def _grid(self, shelf_id, **params):
        params['shelf_id'] = shelf_id
        return self.client.post('/admin-portal/set-shelf-grid/', params).json()

    def _labels(self, shelf_id):
        return [lv.label for lv in ShelfLevel.objects
                .filter(shelf_id=shelf_id).order_by('level_number', 'column_number')]

    def test_a_new_shelf_gets_the_levels_it_was_asked_for(self):
        r = self._add(levels=5, columns=1)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['levels_created'], 5)
        self.assertEqual(self._labels(r['shelf_id']),
                         ['Level 1', 'Level 2', 'Level 3', 'Level 4', 'Level 5'])

    def test_a_single_column_bay_is_not_labelled_column_one(self):
        """"Level 3, Column 1" is noise on a bookcase that has no bays."""
        r = self._add(levels=2, columns=1)
        for label in self._labels(r['shelf_id']):
            self.assertNotIn('Column', label)

    def test_a_divided_bay_gets_every_cell(self):
        r = self._add(levels=3, columns=3)
        self.assertEqual(r['levels_created'], 9)
        labels = self._labels(r['shelf_id'])
        self.assertIn('Level 1, Column 1', labels)
        self.assertIn('Level 3, Column 3', labels)

    def test_asking_for_no_levels_places_an_empty_bay(self):
        r = self._add(levels=0)
        self.assertTrue(r['success'], r)
        self.assertEqual(ShelfLevel.objects.filter(shelf_id=r['shelf_id']).count(), 0)

    def test_growing_a_shelf_adds_only_what_is_missing(self):
        shelf_id = self._add(levels=2, columns=1)['shelf_id']
        r = self._grid(shelf_id, levels=4, columns=1)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['added'], 2)
        self.assertEqual(len(self._labels(shelf_id)), 4)

    def test_running_the_same_grid_twice_changes_nothing(self):
        shelf_id = self._add(levels=3, columns=2)['shelf_id']
        before = self._labels(shelf_id)
        r = self._grid(shelf_id, levels=3, columns=2)
        self.assertEqual(r['added'], 0)
        self.assertEqual(self._labels(shelf_id), before)

    def test_shrinking_never_destroys_a_level_that_holds_books(self):
        """Reported, not deleted: a mistyped number must not lose books."""
        shelf_id = self._add(levels=4, columns=1)['shelf_id']
        top = ShelfLevel.objects.get(shelf_id=shelf_id, level_number=4)
        Book.objects.create(title='Kept', author='A', shelf_level=top)

        r = self._grid(shelf_id, levels=2, columns=1)
        self.assertTrue(r['success'], r)
        self.assertEqual(ShelfLevel.objects.filter(shelf_id=shelf_id).count(), 4)
        surplus = {row['label']: row['books'] for row in r['surplus']}
        self.assertIn('Level 4', surplus)
        self.assertEqual(surplus['Level 4'], 1)
        self.assertTrue(Book.objects.filter(title='Kept').exists())

    def test_absurd_counts_are_clamped_not_obeyed(self):
        r = self._add(levels=9999, columns=9999)
        self.assertTrue(r['success'], r)
        from library.views import MAX_SHELF_LEVELS, MAX_SHELF_COLUMNS
        self.assertEqual(ShelfLevel.objects.filter(shelf_id=r['shelf_id']).count(),
                         MAX_SHELF_LEVELS * MAX_SHELF_COLUMNS)

    def test_nonsense_counts_fall_back_rather_than_erroring(self):
        r = self._add(levels='four', columns='')
        self.assertTrue(r['success'], r)
        self.assertEqual(ShelfLevel.objects.filter(shelf_id=r['shelf_id']).count(), 1)


class LabelPickerGroupingTests(TestCase):
    """The picker groups by shelf and by level entirely client-side."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.bay = Shelf.objects.create(room=room, name='Bay A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=self.bay, level_number=2,
                                               category='Maths')
        Book.objects.create(title='On a shelf', author='A', shelf_level=self.level)
        Book.objects.create(title='Nowhere yet', author='B')

    def _rows(self):
        d = self.client.get('/admin-portal/book-label-picker/').json()
        self.assertTrue(d['success'], d)
        return {b['title']: b for b in d['books']}

    def test_a_shelved_book_carries_both_grouping_keys(self):
        row = self._rows()['On a shelf']
        self.assertEqual(row['shelf_id'], self.bay.shelf_id)
        self.assertEqual(row['shelf_name'], 'Bay A')
        self.assertEqual(row['level_id'], self.level.shelf_level_id)
        self.assertEqual(row['level_number'], 2)

    def test_an_unshelved_book_says_so_with_nulls(self):
        """Null rather than absent: the picker buckets on it."""
        row = self._rows()['Nowhere yet']
        self.assertIsNone(row['shelf_id'])
        self.assertIsNone(row['level_id'])

    def test_the_location_string_is_there_for_the_print_order(self):
        self.assertTrue(self._rows()['On a shelf']['location'])


class ReshapeByCornerTests(TestCase):
    """Dragging a corner is the one edit width and depth cannot express."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        self.room = Room.objects.create(floor_plan=self.plan, name='Main', map_x=0, map_y=0)
        self.L = [[0, 0], [100, 0], [100, 40], [40, 40], [40, 100], [0, 100]]

    def _reshape_shelf(self, shelf, geom):
        import json
        return self.client.post('/admin-portal/edit-shelf/', {
            'shelf_id': shelf.shelf_id, 'geometry': json.dumps(geom)}).json()

    def test_a_sized_shelf_becomes_traced_when_its_corners_are_moved(self):
        """That is what dragging a corner means; refusing it would be worse."""
        shelf = Shelf.objects.create(room=self.room, name='Plain', map_x=50,
                                     map_y=50, width=46, depth=14)
        self.assertIsNone(shelf.geometry)
        r = self._reshape_shelf(shelf, self.L)
        self.assertTrue(r['success'], r)
        shelf.refresh_from_db()
        self.assertEqual(shelf.geometry, self.L)

    def test_the_anchor_follows_the_new_outline(self):
        """Otherwise the label and the route target sit off the shape."""
        shelf = Shelf.objects.create(room=self.room, name='Plain', map_x=999,
                                     map_y=999, width=46, depth=14)
        self._reshape_shelf(shelf, self.L)
        shelf.refresh_from_db()
        self.assertLess(shelf.map_x, 100)
        self.assertLess(shelf.map_y, 100)

    def test_a_broken_outline_is_refused_and_changes_nothing(self):
        shelf = Shelf.objects.create(room=self.room, name='Plain', map_x=50,
                                     map_y=50, geometry=self.L)
        import json
        r = self.client.post('/admin-portal/edit-shelf/', {
            'shelf_id': shelf.shelf_id,
            'geometry': json.dumps([[0, 0], [1, 1]])}).json()
        self.assertFalse(r['success'], r)
        shelf.refresh_from_db()
        self.assertEqual(shelf.geometry, self.L)

    def test_furniture_can_be_reshaped_too(self):
        import json
        o = Obstacle.objects.create(floor_plan=self.plan, kind='Table',
                                    name='Long table', map_x=0, map_y=0,
                                    geometry=[[0, 0], [10, 0], [10, 10], [0, 10]])
        r = self.client.post('/admin-portal/edit-obstacle/', {
            'obstacle_id': o.obstacle_id, 'geometry': json.dumps(self.L)}).json()
        self.assertTrue(r['success'], r)
        o.refresh_from_db()
        self.assertEqual(o.geometry, self.L)

    def test_renaming_a_shelf_without_geometry_leaves_the_outline_alone(self):
        shelf = Shelf.objects.create(room=self.room, name='Corner', map_x=50,
                                     map_y=50, geometry=self.L)
        r = self.client.post('/admin-portal/edit-shelf/', {
            'shelf_id': shelf.shelf_id, 'name': 'Renamed'}).json()
        self.assertTrue(r['success'], r)
        shelf.refresh_from_db()
        self.assertEqual(shelf.name, 'Renamed')
        self.assertEqual(shelf.geometry, self.L)


class ShelfTopSurfaceTests(TestCase):
    """Books stacked on top of a bay, and fixing a board recorded wrongly."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Bay A', map_x=0, map_y=0)

    def _grid(self, **params):
        params.setdefault('shelf_id', self.shelf.shelf_id)
        params.setdefault('levels', 3)
        params.setdefault('columns', 1)
        return self.client.post('/admin-portal/set-shelf-grid/', params).json()

    def _tops(self):
        return ShelfLevel.objects.filter(shelf=self.shelf, is_top=True)

    def test_ticking_it_adds_a_top_surface(self):
        r = self._grid(has_top='1')
        self.assertTrue(r['success'], r)
        self.assertTrue(r['has_top'])
        self.assertEqual(self._tops().count(), 1)
        self.assertEqual(self._tops().first().label, 'Top')

    def test_the_top_is_not_one_of_the_numbered_boards(self):
        self._grid(levels=3, has_top='1')
        numbered = ShelfLevel.objects.filter(shelf=self.shelf, is_top=False, is_under=False)
        self.assertEqual(numbered.count(), 3)
        self.assertEqual(ShelfLevel.objects.filter(shelf=self.shelf).count(), 4)

    def test_ticking_it_twice_does_not_stack_two_tops(self):
        self._grid(has_top='1')
        self._grid(has_top='1')
        self.assertEqual(self._tops().count(), 1)

    def test_unticking_removes_an_empty_top(self):
        self._grid(has_top='1')
        r = self._grid(has_top='0')
        self.assertFalse(r['has_top'])
        self.assertEqual(self._tops().count(), 0)
        self.assertIn('removed', r['top_note'])

    def test_unticking_keeps_a_top_that_holds_books(self):
        """The rule everywhere here: a number in a box never deletes books."""
        self._grid(has_top='1')
        Book.objects.create(title='On top', author='A', shelf_level=self._tops().first())
        r = self._grid(has_top='0')
        self.assertEqual(self._tops().count(), 1)
        self.assertIn('kept', r['top_note'])
        self.assertTrue(Book.objects.filter(title='On top').exists())

    def test_shrinking_the_grid_never_reports_the_top_as_surplus(self):
        """It has no level number to be outside the range."""
        self._grid(levels=4, has_top='1')
        r = self._grid(levels=2)
        labels = [row['label'] for row in r['surplus']]
        self.assertNotIn('Top', labels)

    def test_a_board_can_be_turned_into_the_top_afterwards(self):
        self._grid(levels=2)
        board = ShelfLevel.objects.filter(shelf=self.shelf, level_number=2).first()
        Book.objects.create(title='Stays put', author='A', shelf_level=board)
        r = self.client.post('/admin-portal/edit-shelf-level/', {
            'shelf_level_id': board.shelf_level_id,
            'level_number': board.level_number,
            'is_top': '1', 'is_under': '0'}).json()
        self.assertTrue(r['success'], r)
        board.refresh_from_db()
        self.assertTrue(board.is_top)
        self.assertEqual(board.label, 'Top')
        # And its books came with it, rather than being deleted and remade.
        self.assertEqual(board.book_set.count(), 1)

    def test_a_board_cannot_be_both_top_and_underneath(self):
        self._grid(levels=1)
        board = ShelfLevel.objects.filter(shelf=self.shelf).first()
        r = self.client.post('/admin-portal/edit-shelf-level/', {
            'shelf_level_id': board.shelf_level_id,
            'level_number': 1, 'is_top': '1', 'is_under': '1'}).json()
        self.assertFalse(r['success'])
        board.refresh_from_db()
        self.assertFalse(board.is_top)

    def test_clearing_both_returns_it_to_a_numbered_board(self):
        self._grid(has_top='1')
        top = self._tops().first()
        r = self.client.post('/admin-portal/edit-shelf-level/', {
            'shelf_level_id': top.shelf_level_id,
            'level_number': 9, 'is_top': '0', 'is_under': '0'}).json()
        self.assertTrue(r['success'], r)
        top.refresh_from_db()
        self.assertFalse(top.is_top)
        self.assertEqual(top.label, 'Level 9')


class CallNumberTests(TestCase):
    """The spine number: class, author mark, title letter, year."""

    def _n(self, genre, author, title, year=None):
        return Book(genre=genre, author=author, title=title,
                    publication_year=year).derive_call_number()

    def test_it_reads_in_the_house_format(self):
        self.assertEqual(self._n('Fiction', 'Austen, Jane', 'Pride and Prejudice', 1963),
                         'FIC A31p 1963')

    def test_the_mark_comes_from_the_surname_not_the_forename(self):
        """Austen is A, whichever way round the name was typed."""
        a = self._n('Fiction', 'Austen, Jane', 'Quiet Sea', 1963)
        b = self._n('Fiction', 'Jane Austen', 'Quiet Sea', 1963)
        self.assertEqual(a, b)
        self.assertTrue(a.split()[1].startswith('A'), a)

    def test_a_leading_article_does_not_decide_the_letter(self):
        """"The Quiet Sea" files under q, the way a shelf is ordered."""
        self.assertTrue(self._n('Fiction', 'Austen, Jane', 'The Quiet Sea').endswith('q'))
        self.assertTrue(self._n('Fiction', 'Austen, Jane', 'A Quiet Sea').endswith('q'))
        self.assertTrue(self._n('Fiction', 'Austen, Jane', 'An Ocean').endswith('o'))

    def test_the_class_is_the_first_three_letters_of_the_genre(self):
        self.assertTrue(self._n('Science', 'Hawking, S', 'Time', 1988).startswith('SCI '))
        self.assertTrue(self._n('History', 'Rizal, J', 'Noli', 1887).startswith('HIS '))

    def test_a_book_with_nothing_to_go_on_still_gets_something(self):
        self.assertEqual(self._n('', '', 'Untitled'), 'GEN')

    def test_the_year_is_left_off_when_unknown(self):
        self.assertNotIn('None', self._n('Fiction', 'Austen, Jane', 'Emma'))

    def test_it_is_filled_in_on_save_and_never_overwritten(self):
        book = Book.objects.create(title='Emma', author='Austen, Jane',
                                   genre='Fiction', publication_year=1815)
        self.assertEqual(book.call_number, 'FIC A31e 1815')

        book.call_number = 'FIC AUS 1815'      # a librarian's own number
        book.save()
        book.refresh_from_db()
        self.assertEqual(book.call_number, 'FIC AUS 1815')

        # And correcting the genre later does not move a spine already labelled.
        book.genre = 'Classics'
        book.save()
        book.refresh_from_db()
        self.assertEqual(book.call_number, 'FIC AUS 1815')

    def test_clearing_it_renumbers(self):
        book = Book.objects.create(title='Emma', author='Austen, Jane',
                                   genre='Fiction', publication_year=1815)
        book.call_number = ''
        book.genre = 'Classics'
        book.save()
        book.refresh_from_db()
        self.assertEqual(book.call_number, 'CLA A31e 1815')


class BookConditionTests(TestCase):
    """Three steps, and anything else means Good rather than a failed row."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)

    def test_a_new_copy_is_good_until_said_otherwise(self):
        self.assertEqual(Book.objects.create(title='T', author='A').condition, 'Good')

    def test_the_three_values_are_the_whole_list(self):
        self.assertEqual([c[0] for c in Book.CONDITION_CHOICES],
                         ['Good', 'Worn', 'Damaged'])

    def test_editing_accepts_a_valid_value_and_ignores_a_bad_one(self):
        book = Book.objects.create(title='T', author='A')
        base = {'title': 'T', 'author': 'A'}
        self.client.post('/admin-portal/edit-book/%d/' % book.book_id,
                         dict(base, condition='Damaged'),
                         HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        book.refresh_from_db()
        self.assertEqual(book.condition, 'Damaged')

        self.client.post('/admin-portal/edit-book/%d/' % book.book_id,
                         dict(base, condition='Pristine'),
                         HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        book.refresh_from_db()
        self.assertEqual(book.condition, 'Damaged', 'a bad value must not stick')


class CatalogueImportTests(TestCase):
    """The template, the importer and the labels agree on the columns."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)

    def _template(self):
        import io as _io
        import openpyxl
        r = self.client.get('/admin-portal/download-book-template/')
        self.assertEqual(r.status_code, 200)
        return list(openpyxl.load_workbook(_io.BytesIO(r.content)).active.iter_rows(values_only=True))

    def test_the_template_offers_both_new_columns(self):
        headers = self._template()[0]
        self.assertIn('Condition', headers)
        self.assertIn('Code Label', headers)

    def _import(self, rows, **extra):
        import io as _io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        for row in rows:
            wb.active.append(row)
        buf = _io.BytesIO()
        wb.save(buf)
        up = SimpleUploadedFile('b.xlsx', buf.getvalue(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        data = {'excel_file': up}
        data.update(extra)          # preview=1, copy_rows=... when a test needs them
        return self.client.post('/admin-portal/import-books/', data).json()

    def test_the_template_can_be_filled_in_and_imported_straight_back(self):
        rows = self._template()
        r = self._import(rows)
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.filter(condition='Worn').count(), 1)
        self.assertTrue(Book.objects.filter(call_number='FIC A31p 1963').exists())

    def test_a_blank_code_label_is_worked_out(self):
        r = self._import([['Title', 'Author', 'Genre', 'Publication Year'],
                          ['Emma', 'Austen, Jane', 'Fiction', 1815]])
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.get(title='Emma').call_number, 'FIC A31e 1815')

    def test_an_unknown_condition_does_not_fail_the_row(self):
        r = self._import([['Title', 'Author', 'Condition'],
                          ['Emma', 'Austen, Jane', 'absolutely fine']])
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.get(title='Emma').condition, 'Good')

    def test_a_sheet_with_no_such_columns_still_imports(self):
        """Every sheet written before these columns existed must still work."""
        r = self._import([['Title', 'Author'], ['Emma', 'Austen, Jane']])
        self.assertTrue(r['success'], r)
        book = Book.objects.get(title='Emma')
        self.assertEqual(book.condition, 'Good')
        self.assertTrue(book.call_number)

    def test_the_label_sheet_prints_the_code_label(self):
        import io as _io
        from pypdf import PdfReader
        book = Book.objects.create(title='Emma', author='Austen, Jane',
                                   genre='Fiction', publication_year=1815)
        r = self.client.get('/admin-portal/book-qr-labels/', {'ids': str(book.book_id)})
        self.assertEqual(r['Content-Type'], 'application/pdf')
        text = PdfReader(_io.BytesIO(r.content)).pages[0].extract_text()
        self.assertIn('FIC A31e 1815', text)
        self.assertIn('Emma', text)


class BookLocationLabelTests(TestCase):
    """Shelf, then across, then up -- the order somebody walks it."""

    def setUp(self):
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)

    def _book(self, **level_kw):
        level = ShelfLevel.objects.create(shelf=self.shelf, **level_kw)
        return Book.objects.create(title='T', author='A', shelf_level=level)

    def test_the_long_form_reads_as_a_sentence(self):
        b = self._book(level_number=2, column_number=1)
        self.assertEqual(b.location_label(), 'Shelf A Column 1 Level 2')

    def test_the_short_form_drops_the_word_shelf(self):
        b = self._book(level_number=2, column_number=1)
        self.assertEqual(b.location_label(short=True), 'A C1 L2')

    def test_a_bay_with_no_columns_says_nothing_about_them(self):
        b = self._book(level_number=3)
        self.assertEqual(b.location_label(), 'Shelf A Level 3')
        self.assertEqual(b.location_label(short=True), 'A L3')

    def test_a_top_is_named_not_numbered(self):
        """"L5" sends somebody to the fifth board of a case whose books are on its lid."""
        b = self._book(level_number=9, is_top=True)
        self.assertEqual(b.location_label(), 'Shelf A Top')
        self.assertNotIn('L9', b.location_label(short=True))

    def test_an_unshelved_copy_has_no_location_at_all(self):
        b = Book.objects.create(title='T', author='A')
        self.assertEqual(b.location_label(), '')
        self.assertEqual(b.location_label(short=True), '')


class LocationParsingTests(TestCase):
    """A sheet carries whichever form the person typing it preferred."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        for lv in (1, 2, 3):
            for col in (1, 2):
                ShelfLevel.objects.create(shelf=self.shelf, level_number=lv,
                                          column_number=col)

    def test_both_forms_read_the_same(self):
        from library.views import parse_location
        self.assertEqual(parse_location('Shelf A Column 1 Level 2'),
                         ('Shelf A', 2, 1, False, False))
        self.assertEqual(parse_location('A C1 L2'), ('A', 2, 1, False, False))

    def test_separators_do_not_matter(self):
        from library.views import parse_location
        for text in ('Shelf A · Column 1 · Level 2',
                     'Shelf A, Column 1, Level 2',
                     'Shelf A/Column 1/Level 2'):
            self.assertEqual(parse_location(text), ('Shelf A', 2, 1, False, False), text)

    def test_a_plain_storage_area_still_means_nothing_positional(self):
        """Old sheets said "Zone 3" and must keep working."""
        from library.views import parse_location
        self.assertEqual(parse_location('Zone 3'), ('Zone 3', None, None, False, False))

    def _import(self, rows, **extra):
        import io as _io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        for row in rows:
            wb.active.append(row)
        buf = _io.BytesIO()
        wb.save(buf)
        up = SimpleUploadedFile('b.xlsx', buf.getvalue(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        data = {'excel_file': up}
        data.update(extra)          # preview=1, copy_rows=... when a test needs them
        return self.client.post('/admin-portal/import-books/', data).json()

    def test_importing_the_long_form_lands_on_the_right_board(self):
        r = self._import([['Title', 'Author', 'Location'],
                          ['Emma', 'Austen, Jane', 'Shelf A Column 1 Level 2']])
        self.assertTrue(r['success'], r)
        book = Book.objects.get(title='Emma')
        self.assertEqual(book.shelf_level.level_number, 2)
        self.assertEqual(book.shelf_level.column_number, 1)
        self.assertEqual(book.location_label(), 'Shelf A Column 1 Level 2')

    def test_importing_the_short_form_lands_on_the_same_board(self):
        r = self._import([['Title', 'Author', 'Location'],
                          ['Emma', 'Austen, Jane', 'A C1 L2']])
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.get(title='Emma').location_label(),
                         'Shelf A Column 1 Level 2')

    def test_a_board_that_does_not_exist_yet_is_made(self):
        """The location travels with the import rather than being dropped."""
        r = self._import([['Title', 'Author', 'Location'],
                          ['Emma', 'Austen, Jane', 'Shelf A Column 2 Level 9']])
        self.assertTrue(r['success'], r)
        book = Book.objects.get(title='Emma')
        self.assertEqual(book.shelf_level.level_number, 9)
        self.assertEqual(book.shelf_level.column_number, 2)

    def test_a_top_imports_as_a_top(self):
        r = self._import([['Title', 'Author', 'Location'],
                          ['Emma', 'Austen, Jane', 'Shelf A Top']])
        self.assertTrue(r['success'], r)
        self.assertTrue(Book.objects.get(title='Emma').shelf_level.is_top)

    def test_the_template_round_trips_with_its_own_locations(self):
        import io as _io
        import openpyxl
        tpl = self.client.get('/admin-portal/download-book-template/')
        rows = list(openpyxl.load_workbook(_io.BytesIO(tpl.content)).active.iter_rows(values_only=True))
        self.assertIn('Location', rows[0])
        r = self._import(rows)
        self.assertTrue(r['success'], r)
        # Both example rows name the same board, one long and one short.
        placed = [b.location_label() for b in Book.objects.all() if b.shelf_level]
        self.assertTrue(placed, 'nothing was shelved by the import')
        for where in placed:
            self.assertEqual(where, 'Shelf A Column 1 Level 2')


class ImportRepeatAndConditionTests(TestCase):
    """The two things a real sheet exposed: repeats, and unread conditions."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf C', map_x=0, map_y=0)
        ShelfLevel.objects.create(shelf=shelf, level_number=3, column_number=1)

    def _import(self, rows, **extra):
        import io as _io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        for row in rows:
            wb.active.append(row)
        buf = _io.BytesIO()
        wb.save(buf)
        up = SimpleUploadedFile('b.xlsx', buf.getvalue(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        data = {'excel_file': up}
        data.update(extra)          # preview=1, copy_rows=... when a test needs them
        return self.client.post('/admin-portal/import-books/', data).json()

    # Shaped like the real sheet: a blank ISBN column and "GOOD/BAD CONDITION".
    SHEET = [
        ['Title', 'Author', 'Publication Year', 'ISBN', 'Genre', 'Condition',
         'Code Label', 'Quantity', 'Location'],
        ['Communication for the Common Good', 'Florangel Rosario-Braid', 1990, '',
         'EDUCATION', 'BAD CONDITION', 'SR 001.54 C65 1990', 1,
         'Shelf C Column 1 Level 3'],
        ['A Course on Words', 'Waldo E. Sweet', 1989, '0-472-08101-2',
         'EDUCATION', 'GOOD CONDITION', 'P 305 S93 1982', 1,
         'Shelf C Column 1 Level 3'],
    ]

    def test_the_sheets_own_words_are_understood(self):
        from library.views import parse_condition
        self.assertEqual(parse_condition('GOOD CONDITION'), 'Good')
        self.assertEqual(parse_condition('BAD CONDITION'), 'Damaged')
        self.assertEqual(parse_condition('Fair'), 'Worn')
        self.assertIsNone(parse_condition('sparkly'))

    def test_a_bad_condition_row_is_not_filed_as_good(self):
        r = self._import(self.SHEET)
        self.assertTrue(r['success'], r)
        book = Book.objects.get(title='Communication for the Common Good')
        self.assertEqual(book.condition, 'Damaged')

    def test_importing_the_same_sheet_twice_adds_nothing(self):
        """The bug: 380 of 480 books had no ISBN, so every run copied them."""
        first = self._import(self.SHEET)
        self.assertEqual(first['imported'], 2, first)
        after_first = Book.objects.count()

        second = self._import(self.SHEET)
        self.assertTrue(second['success'], second)
        self.assertEqual(Book.objects.count(), after_first,
                         'a second run of the same sheet added books again')
        self.assertIn('already in the catalogue', second['message'])

    # Duplicate rows need a decision.

    def test_a_repeat_is_listed_by_name_not_just_counted(self):
        """A count says a decision is waiting; only the title says which way."""
        self._import(self.SHEET)
        preview = self._import(self.SHEET, preview='1')
        self.assertTrue(preview['success'], preview)
        titles = sorted(c['title'] for c in preview['clashes'])
        self.assertEqual(titles, ['A Course on Words',
                                  'Communication for the Common Good'])
        for c in preview['clashes']:
            self.assertEqual(c['have'], 1)
            self.assertEqual(c['where'], 'Shelf C')
            self.assertGreaterEqual(c['row'], 2)

    def test_a_repeat_is_rejected_unless_it_is_asked_for(self):
        """Skipping stays the default: the commoner accident is a double import."""
        self._import(self.SHEET)
        before = Book.objects.count()
        again = self._import(self.SHEET)
        self.assertEqual(Book.objects.count(), before)
        self.assertEqual(again['copies_of_existing'], 0)

    def test_a_ticked_row_is_imported_as_a_second_copy(self):
        self._import(self.SHEET)
        preview = self._import(self.SHEET, preview='1')
        wanted = [c for c in preview['clashes']
                  if c['title'] == 'A Course on Words']
        self.assertEqual(len(wanted), 1)

        done = self._import(self.SHEET, copy_rows=str(wanted[0]['row']))
        self.assertTrue(done['success'], done)
        self.assertEqual(done['copies_of_existing'], 1)
        self.assertEqual(Book.objects.filter(title='A Course on Words').count(), 2)
        # The row that was not ticked stayed out.
        self.assertEqual(
            Book.objects.filter(title='Communication for the Common Good').count(), 1)

    def test_a_second_copy_does_not_carry_the_first_ones_isbn(self):
        """The ISBN names the edition; the record already on file holds it."""
        self._import(self.SHEET)
        preview = self._import(self.SHEET, preview='1')
        rows = ','.join(str(c['row']) for c in preview['clashes'])
        self._import(self.SHEET, copy_rows=rows)

        copies = Book.objects.filter(title='A Course on Words').order_by('book_id')
        self.assertEqual(copies.count(), 2)
        self.assertEqual(copies[0].ISBN, '0-472-08101-2')
        self.assertIsNone(copies[1].ISBN)

    def test_a_copy_bound_for_another_shelf_is_still_a_repeat(self):
        """Re-importing a book for another shelf is flagged as a duplicate."""
        shelf = Shelf.objects.get(name='Shelf C')
        ShelfLevel.objects.create(shelf=shelf, level_number=4, column_number=1)
        self._import(self.SHEET)

        elsewhere = [self.SHEET[0],
                     ['A Course on Words', 'Waldo E. Sweet', 1989, '',
                      'EDUCATION', 'GOOD CONDITION', '', 1,
                      'Shelf C Column 1 Level 4']]
        preview = self._import(elsewhere, preview='1')
        self.assertEqual(len(preview['clashes']), 1, preview)
        self.assertEqual(preview['clashes'][0]['title'], 'A Course on Words')

    def test_a_preview_writes_nothing_whichever_way_the_rows_are_ticked(self):
        self._import(self.SHEET)
        before = Book.objects.count()
        preview = self._import(
            self.SHEET, preview='1',
            copy_rows=','.join(str(c['row'])
                               for c in self._import(self.SHEET, preview='1')['clashes']))
        self.assertTrue(preview['success'], preview)
        self.assertEqual(Book.objects.count(), before)

    def test_the_confirmed_import_does_not_ask_the_same_question_twice(self):
        """Clashes are a preview's business. After the confirm they are decided."""
        self._import(self.SHEET)
        done = self._import(self.SHEET)
        self.assertEqual(done['clashes'], [])

    def test_a_row_number_that_names_nothing_is_ignored(self):
        """The sheet is re-uploaded, so a stale or invented number is possible."""
        self._import(self.SHEET)
        before = Book.objects.count()
        done = self._import(self.SHEET, copy_rows='999,,abc,-4')
        self.assertTrue(done['success'], done)
        self.assertEqual(Book.objects.count(), before)

    def test_quantity_still_makes_copies_of_a_book_that_is_new(self):
        """Regression: the snapshot must not let a row block its own copies."""
        rows = [self.SHEET[0],
                ['Brand New', 'Nobody', 1990, '', 'EDUCATION', 'GOOD CONDITION',
                 '', 3, 'Shelf C Column 1 Level 3']]
        r = self._import(rows)
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.filter(title='Brand New').count(), 3)

    def test_the_listing_stops_and_says_so_when_a_sheet_is_all_repeats(self):
        """Past the cap the sheet is not a delivery, it is the same file again."""
        from unittest.mock import patch
        rows = [self.SHEET[0]] + [
            ['Repeat %d' % i, 'Someone', 1990, '', 'EDUCATION', 'GOOD CONDITION',
             '', 1, 'Shelf C Column 1 Level 3'] for i in range(5)]
        self._import(rows)

        with patch('library.views.MAX_CLASH_ROWS', 2):
            preview = self._import(rows, preview='1')
        self.assertEqual(len(preview['clashes']), 2, preview)
        self.assertTrue(preview['clashes_truncated'])

        # Under a cap the sheet fits inside, nothing is claimed to be missing.
        preview = self._import(rows, preview='1')
        self.assertEqual(len(preview['clashes']), 5)
        self.assertFalse(preview['clashes_truncated'])

    def test_two_rows_for_one_book_in_a_sheet_are_two_copies(self):
        """A sheet listing the same book twice is how a library says it holds two."""
        rows = [self.SHEET[0],
                ['Twice Over', 'Someone', 1990, '', 'EDUCATION', 'GOOD CONDITION',
                 '', 1, 'Shelf C Column 1 Level 3'],
                ['Twice Over', 'Someone', 1990, '', 'EDUCATION', 'GOOD CONDITION',
                 '', 1, 'Shelf C Column 1 Level 3']]
        r = self._import(rows)
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.filter(title='Twice Over').count(), 2)

        # And running that same sheet again still adds nothing.
        self._import(rows)
        self.assertEqual(Book.objects.filter(title='Twice Over').count(), 2)

    def test_a_quantity_still_creates_that_many_copies(self):
        """The dedupe must not mistake a book's own copies for repeats."""
        rows = [self.SHEET[0],
                ['Three Copies', 'Someone', 1990, '', 'EDUCATION', 'GOOD CONDITION',
                 '', 3, 'Shelf C Column 1 Level 3']]
        r = self._import(rows)
        self.assertTrue(r['success'], r)
        self.assertEqual(Book.objects.filter(title='Three Copies').count(), 3)

    def test_the_same_book_on_a_different_shelf_is_still_asked_about(self):
        """Re-importing a book for another shelf needs confirmation."""
        rows = [self.SHEET[0],
                ['Communication for the Common Good', 'Florangel Rosario-Braid', 1990,
                 '', 'EDUCATION', 'BAD CONDITION', '', 1, 'Shelf C Column 1 Level 9']]
        self._import(self.SHEET)
        r = self._import(rows)
        self.assertTrue(r['success'], r)
        self.assertEqual(
            Book.objects.filter(title='Communication for the Common Good').count(), 1)

        preview = self._import(rows, preview='1')
        self.assertEqual(len(preview['clashes']), 1, preview)
        asked = self._import(rows, copy_rows=str(preview['clashes'][0]['row']))
        self.assertTrue(asked['success'], asked)
        self.assertEqual(
            Book.objects.filter(title='Communication for the Common Good').count(), 2)

    def test_an_unreadable_condition_is_named_in_the_report(self):
        rows = [self.SHEET[0],
                ['Odd One', 'Someone', 1990, '', 'EDUCATION', 'sparkly', '', 1, '']]
        r = self._import(rows)
        self.assertIn('could not read the condition', r['message'])
        self.assertIn('sparkly', r['message'])
        self.assertEqual(Book.objects.get(title='Odd One').condition, 'Good')


class ImportPreviewTests(TestCase):
    """A preview reports what happened, then puts it all back."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf C', map_x=0, map_y=0)
        ShelfLevel.objects.create(shelf=shelf, level_number=3, column_number=1)

    SHEET = [
        ['Title', 'Author', 'Publication Year', 'ISBN', 'Genre', 'Condition',
         'Code Label', 'Quantity', 'Location'],
        ['One', 'Someone', 1990, '', 'EDUCATION', 'BAD CONDITION', '', 2,
         'Shelf C Column 1 Level 3'],
        ['Two', 'Another', 1991, '111-1', 'EDUCATION', 'GOOD CONDITION', '', 1,
         'Shelf C Column 1 Level 3'],
    ]

    def _post(self, rows, **extra):
        import io as _io
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        for row in rows:
            wb.active.append(row)
        buf = _io.BytesIO()
        wb.save(buf)
        up = SimpleUploadedFile('b.xlsx', buf.getvalue(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        data = {'excel_file': up}
        data.update(extra)
        return self.client.post('/admin-portal/import-books/', data).json()

    def test_a_preview_writes_nothing(self):
        before = Book.objects.count()
        r = self._post(self.SHEET, preview='1')
        self.assertTrue(r['success'], r)
        self.assertTrue(r['preview'])
        self.assertEqual(Book.objects.count(), before,
                         'the preview left books behind')

    def test_a_preview_counts_what_the_real_run_would_do(self):
        preview = self._post(self.SHEET, preview='1')
        real = self._post(self.SHEET)
        for key in ('imported', 'copies_created', 'skipped_dup',
                    'skipped_repeat', 'skipped_blank'):
            self.assertEqual(preview[key], real[key], key)
        self.assertEqual(Book.objects.count(), real['imported'])

    def test_a_preview_says_what_will_happen_not_what_did(self):
        preview = self._post(self.SHEET, preview='1')
        self.assertIn('will be imported', preview['message'])
        self.assertNotIn('Imported', preview['message'])
        self.assertIn('Imported', self._post(self.SHEET)['message'])

    def test_a_preview_does_not_leave_shelf_levels_behind_either(self):
        """It creates levels for locations it cannot find; those roll back too."""
        before = ShelfLevel.objects.count()
        self._post([self.SHEET[0],
                    ['Three', 'Someone', 1990, '', 'EDUCATION', '', '', 1,
                     'Shelf C Column 1 Level 8']], preview='1')
        self.assertEqual(ShelfLevel.objects.count(), before)

    def test_the_summary_breaks_the_numbers_out_for_the_confirm_box(self):
        r = self._post(self.SHEET, preview='1')
        for key in ('imported', 'copies_created', 'skipped_dup',
                    'skipped_repeat', 'skipped_blank'):
            self.assertIn(key, r)

    def test_a_preview_hands_out_no_book_ids(self):
        """They would be ids of rows that no longer exist."""
        r = self._post(self.SHEET, preview='1')
        self.assertEqual(r['created_ids'], [])
        self.assertTrue(self._post(self.SHEET)['created_ids'])

    def test_importing_without_asking_for_a_preview_still_writes(self):
        r = self._post(self.SHEET)
        self.assertFalse(r.get('preview'))
        self.assertEqual(Book.objects.count(), 3)


class CopiesColumnTests(TestCase):
    """The Copies column counts copies, and can be searched on."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        for i in range(4):
            Book.objects.create(title='Four Of These', author='Ann Author',
                                genre='FICTION', status='Available')
        Book.objects.create(title='Only One', author='Bob Byline',
                            genre='FICTION', status='Available')
        for i in range(2):
            Book.objects.create(title='Two Of These', author='Cee Creator',
                                genre='FICTION', status='Available')
        # Same title, different author: a different book, not another copy.
        Book.objects.create(title='Only One', author='Dee Different',
                            genre='FICTION', status='Available')

    def _rows(self, query=''):
        r = self.client.get('/admin-portal/management/' + query)
        self.assertEqual(r.status_code, 200)
        return r.context['books'].paginator.count

    def test_the_column_counts_the_copies_rather_than_always_saying_one(self):
        r = self.client.get('/admin-portal/management/?q=Four Of These')
        counts = {b.copies for b in r.context['books']}
        self.assertEqual(counts, {4})

    def test_a_different_author_is_a_different_book(self):
        r = self.client.get('/admin-portal/management/?q=Only One')
        self.assertEqual({b.copies for b in r.context['books']}, {1})

    def test_searching_for_single_copies(self):
        self.assertEqual(self._rows('?copies=1'), 2)

    def test_searching_for_an_exact_number(self):
        self.assertEqual(self._rows('?copies=4'), 4)
        self.assertEqual(self._rows('?copies=2'), 2)

    def test_searching_for_that_many_or_more(self):
        self.assertEqual(self._rows('?copies=2%2B'), 6)
        self.assertEqual(self._rows('?copies=4%2B'), 4)

    def test_nonsense_is_ignored_rather_than_emptying_the_table(self):
        for bad in ('?copies=banana', '?copies=0', '?copies=-3', '?copies=%2B',
                    '?copies=999999'):
            self.assertEqual(self._rows(bad), 8, bad)

    def test_the_filter_survives_paging(self):
        r = self.client.get('/admin-portal/management/?copies=2%2B')
        self.assertIn('copies=2%2B', r.context['querystring'])

    def test_it_combines_with_the_other_filters(self):
        Book.objects.create(title='Four Of These', author='Ann Author',
                            genre='FICTION', status='Borrowed')
        # Five rows share the title now; only the four available ones match.
        self.assertEqual(self._rows('?copies=5&status=Available'), 4)

    def test_the_staff_page_can_search_on_copies_too(self):
        staff = User.objects.create(
            fullname='Smoke Staff', email='smoke-staff@example.invalid',
            password_hash=hash_password('SmokeTest123'),
            role='Staff', account_status='Active', modules='books')
        r = _signed_in(staff).get('/library-staff/books/?copies=4')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context['books'].paginator.count, 4)


class ParseCopiesFilterTests(TestCase):
    """What the dropdown sends, and what someone types into the address bar."""

    def setUp(self):
        from library.views import COPIES_FILTER_CHOICES, parse_copies_filter
        self.parse = parse_copies_filter
        self.choices = COPIES_FILTER_CHOICES

    def test_a_bare_number_means_exactly_that_many(self):
        self.assertEqual(self.parse('3'), {'copies': 3})

    def test_a_trailing_plus_means_that_many_or_more(self):
        self.assertEqual(self.parse('3+'), {'copies__gte': 3})

    def test_whitespace_is_forgiven(self):
        self.assertEqual(self.parse('  2+  '), {'copies__gte': 2})

    def test_nothing_recognisable_means_no_filter(self):
        for raw in ('', None, 'banana', '0', '-1', '+', '1.5', '100', '3++'):
            self.assertIsNone(self.parse(raw), raw)

    def test_every_dropdown_option_parses(self):
        for value, label in self.choices:
            self.assertIsNotNone(self.parse(value), value)


class SlowRequestLoggingTests(TestCase):
    """A slow page leaves a line behind; a normal one costs nothing."""

    def _run(self, seconds):
        import logging
        from library.middleware import SlowRequestLoggingMiddleware

        class FakeResponse:
            status_code = 200

        def get_response(request):
            # Fake the clock instead of sleeping.
            return FakeResponse()

        mw = SlowRequestLoggingMiddleware(get_response)
        ticks = iter([0.0, seconds])
        with mock.patch('library.middleware.time.monotonic',
                        side_effect=lambda: next(ticks)):
            with self.assertLogs('library.middleware', level='WARNING') as caught:
                logging.getLogger('library.middleware').warning('probe')
                mw(_FakeRequest())
        return caught.output

    def test_a_slow_request_is_logged_with_its_path_and_duration(self):
        lines = self._run(3.0)
        slow = [l for l in lines if 'Slow request' in l]
        self.assertEqual(len(slow), 1, lines)
        self.assertIn('/probe/', slow[0])
        self.assertIn('3.00s', slow[0])

    def test_a_normal_request_logs_nothing(self):
        lines = self._run(0.2)
        self.assertEqual([l for l in lines if 'Slow request' in l], [])

    def test_the_threshold_is_inclusive(self):
        from library.middleware import SLOW_REQUEST_SECONDS
        lines = self._run(SLOW_REQUEST_SECONDS)
        self.assertTrue(any('Slow request' in l for l in lines))

    def test_it_is_registered_first_so_it_times_the_whole_stack(self):
        from django.conf import settings
        self.assertEqual(settings.MIDDLEWARE[0],
                         'library.middleware.SlowRequestLoggingMiddleware')


class ImportFailureTests(TestCase):
    """An unexpected fault is logged in full and reported without its guts."""

    def test_the_browser_is_given_a_reference_not_the_exception(self):
        from library.views import import_failed
        secret = 'relation "library_book" does not exist at /srv/ayla/db.sock'
        with self.assertLogs('library.views', level='ERROR'):
            response = import_failed('Book import', RuntimeError(secret))
        body = json.loads(response.content)
        self.assertFalse(body['success'])
        self.assertNotIn(secret, body['error'])
        self.assertNotIn('library_book', body['error'])
        self.assertNotIn('/srv/', body['error'])

    def test_the_reference_in_the_message_is_the_one_in_the_log(self):
        from library.views import import_failed
        with self.assertLogs('library.views', level='ERROR') as caught:
            response = import_failed('Patron import', ValueError('boom'))
        reference = json.loads(response.content)['error'].split()[-1].rstrip('.')
        self.assertEqual(len(reference), 8)
        self.assertTrue(any(reference in line for line in caught.output),
                        'the reference shown to the user is not in the log')

    def test_the_log_keeps_what_the_browser_does_not(self):
        from library.views import import_failed
        with self.assertLogs('library.views', level='ERROR') as caught:
            import_failed('Log import', RuntimeError('the real cause'))
        joined = '\n'.join(caught.output)
        self.assertIn('the real cause', joined)
        self.assertIn('Log import', joined)

    def test_every_import_view_routes_faults_through_it(self):
        """No import endpoint may go back to returning str(e)."""
        import inspect
        from library import views
        source = inspect.getsource(views)
        self.assertNotIn("'error': str(e)", source,
                         'an import view is leaking exception text again')


class AuditFailuresAreLoggedTests(TestCase):
    """Auditing still never breaks the operation -- but no longer vanishes."""

    def test_a_failing_audit_write_is_recorded_somewhere(self):
        from library import audit
        with mock.patch.object(audit.SystemLog.objects, 'create',
                               side_effect=RuntimeError('table is gone')):
            with self.assertLogs('library.audit', level='ERROR') as caught:
                audit.log_admin_action(_FakeRequest(), 'Something', 'detail')
        self.assertIn('audit record', '\n'.join(caught.output))

    def test_the_operation_still_survives_a_failing_audit(self):
        from library import audit
        with mock.patch.object(audit.SystemLog.objects, 'create',
                               side_effect=RuntimeError('table is gone')):
            with self.assertLogs('library.audit', level='ERROR'):
                audit.log_admin_action(_FakeRequest(), 'Something', 'detail')
        # Reaching here at all is the assertion: no exception escaped.


class ExportLibraryTests(TestCase):
    """The fixture that carries the library to a new host."""

    def setUp(self):
        import tempfile
        self.dir = tempfile.mkdtemp()
        self.path = os.path.join(self.dir, 'export.json')
        self.book = Book.objects.create(
            # Title with an en dash.
            title='The Science Library (Volumes 1\u20136)',
            author='Unknown', genre='REFERENCE', status='Available')

    def _export(self, *extra):
        from django.core.management import call_command
        out = StringIO()
        call_command('export_library', '-o', self.path, *extra, stdout=out)
        return out.getvalue()

    def _objects(self):
        with io.open(self.path, encoding='utf-8') as fh:
            return json.load(fh)

    def test_the_file_is_utf8_and_keeps_the_en_dash(self):
        self._export()
        titles = [o['fields']['title'] for o in self._objects()
                  if o['model'] == 'library.book']
        self.assertIn('The Science Library (Volumes 1\u20136)', titles)

    def test_it_is_decodable_as_utf8_end_to_end(self):
        """The failure this guards against is a file the server cannot read."""
        self._export()
        raw = io.open(self.path, 'rb').read()
        raw.decode('utf-8')          # raises if the encoding regressed

    def test_transient_security_state_never_travels(self):
        LoginAttempt.objects.create(scope='admin', identifier='someone@example.com')
        self._export()
        models = {o['model'] for o in self._objects()}
        self.assertNotIn('library.loginattempt', models)
        self.assertNotIn('library.passwordresetotp', models)

    def test_the_activity_log_is_left_out_unless_asked_for(self):
        SystemLog.objects.create(actor_role='Admin', admin_name='Someone',
                                 action='Did a thing', entity_type='Book')
        self._export()
        self.assertNotIn('library.systemlog',
                         {o['model'] for o in self._objects()})
        self._export('--with-logs')
        self.assertIn('library.systemlog',
                      {o['model'] for o in self._objects()})

    def test_book_ids_travel_with_the_rows(self):
        """Printed QR labels encode book_id; a book must not be renumbered."""
        self._export()
        exported = {o['pk'] for o in self._objects() if o['model'] == 'library.book'}
        self.assertIn(self.book.book_id, exported)

    def test_it_reports_what_it_wrote(self):
        output = self._export()
        self.assertIn('Book', output)
        self.assertIn('Not included', output)
        self.assertIn('media/credentials', output)

    def test_it_refuses_a_directory_that_does_not_exist(self):
        from django.core.management import call_command
        from django.core.management.base import CommandError
        with self.assertRaises(CommandError):
            call_command('export_library', '-o',
                         os.path.join(self.dir, 'nope', 'x.json'), stdout=StringIO())


class WrittenLocationReusesBoardsTests(TestCase):
    """An import must file books on the boards that exist, not make twins."""

    def setUp(self):
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf C', map_x=0, map_y=0)

    def _resolve(self, text):
        from library.views import _resolve_written_location
        return _resolve_written_location(text)

    def test_column_1_finds_a_board_stored_without_a_column(self):
        board = ShelfLevel.objects.create(shelf=self.shelf, level_number=3)
        before = ShelfLevel.objects.count()
        self.assertEqual(self._resolve('Shelf C Column 1 Level 3'), board)
        self.assertEqual(ShelfLevel.objects.count(), before,
                         'a duplicate board was created')

    def test_no_column_finds_a_board_stored_as_column_1(self):
        board = ShelfLevel.objects.create(shelf=self.shelf, level_number=3,
                                          column_number=1)
        before = ShelfLevel.objects.count()
        self.assertEqual(self._resolve('Shelf C Level 3'), board)
        self.assertEqual(ShelfLevel.objects.count(), before)

    def test_importing_the_same_sheet_twice_adds_no_boards(self):
        ShelfLevel.objects.create(shelf=self.shelf, level_number=3)
        first = self._resolve('Shelf C Column 1 Level 3')
        count = ShelfLevel.objects.count()
        for _ in range(4):
            self.assertEqual(self._resolve('Shelf C Column 1 Level 3'), first)
        self.assertEqual(ShelfLevel.objects.count(), count)

    def test_a_genuinely_divided_shelf_still_separates_its_bays(self):
        c1 = ShelfLevel.objects.create(shelf=self.shelf, level_number=1, column_number=1)
        c2 = ShelfLevel.objects.create(shelf=self.shelf, level_number=1, column_number=2)
        self.assertEqual(self._resolve('Shelf C Column 1 Level 1'), c1)
        self.assertEqual(self._resolve('Shelf C Column 2 Level 1'), c2)

    def test_column_2_never_borrows_column_1(self):
        ShelfLevel.objects.create(shelf=self.shelf, level_number=1, column_number=1)
        before = ShelfLevel.objects.count()
        made = self._resolve('Shelf C Column 2 Level 1')
        self.assertEqual(made.column_number, 2)
        self.assertEqual(ShelfLevel.objects.count(), before + 1,
                         'column 2 should be created, not matched to column 1')

    def test_a_bare_level_on_a_divided_shelf_takes_the_first_bay(self):
        c1 = ShelfLevel.objects.create(shelf=self.shelf, level_number=2, column_number=1)
        ShelfLevel.objects.create(shelf=self.shelf, level_number=2, column_number=2)
        before = ShelfLevel.objects.count()
        self.assertEqual(self._resolve('Shelf C Level 2'), c1)
        self.assertEqual(ShelfLevel.objects.count(), before)

    def test_the_top_is_found_whatever_column_the_sheet_mentions(self):
        top = ShelfLevel.objects.create(shelf=self.shelf, level_number=6, is_top=True)
        before = ShelfLevel.objects.count()
        self.assertEqual(self._resolve('Shelf C Top'), top)
        self.assertEqual(self._resolve('Shelf C Column 1 Top'), top)
        self.assertEqual(ShelfLevel.objects.count(), before)

    def test_a_board_that_really_is_missing_is_still_created(self):
        before = ShelfLevel.objects.count()
        made = self._resolve('Shelf C Level 4')
        self.assertIsNotNone(made)
        self.assertEqual(made.level_number, 4)
        self.assertEqual(ShelfLevel.objects.count(), before + 1)

    def test_an_import_files_books_on_the_existing_board(self):
        """End to end: the sheet's location must not orphan the real board."""
        board = ShelfLevel.objects.create(shelf=self.shelf, level_number=3)
        user = _admin(modules='books')
        client = _signed_in(user)

        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        wb.active.append(['Title', 'Author', 'Genre', 'Location'])
        wb.active.append(['A Book', 'An Author', 'EDUCATION',
                          'Shelf C Column 1 Level 3'])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile(
            'b.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        before = ShelfLevel.objects.count()
        r = client.post('/admin-portal/import-books/', {'excel_file': upload})
        self.assertTrue(r.json()['success'], r.json())
        self.assertEqual(ShelfLevel.objects.count(), before,
                         'the import created a duplicate board')
        self.assertEqual(Book.objects.get(title='A Book').shelf_level, board)


class TemplateCommentsDoNotRenderTests(TestCase):
    """{# #} does not span lines, and a multi-line one prints into the page."""

    def test_no_template_has_a_multi_line_short_comment(self):
        import re
        from django.conf import settings

        # Pattern for a comment tag that spans lines.
        pattern = re.compile(r"\{#(?:(?!#\}).)*\n(?:(?!#\}).)*#\}", re.S)

        offenders = []
        root = os.path.join(settings.BASE_DIR, "templates")
        for folder, _dirs, files in os.walk(root):
            for name in files:
                if not name.endswith(".html"):
                    continue
                path = os.path.join(folder, name)
                with io.open(path, encoding="utf-8") as handle:
                    for hit in pattern.findall(handle.read()):
                        offenders.append("%s: %s" % (
                            os.path.relpath(path, root),
                            " ".join(hit.split())[:60]))

        self.assertEqual(offenders, [],
                         "Multi-line {# #} renders into the page; "
                         "use {% comment %} instead:\n" + "\n".join(offenders))

    def test_the_check_would_actually_catch_one(self):
        """A guard that cannot fail is not a guard."""
        import re
        pattern = re.compile(r"\{#(?:(?!#\}).)*\n(?:(?!#\}).)*#\}", re.S)
        self.assertTrue(pattern.search("{# first line\n   second line #}"))
        self.assertFalse(pattern.search("{# all on one line #}"))


class ImportShelfOrderTests(TestCase):
    """The sheet's row order is the shelf's order."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf C', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=3)

    def _import(self, rows, header=None):
        import openpyxl
        from django.core.files.uploadedfile import SimpleUploadedFile
        wb = openpyxl.Workbook()
        wb.active.append(header or ['Title', 'Author', 'Genre', 'Location'])
        for row in rows:
            wb.active.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile(
            'b.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r = self.client.post('/admin-portal/import-books/', {'excel_file': upload})
        self.assertTrue(r.json()['success'], r.json())
        return r.json()

    def _shelf_order(self):
        """Titles in the order the shelf actually shows them."""
        from django.db.models import F
        return list(Book.objects.filter(shelf_level=self.level)
                    .order_by(F('shelf_slot').asc(nulls_last=True), 'title', 'book_id')
                    .values_list('title', flat=True))

    LOC = 'Shelf C Column 1 Level 3'

    def test_the_sheet_order_survives_instead_of_going_alphabetical(self):
        self._import([
            ['Zebra Handbook', 'A', 'REF', self.LOC],
            ['Apple Growing', 'B', 'REF', self.LOC],
            ['Mango Farming', 'C', 'REF', self.LOC],
        ])
        self.assertEqual(self._shelf_order(),
                         ['Zebra Handbook', 'Apple Growing', 'Mango Farming'])

    def test_slots_are_consecutive_from_one(self):
        self._import([['One', 'A', 'REF', self.LOC],
                      ['Two', 'B', 'REF', self.LOC],
                      ['Three', 'C', 'REF', self.LOC]])
        slots = list(Book.objects.filter(shelf_level=self.level)
                     .order_by('shelf_slot').values_list('shelf_slot', flat=True))
        self.assertEqual(slots, [1, 2, 3])

    def test_an_explicit_slot_column_wins(self):
        self._import(
            [['First', 'A', 'REF', self.LOC, 3],
             ['Second', 'B', 'REF', self.LOC, 1],
             ['Third', 'C', 'REF', self.LOC, 2]],
            header=['Title', 'Author', 'Genre', 'Location', 'Slot'])
        self.assertEqual(self._shelf_order(), ['Second', 'Third', 'First'])

    def test_a_later_import_appends_rather_than_jumping_the_queue(self):
        self._import([['One', 'A', 'REF', self.LOC],
                      ['Two', 'B', 'REF', self.LOC]])
        self._import([['Three', 'C', 'REF', self.LOC],
                      ['Four', 'D', 'REF', self.LOC]])
        self.assertEqual(self._shelf_order(), ['One', 'Two', 'Three', 'Four'])

    def test_books_already_on_the_board_keep_their_visible_order(self):
        """Numbering the unnumbered must not appear to move anything."""
        Book.objects.create(title='Bravo', author='X', shelf_level=self.level)
        Book.objects.create(title='Alpha', author='X', shelf_level=self.level)
        before = self._shelf_order()
        self.assertEqual(before, ['Alpha', 'Bravo'])       # alphabetical fallback

        self._import([['New Arrival', 'C', 'REF', self.LOC]])
        self.assertEqual(self._shelf_order(), ['Alpha', 'Bravo', 'New Arrival'])

    def test_copies_of_one_row_stand_together(self):
        self._import(
            [['Doubled', 'A', 'REF', self.LOC, 2],
             ['After It', 'B', 'REF', self.LOC, 1]],
            header=['Title', 'Author', 'Genre', 'Location', 'Quantity'])
        self.assertEqual(self._shelf_order(), ['Doubled', 'Doubled', 'After It'])

    def test_books_with_no_location_get_no_slot(self):
        """A book nobody has placed must not claim the first space anywhere."""
        self._import([['Unplaced', 'A', 'REF', '']])
        book = Book.objects.get(title='Unplaced')
        self.assertIsNone(book.shelf_level)
        self.assertIsNone(book.shelf_slot)

    def test_two_boards_are_numbered_independently(self):
        other = ShelfLevel.objects.create(shelf=self.shelf, level_number=4)
        self._import([['A1', 'A', 'REF', self.LOC],
                      ['B1', 'B', 'REF', 'Shelf C Column 1 Level 4'],
                      ['A2', 'C', 'REF', self.LOC]])
        self.assertEqual(
            sorted(Book.objects.filter(shelf_level=self.level)
                   .values_list('title', 'shelf_slot')),
            [('A1', 1), ('A2', 2)])
        self.assertEqual(
            list(Book.objects.filter(shelf_level=other)
                 .values_list('title', 'shelf_slot')),
            [('B1', 1)])

    def test_an_alphabetical_sheet_still_comes_out_alphabetical(self):
        """Honouring the sheet costs nothing when the sheet is already sorted."""
        self._import([['Apple', 'A', 'REF', self.LOC],
                      ['Mango', 'B', 'REF', self.LOC],
                      ['Zebra', 'C', 'REF', self.LOC]])
        self.assertEqual(self._shelf_order(), ['Apple', 'Mango', 'Zebra'])


class LabelPrintOrderTests(TestCase):
    """A pile of labels comes out either in walking order or in A-Z order."""

    def setUp(self):
        from library import labels
        self.labels = labels
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.a = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.b = Shelf.objects.create(room=room, name='Shelf B', map_x=0, map_y=0)
        self.a1 = ShelfLevel.objects.create(shelf=self.a, level_number=1)
        self.a2 = ShelfLevel.objects.create(shelf=self.a, level_number=2)
        self.top = ShelfLevel.objects.create(shelf=self.a, level_number=3, is_top=True)
        self.b1 = ShelfLevel.objects.create(shelf=self.b, level_number=1)

    def _book(self, title, level=None, slot=None, author='X'):
        return Book.objects.create(title=title, author=author, genre='REF',
                                   shelf_level=level, shelf_slot=slot,
                                   status='Available')

    def _titles(self, books):
        return [b.title for b in books]

    def test_shelf_order_follows_the_slot_not_the_alphabet(self):
        """This is the whole point: a shelf is not in alphabetical order."""
        z = self._book('Zebra', self.a1, slot=1)
        a = self._book('Apple', self.a1, slot=2)
        m = self._book('Mango', self.a1, slot=3)
        got = self.labels.sort_for_printing([a, m, z])
        self.assertEqual(self._titles(got), ['Zebra', 'Apple', 'Mango'])

    def test_shelf_order_walks_shelves_then_boards(self):
        b = self._book('On Shelf B', self.b1, slot=1)
        a2 = self._book('Second board', self.a2, slot=1)
        a1 = self._book('First board', self.a1, slot=1)
        got = self.labels.sort_for_printing([b, a2, a1])
        self.assertEqual(self._titles(got),
                         ['First board', 'Second board', 'On Shelf B'])

    def test_the_top_of_a_case_comes_last_within_it(self):
        top = self._book('On the top', self.top, slot=1)
        low = self._book('Inside it', self.a1, slot=1)
        got = self.labels.sort_for_printing([top, low])
        self.assertEqual(self._titles(got), ['Inside it', 'On the top'])

    def test_a_book_with_no_slot_falls_back_to_its_title(self):
        no_slot_b = self._book('Bravo', self.a1)
        no_slot_a = self._book('Alpha', self.a1)
        got = self.labels.sort_for_printing([no_slot_b, no_slot_a])
        self.assertEqual(self._titles(got), ['Alpha', 'Bravo'])

    def test_slotted_books_come_before_unslotted_on_the_same_board(self):
        placed = self._book('Zulu', self.a1, slot=1)
        unplaced = self._book('Alpha', self.a1)
        got = self.labels.sort_for_printing([unplaced, placed])
        self.assertEqual(self._titles(got), ['Zulu', 'Alpha'])

    def test_books_on_no_shelf_at_all_come_last(self):
        loose = self._book('Not placed yet')
        shelved = self._book('Zulu', self.a1, slot=1)
        got = self.labels.sort_for_printing([loose, shelved])
        self.assertEqual(self._titles(got), ['Zulu', 'Not placed yet'])

    def test_alphabetical_ignores_the_shelf_entirely(self):
        z = self._book('Zebra', self.a1, slot=1)
        a = self._book('Apple', self.b1, slot=1)
        m = self._book('Mango', self.a2, slot=1)
        got = self.labels.sort_alphabetically([z, a, m])
        self.assertEqual(self._titles(got), ['Apple', 'Mango', 'Zebra'])

    def test_alphabetical_breaks_ties_on_author(self):
        second = self._book('Same Title', self.a1, slot=1, author='Zed')
        first = self._book('Same Title', self.a1, slot=2, author='Adams')
        got = self.labels.sort_alphabetically([second, first])
        self.assertEqual([b.author for b in got], ['Adams', 'Zed'])


class LabelSheetSortParamTests(TestCase):
    """The dialog's choice reaches the PDF."""

    def setUp(self):
        self.client = _signed_in(_admin(modules='books'))
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.z = Book.objects.create(title='Zebra', author='A', genre='REF',
                                     shelf_level=level, shelf_slot=1,
                                     status='Available', qr_code=str(uuid4()))
        self.a = Book.objects.create(title='Apple', author='B', genre='REF',
                                     shelf_level=level, shelf_slot=2,
                                     status='Available', qr_code=str(uuid4()))

    def _pdf(self, sort=None):
        url = '/admin-portal/book-qr-labels/?ids=%d,%d' % (self.a.book_id, self.z.book_id)
        if sort:
            url += '&sort=' + sort
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200, r.content[:200])
        self.assertEqual(r['Content-Type'], 'application/pdf')
        return r.content

    def test_every_order_produces_a_pdf(self):
        for sort in (None, 'location', 'title', 'alphabetical', 'az', 'id'):
            self.assertTrue(self._pdf(sort).startswith(b'%PDF'), sort)

    def _titles_in_order(self, sort=None):
        """The titles as they appear on the sheet, in the order printed."""
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(self._pdf(sort)))
        text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        seen = []
        for line in text.splitlines():
            line = line.strip()
            for title in ('Zebra', 'Apple'):
                if line.startswith(title) and title not in seen:
                    seen.append(title)
        return seen

    def test_shelf_order_prints_by_slot(self):
        """Zebra is slot 1 and Apple slot 2, so Zebra prints first."""
        self.assertEqual(self._titles_in_order('location'), ['Zebra', 'Apple'])

    def test_alphabetical_prints_a_before_z(self):
        self.assertEqual(self._titles_in_order('title'), ['Apple', 'Zebra'])

    def test_the_default_is_shelf_order(self):
        self.assertEqual(self._titles_in_order(), ['Zebra', 'Apple'])

    def test_an_unknown_sort_falls_back_to_shelf_order(self):
        self.assertEqual(self._titles_in_order('nonsense'), ['Zebra', 'Apple'])


class ShelfManagerBookActionsTests(TestCase):
    """Moving and unshelving from the page that draws the shelves."""

    def test_shelf_manager_carries_the_mover(self):
        """The books are dealt with on the page that draws the shelves."""
        client = _signed_in(_admin(email='shelf-page@example.invalid', modules=''))
        html = client.get('/admin-portal/shelf-manager/').content.decode()
        for handle in ('paneLList', 'paneRList', 'movePane', 'unshelfPane',
                       'selectTail', 'swapPanes', 'board-books', 'move-books'):
            self.assertIn(handle, html, handle + ' is missing from Shelf Manager')

    def test_both_jobs_are_modals_behind_their_own_button(self):
        """Moving books and reordering a board are separate questions."""
        client = _signed_in(_admin(email='modals@example.invalid', modules=''))
        html = client.get('/admin-portal/shelf-manager/').content.decode()
        for handle in ('moverModal', 'reorderModal', 'openMover(',
                       'openReorder(', 'saveReorder', 'sortReorder'):
            self.assertIn(handle, html, handle + ' is missing')
        # Both panels start closed.
        for modal in ('moverModal', 'reorderModal'):
            at = html.index('id="%s"' % modal)
            self.assertIn('hidden', html[at - 120:at + 120],
                          modal + ' should start hidden')

    def test_manage_books_no_longer_offers_to_move_them(self):
        """Moving books between shelves is done in the shelf manager."""
        client = _signed_in(_admin(email='books-page@example.invalid', modules='books'))
        html = client.get('/admin-portal/management/').content.decode()
        for gone in ('toolbarMoveBtn', 'moveBooksModal', 'openMoveBooks',
                     'confirmMoveBooks'):
            self.assertNotIn(gone, html, gone + ' should have moved to Shelf Manager')
        self.assertIn('toolbarDeleteBtn', html, 'bulk delete stays in Manage Books')


class MoveBooksTests(TestCase):
    """Moving books between boards, and off the shelves."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.a = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.b = Shelf.objects.create(room=room, name='Shelf B', map_x=0, map_y=0)
        self.a1 = ShelfLevel.objects.create(shelf=self.a, level_number=1)
        self.b1 = ShelfLevel.objects.create(shelf=self.b, level_number=1)

    def _book(self, title, level=None, slot=None):
        return Book.objects.create(title=title, author='X', genre='REF',
                                   shelf_level=level, shelf_slot=slot,
                                   status='Available')

    def _move(self, books, to):
        ids = ','.join(str(b.book_id) for b in books)
        return self.client.post('/admin-portal/move-books/',
                                {'ids': ids, 'level': to}).json()

    def _on(self, level):
        from django.db.models import F
        return list(Book.objects.filter(shelf_level=level)
                    .order_by(F('shelf_slot').asc(nulls_last=True), 'book_id')
                    .values_list('title', 'shelf_slot'))

    # Endpoints used by both Shelf Manager and Manage Books.

    def test_an_administrator_holding_no_modules_can_still_move_books(self):
        """The trap: Shelf Manager opens for them, so its buttons must work."""
        bare = _admin(email='bare-admin@example.invalid', modules='')
        client = _signed_in(bare)
        book = self._book('Maths', self.b1, slot=1)

        page = client.get('/admin-portal/shelf-manager/')
        self.assertEqual(page.status_code, 200, 'they can open the page')

        r = client.post('/admin-portal/move-books/',
                        {'ids': str(book.book_id), 'level': self.a1.shelf_level_id})
        self.assertEqual(r.status_code, 200, 'a redirect here would be a dead button')
        self.assertTrue(r.json()['success'], r.json())
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.a1)

    def test_shelf_staff_can_move_books_without_the_catalogue_module(self):
        """Shelving is shelf work. It is the module for standing at the bay."""
        from library.models import User
        staff = User.objects.create(
            fullname='Shelver', email='shelver@example.invalid',
            password_hash=hash_password('SmokeTest123'),
            role='Staff', account_status='Active', modules='shelf')
        client = _signed_in(staff)
        book = self._book('Maths', self.b1, slot=1)
        r = client.post('/admin-portal/move-books/',
                        {'ids': str(book.book_id), 'level': self.a1.shelf_level_id})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.json()['success'], r.json())

    def test_staff_with_neither_module_are_refused(self):
        from library.models import User
        nobody = User.objects.create(
            fullname='Desk', email='desk@example.invalid',
            password_hash=hash_password('SmokeTest123'),
            role='Staff', account_status='Active', modules='chat')
        client = _signed_in(nobody)
        book = self._book('Maths', self.b1, slot=1)
        r = client.post('/admin-portal/move-books/',
                        {'ids': str(book.book_id), 'level': self.a1.shelf_level_id})
        self.assertEqual(r.status_code, 302)
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.b1, 'the book must not have moved')

    def test_unshelving_is_reported_as_taking_off_not_moving_to(self):
        """"moved to off the shelves" is not a sentence."""
        book = self._book('Maths', self.b1, slot=1)
        r = self._move([book], 'none')
        self.assertTrue(r['success'], r)
        self.assertEqual(r['message'], '1 book taken off the shelves.')
        self.assertNotIn('moved to', r['message'])

    def test_moving_onto_a_board_still_names_the_board(self):
        book = self._book('Maths', self.b1, slot=1)
        r = self._move([book], self.a1.shelf_level_id)
        self.assertIn('moved to Shelf A', r['message'])

    def test_a_book_moves_to_another_shelf(self):
        book = self._book('Maths', self.b1, slot=1)
        r = self._move([book], self.a1.shelf_level_id)
        self.assertTrue(r['success'], r)
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.a1)

    def test_arrivals_land_after_what_is_already_there(self):
        """The bug assign_books_to_level has: two books claiming one slot."""
        self._book('Resident one', self.a1, slot=1)
        self._book('Resident two', self.a1, slot=2)
        incoming = self._book('Arriving', self.b1, slot=1)
        self._move([incoming], self.a1.shelf_level_id)
        self.assertEqual(self._on(self.a1),
                         [('Resident one', 1), ('Resident two', 2), ('Arriving', 3)])

    def test_no_two_books_share_a_slot_after_a_move(self):
        for i in range(3):
            self._book('Resident %d' % i, self.a1, slot=i + 1)
        movers = [self._book('Mover %d' % i, self.b1, slot=i + 1) for i in range(3)]
        self._move(movers, self.a1.shelf_level_id)
        slots = [slot for _, slot in self._on(self.a1)]
        self.assertEqual(sorted(slots), list(range(1, 7)))
        self.assertEqual(len(set(slots)), 6, 'two books share a position')

    def test_a_whole_board_keeps_its_order_when_moved(self):
        movers = [self._book('Book %d' % i, self.b1, slot=i + 1) for i in range(4)]
        self._move(movers, self.a1.shelf_level_id)
        self.assertEqual([t for t, _ in self._on(self.a1)],
                         ['Book 0', 'Book 1', 'Book 2', 'Book 3'])

    def test_moving_to_none_unshelves_and_clears_the_slot(self):
        book = self._book('Loose', self.a1, slot=1)
        r = self._move([book], 'none')
        self.assertTrue(r['success'], r)
        book.refresh_from_db()
        self.assertIsNone(book.shelf_level)
        self.assertIsNone(book.shelf_slot)

    def test_the_move_is_recorded(self):
        book = self._book('Tracked', self.b1, slot=1)
        before = SystemLog.objects.count()
        self._move([book], self.a1.shelf_level_id)
        self.assertEqual(SystemLog.objects.count(), before + 1)
        entry = SystemLog.objects.order_by('-log_id').first()
        self.assertIn('Shelf A', entry.detail)

    def test_the_reply_names_where_they_went(self):
        book = self._book('Named', self.b1, slot=1)
        r = self._move([book], self.a1.shelf_level_id)
        self.assertIn('Shelf A', r['message'])
        self.assertEqual(r['moved'], 1)

    def test_nothing_selected_is_refused(self):
        r = self.client.post('/admin-portal/move-books/',
                             {'ids': '', 'level': self.a1.shelf_level_id}).json()
        self.assertFalse(r['success'])

    def test_a_destination_that_does_not_exist_is_refused(self):
        book = self._book('Stays', self.b1, slot=1)
        r = self._move([book], 999999)
        self.assertFalse(r['success'])
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.b1, 'the book moved anyway')

    def test_no_destination_at_all_is_refused(self):
        book = self._book('Stays', self.b1, slot=1)
        r = self.client.post('/admin-portal/move-books/',
                             {'ids': str(book.book_id)}).json()
        self.assertFalse(r['success'])
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.b1)

    def test_a_borrowed_book_moves_with_its_record_intact(self):
        book = self._book('Out on loan', self.b1, slot=1)
        book.status = 'Borrowed'
        book.save()
        self._move([book], self.a1.shelf_level_id)
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.a1)
        self.assertEqual(book.status, 'Borrowed')

    def test_it_is_guarded_against_staff_holding_neither_module(self):
        """The gate widened, deliberately, and this is what it still refuses."""
        from library.models import User
        outsider = _signed_in(User.objects.create(
            fullname='Front desk', email='no-books@example.invalid',
            password_hash=hash_password('SmokeTest123'),
            role='Staff', account_status='Active', modules='chat'))
        book = self._book('Guarded', self.b1, slot=1)
        r = outsider.post('/admin-portal/move-books/',
                          {'ids': str(book.book_id), 'level': self.a1.shelf_level_id})
        self.assertIn(r.status_code, (302, 403))
        book.refresh_from_db()
        self.assertEqual(book.shelf_level, self.b1)


class BooksLocationFilterTests(TestCase):
    """Finding books by where they are -- including the ones that are nowhere."""

    def setUp(self):
        self.client = _signed_in(_admin(modules='books'))
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        self.a = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.a1 = ShelfLevel.objects.create(shelf=self.a, level_number=1)
        self.a2 = ShelfLevel.objects.create(shelf=self.a, level_number=2)
        b = Shelf.objects.create(room=room, name='Shelf B', map_x=0, map_y=0)
        self.b1 = ShelfLevel.objects.create(shelf=b, level_number=1)
        for level, n in ((self.a1, 3), (self.a2, 2), (self.b1, 4)):
            for i in range(n):
                Book.objects.create(title='On %s %d' % (level.shelf_level_id, i),
                                    author='X', genre='REF', shelf_level=level)
        for i in range(5):
            Book.objects.create(title='Loose %d' % i, author='X', genre='REF')

    def _rows(self, query=''):
        r = self.client.get('/admin-portal/management/' + query)
        self.assertEqual(r.status_code, 200)
        return r.context['books'].paginator.count

    def test_filtering_by_shelf(self):
        self.assertEqual(self._rows('?shelf=%d' % self.a.shelf_id), 5)

    def test_filtering_by_board(self):
        self.assertEqual(self._rows('?level=%d' % self.a1.shelf_level_id), 3)

    def test_the_board_wins_over_its_shelf(self):
        self.assertEqual(
            self._rows('?shelf=%d&level=%d' % (self.a.shelf_id, self.a2.shelf_level_id)), 2)

    def test_filtering_for_the_unshelved(self):
        self.assertEqual(self._rows('?shelf=none'), 5)

    def test_nonsense_is_ignored_rather_than_emptying_the_table(self):
        for bad in ('?shelf=banana', '?shelf=', '?level=notanumber'):
            self.assertEqual(self._rows(bad), 14, bad)

    def test_it_combines_with_the_other_filters(self):
        Book.objects.filter(shelf_level=self.a1).update(status='Borrowed')
        self.assertEqual(self._rows('?shelf=%d&status=Borrowed' % self.a.shelf_id), 3)

    def test_the_filter_survives_paging(self):
        r = self.client.get('/admin-portal/management/?shelf=none')
        self.assertIn('shelf=none', r.context['querystring'])

    def test_the_page_offers_every_shelf_and_a_count_of_the_unshelved(self):
        r = self.client.get('/admin-portal/management/')
        names = [s['name'] for s in r.context['shelf_choices']]
        self.assertIn('Shelf A', names)
        self.assertIn('Shelf B', names)
        self.assertEqual(r.context['unshelved_count'], 5)


class DeleteBooksTests(TestCase):
    """Deleting the ticked books, and what it takes with them."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Main', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=0, map_y=0)
        self.level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        self.patron = Patron.objects.create(
            first_name='Pat', last_name='Ron', email='pat@example.invalid',
            patron_type='Student')

    def _book(self, title, status='Available'):
        return Book.objects.create(title=title, author='X', genre='REF',
                                   shelf_level=self.level, status=status)

    def _post(self, books, confirm=False):
        data = {'ids': ','.join(str(b.book_id) for b in books)}
        if confirm:
            data['confirm'] = '1'
        return self.client.post('/admin-portal/delete-books/', data).json()

    def test_asking_first_deletes_nothing(self):
        books = [self._book('One'), self._book('Two')]
        before = Book.objects.count()
        r = self._post(books)
        self.assertTrue(r['success'], r)
        self.assertTrue(r['preview'])
        self.assertEqual(r['deletable'], 2)
        self.assertEqual(Book.objects.count(), before, 'the preview deleted books')

    def test_confirming_deletes_them(self):
        books = [self._book('One'), self._book('Two')]
        r = self._post(books, confirm=True)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['deleted'], 2)
        self.assertEqual(Book.objects.filter(title__in=['One', 'Two']).count(), 0)

    def test_a_borrowed_book_is_kept_not_deleted(self):
        """Deleting it would lose the only record that it is out."""
        out = self._book('Out on loan', status='Borrowed')
        ok = self._book('On the shelf')
        r = self._post([out, ok], confirm=True)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['deleted'], 1)
        self.assertTrue(Book.objects.filter(pk=out.pk).exists())
        self.assertFalse(Book.objects.filter(pk=ok.pk).exists())

    def test_an_overdue_book_is_kept_too(self):
        out = self._book('Overdue', status='Overdue')
        r = self._post([out], confirm=True)
        self.assertFalse(r['success'])
        self.assertTrue(Book.objects.filter(pk=out.pk).exists())

    def test_the_preview_names_the_books_it_will_keep(self):
        out = self._book('Out on loan', status='Borrowed')
        r = self._post([out, self._book('Fine')])
        self.assertEqual(r['on_loan_count'], 1)
        self.assertEqual(r['on_loan'][0]['title'], 'Out on loan')

    def test_the_preview_counts_the_loan_history_that_would_go(self):
        book = self._book('Has history')
        for _ in range(3):
            Transaction.objects.create(book=book, patron=self.patron,
                                       transaction_type='Borrow')
        r = self._post([book])
        self.assertEqual(r['loan_records'], 3)

    def test_deleting_really_does_take_the_loan_history(self):
        """Not a warning about a hypothetical: this is what CASCADE does."""
        book = self._book('Has history')
        Transaction.objects.create(book=book, patron=self.patron,
                                   transaction_type='Borrow')
        self.assertEqual(Transaction.objects.count(), 1)
        self._post([book], confirm=True)
        self.assertEqual(Transaction.objects.count(), 0)

    def test_everything_selected_being_on_loan_is_refused(self):
        out = [self._book('A', status='Borrowed'), self._book('B', status='Borrowed')]
        r = self._post(out, confirm=True)
        self.assertFalse(r['success'])
        self.assertEqual(Book.objects.filter(status='Borrowed').count(), 2)

    def test_the_deletion_is_recorded(self):
        book = self._book('Tracked')
        before = SystemLog.objects.count()
        self._post([book], confirm=True)
        self.assertEqual(SystemLog.objects.count(), before + 1)
        self.assertIn('Tracked', SystemLog.objects.order_by('-log_id').first().detail)

    def test_nothing_selected_is_refused(self):
        r = self.client.post('/admin-portal/delete-books/', {'ids': ''}).json()
        self.assertFalse(r['success'])

    def test_it_needs_the_books_module(self):
        outsider = _signed_in(_admin(email='no-books@example.invalid', modules=''))
        book = self._book('Guarded')
        r = outsider.post('/admin-portal/delete-books/',
                          {'ids': str(book.book_id), 'confirm': '1'})
        self.assertIn(r.status_code, (302, 403))
        self.assertTrue(Book.objects.filter(pk=book.pk).exists())

    def test_too_many_at_once_is_refused(self):
        from library.views import MAX_BOOKS_PER_MOVE
        ids = ','.join(str(i) for i in range(MAX_BOOKS_PER_MOVE + 2))
        r = self.client.post('/admin-portal/delete-books/',
                             {'ids': ids, 'confirm': '1'}).json()
        self.assertFalse(r['success'])


class RowDeleteAsksFirstTests(TestCase):
    """The trash icon on a row used to delete on one click, with no question."""

    def test_the_row_delete_form_carries_a_confirmation(self):
        user = _admin(modules='books')
        client = _signed_in(user)
        plan = FloorPlan.objects.create(name='G', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='R', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='S', map_x=0, map_y=0)
        level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        Book.objects.create(title='Deletable', author='X', genre='REF',
                            shelf_level=level, status='Available')

        html = client.get('/admin-portal/management/').content.decode('utf-8', 'replace')
        # The row's delete form must opt in to ayla-dialog's declarative confirmation.
        self.assertIn('data-confirm', html, 'the row delete has no confirmation')
        self.assertIn('Deletable', html, 'the book is not even on the page')


class AnalyticsTests(TestCase):
    """The charts, and the arithmetic behind them."""

    def setUp(self):
        from library import analytics
        self.analytics = analytics
        self.start = date(2026, 6, 1)
        self.end = date(2026, 6, 30)
        self.patron = Patron.objects.create(
            first_name='Pat', last_name='Ron', email='p@example.invalid',
            patron_type='Student')

    def _borrow(self, book, day):
        """A borrow on a given day."""
        tx = Transaction.objects.create(book=book, patron=self.patron,
                                        transaction_type='Borrow')
        Transaction.objects.filter(pk=tx.pk).update(
            transaction_date=date(2026, 6, day))
        return tx

    def _visit(self, day, in_hour, out_hour=None, auto_closed=False,
               purpose='Study', school='Central', patron=None):
        tz = timezone.get_current_timezone()
        entry = timezone.make_aware(datetime(2026, 6, day, in_hour, 0), tz)
        exit_at = (timezone.make_aware(datetime(2026, 6, day, out_hour, 0), tz)
                   if out_hour is not None else None)
        return PatronLog.objects.create(
            patron=patron or self.patron, entry_time=entry, exit_time=exit_at,
            auto_closed=auto_closed, purpose_of_visit=purpose, school=school)

    # Arrivals vs occupancy

    def test_arrivals_counts_the_moment_people_walk_in(self):
        for _ in range(3):
            self._visit(1, 9, 10)
        self._visit(1, 14, 15)
        out = self.analytics.visits_by_hour(self.start, self.end)
        self.assertEqual(out['peak'], '9 AM')
        self.assertEqual(out['peak_count'], 3)
        self.assertEqual(out['total'], 4)

    def test_occupancy_is_not_the_same_as_arrivals(self):
        """Three arrive at 9 and leave at 9:59; one arrives at 2 and stays till 6."""
        for _ in range(3):
            self._visit(1, 9, 9)
        self._visit(1, 14, 18)
        arrivals = self.analytics.visits_by_hour(self.start, self.end)
        occupancy = self.analytics.occupancy_by_hour(self.start, self.end)
        self.assertEqual(arrivals['peak'], '9 AM')
        rows = dict(occupancy['rows'])
        self.assertEqual(rows['9 AM'], '3.0')
        # The long visit covers 2, 3, 4, 5 and 6 PM.
        for hour in ('2 PM', '3 PM', '4 PM', '5 PM', '6 PM'):
            self.assertEqual(rows[hour], '1.0', hour)

    def test_a_visit_spanning_hours_counts_in_each_one(self):
        self._visit(1, 9, 12)
        rows = dict(self.analytics.occupancy_by_hour(self.start, self.end)['rows'])
        for hour in ('9 AM', '10 AM', '11 AM', '12 PM'):
            self.assertEqual(rows[hour], '1.0', hour)

    def test_guessed_exits_are_left_out_of_occupancy(self):
        """A visit closed by the nightly sweep has an invented exit time."""
        self._visit(1, 9, 17, auto_closed=True)
        out = self.analytics.occupancy_by_hour(self.start, self.end)
        self.assertEqual(out['measured_visits'], 0)
        self.assertEqual(out['excluded_visits'], 1)
        self.assertFalse(out['chart']['has_data'])

    def test_visits_still_open_are_left_out_of_occupancy(self):
        self._visit(1, 9, None)
        out = self.analytics.occupancy_by_hour(self.start, self.end)
        self.assertEqual(out['measured_visits'], 0)
        self.assertEqual(out['excluded_visits'], 1)

    def test_occupancy_averages_over_the_days_measured(self):
        """Two people on each of two days is an average of two, not four."""
        for day in (1, 2):
            self._visit(day, 9, 10)
            self._visit(day, 9, 10)
        out = self.analytics.occupancy_by_hour(self.start, self.end)
        self.assertEqual(out['days_measured'], 2)
        self.assertEqual(dict(out['rows'])['9 AM'], '2.0')

    # The other visit charts

    def test_weekday_chart_finds_the_busiest_day(self):
        self._visit(1, 9, 10)          # 2026-06-01 is a Monday
        self._visit(1, 10, 11)
        self._visit(2, 9, 10)
        out = self.analytics.visits_by_weekday(self.start, self.end)
        self.assertEqual(out['busiest'], 'Monday')

    def test_purpose_chart_ranks_by_count(self):
        for _ in range(3):
            self._visit(1, 9, 10, purpose='Research')
        self._visit(1, 11, 12, purpose='Reading')
        out = self.analytics.visits_by_purpose(self.start, self.end)
        self.assertEqual(out['top'], 'Research')
        self.assertEqual(dict(out['rows'])['Research'], '3')

    def test_visitor_type_counts_per_visit_not_per_person(self):
        """A regular counts each time they come, because the question is who is using the room."""
        for _ in range(4):
            self._visit(1, 9, 10)
        out = self.analytics.visitors_by_type(self.start, self.end)
        self.assertEqual(dict(out['rows'])['Student'], '4')

    def test_school_chart_counts_distinct_schools(self):
        self._visit(1, 9, 10, school='Central')
        self._visit(1, 9, 10, school='Northside')
        out = self.analytics.visitors_by_school(self.start, self.end)
        self.assertEqual(out['distinct'], 2)

    # The collection

    def test_condition_chart_reports_the_damaged_share(self):
        for _ in range(3):
            Book.objects.create(title='Fine', author='X', genre='REF',
                                condition='Good', status='Available')
        Book.objects.create(title='Broken', author='X', genre='REF',
                            condition='Damaged', status='Available')
        out = self.analytics.books_by_condition()
        self.assertEqual(out['damaged'], 1)
        self.assertEqual(out['damaged_share'], 25.0)

    def test_unshelved_summary_counts_books_on_no_shelf(self):
        plan = FloorPlan.objects.create(name='G', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='R', map_x=0, map_y=0)
        shelf = Shelf.objects.create(room=room, name='A', map_x=0, map_y=0)
        level = ShelfLevel.objects.create(shelf=shelf, level_number=1)
        Book.objects.create(title='Placed', author='X', genre='REF', shelf_level=level)
        Book.objects.create(title='Loose', author='X', genre='REF')
        out = self.analytics.unshelved_summary()
        self.assertEqual(out['count'], 1)
        self.assertEqual(out['shelved'], 1)
        self.assertEqual(out['share'], 50.0)

    def test_shelf_occupancy_ranks_fullest_first(self):
        plan = FloorPlan.objects.create(name='G', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='R', map_x=0, map_y=0)
        big = Shelf.objects.create(room=room, name='Big', map_x=0, map_y=0)
        small = Shelf.objects.create(room=room, name='Small', map_x=0, map_y=0)
        Shelf.objects.create(room=room, name='Empty', map_x=0, map_y=0)
        big_lv = ShelfLevel.objects.create(shelf=big, level_number=1)
        small_lv = ShelfLevel.objects.create(shelf=small, level_number=1)
        for i in range(3):
            Book.objects.create(title='B%d' % i, author='X', genre='REF', shelf_level=big_lv)
        Book.objects.create(title='S', author='X', genre='REF', shelf_level=small_lv)
        out = self.analytics.shelf_occupancy()
        self.assertEqual(out['rows'][0], ('Big', '3'))
        self.assertEqual(out['empty_shelves'], 1)

    # Borrowing and fines

    def test_most_borrowed_books_ranks_titles(self):
        popular = Book.objects.create(title='Popular', author='A', genre='FIC')
        quiet = Book.objects.create(title='Quiet', author='B', genre='FIC')
        for _ in range(3):
            self._borrow(popular, 5)
        self._borrow(quiet, 5)
        out = self.analytics.most_borrowed_books(self.start, self.end)
        self.assertEqual(out['rows'][0], ('Popular', 'A', '3'))
        self.assertEqual(out['total'], 4)

    def test_most_borrowed_genres_groups_across_titles(self):
        a = Book.objects.create(title='One', author='A', genre='Science')
        b = Book.objects.create(title='Two', author='B', genre='Science')
        c = Book.objects.create(title='Three', author='C', genre='Poetry')
        for book in (a, b, c):
            self._borrow(book, 5)
        out = self.analytics.most_borrowed_genres(self.start, self.end)
        self.assertEqual(out['rows'][0], ('Science', '2'))

    def test_never_borrowed_finds_stock_that_has_not_moved(self):
        moved = Book.objects.create(title='Moved', author='A', genre='FIC')
        Book.objects.create(title='Still here', author='B', genre='FIC')
        self._borrow(moved, 5)
        out = self.analytics.never_borrowed()
        self.assertEqual(out['count'], 1)
        self.assertEqual(out['rows'][0][0], 'Still here')

    def test_penalties_group_by_day_week_and_month(self):
        book = Book.objects.create(title='Late', author='A', genre='FIC')
        for day, amount in ((1, 10), (2, 5), (15, 20)):
            Transaction.objects.create(
                book=book, patron=self.patron, transaction_type='Borrow',
                transaction_date=date(2026, 6, day),
                return_date=date(2026, 6, day), fine_amount=amount)

        by_day = self.analytics.penalties_over_time(self.start, self.end, 'day')
        by_week = self.analytics.penalties_over_time(self.start, self.end, 'week')
        by_month = self.analytics.penalties_over_time(self.start, self.end, 'month')

        self.assertEqual(by_day['total'], 35.0)
        self.assertEqual(by_week['total'], 35.0)
        self.assertEqual(by_month['total'], 35.0)
        # 1 and 2 June fall in one week; 15 June in another.
        self.assertEqual(by_week['periods_with_fines'], 2)
        self.assertEqual(by_month['periods_with_fines'], 1)
        self.assertEqual(by_day['periods_with_fines'], 3)

    def test_penalties_average_is_over_periods_that_had_fines(self):
        book = Book.objects.create(title='Late', author='A', genre='FIC')
        for day in (1, 2):
            Transaction.objects.create(
                book=book, patron=self.patron, transaction_type='Borrow',
                transaction_date=date(2026, 6, day),
                return_date=date(2026, 6, day), fine_amount=10)
        out = self.analytics.penalties_over_time(self.start, self.end, 'day')
        self.assertEqual(out['average'], 10.0)

    # Empty data

    def test_every_builder_survives_an_empty_library(self):
        PatronLog.objects.all().delete()
        Book.objects.all().delete()
        Transaction.objects.all().delete()
        for call in (
            lambda: self.analytics.visits_by_hour(self.start, self.end),
            lambda: self.analytics.occupancy_by_hour(self.start, self.end),
            lambda: self.analytics.visits_by_weekday(self.start, self.end),
            lambda: self.analytics.visits_by_purpose(self.start, self.end),
            lambda: self.analytics.visitors_by_type(self.start, self.end),
            lambda: self.analytics.visitors_by_school(self.start, self.end),
            lambda: self.analytics.books_by_genre(),
            lambda: self.analytics.books_by_condition(),
            lambda: self.analytics.shelf_occupancy(),
            lambda: self.analytics.unshelved_summary(),
            lambda: self.analytics.most_borrowed_books(self.start, self.end),
            lambda: self.analytics.most_borrowed_genres(self.start, self.end),
            lambda: self.analytics.never_borrowed(),
            lambda: self.analytics.penalties_over_time(self.start, self.end, 'day'),
        ):
            self.assertIsNotNone(call())


class AnalyticsPageTests(TestCase):
    """The page itself."""

    def setUp(self):
        # Logs module so the Log Management page is reachable.
        self.client = _signed_in(_admin(modules='logs'))

    def test_the_page_renders(self):
        r = self.client.get('/admin-portal/analytics/')
        self.assertEqual(r.status_code, 200)

    def test_it_draws_every_section(self):
        html = self.client.get('/admin-portal/analytics/').content.decode('utf-8', 'replace')
        for title in ('Arrivals by hour', 'People inside, by hour',
                      'Visits by day of the week', 'Why people visit',
                      'Who visits', 'Visitors by school', 'Collection by genre',
                      'Condition of the collection', 'Books per shelf',
                      'Never borrowed', 'Most borrowed titles',
                      'Most borrowed genres'):
            self.assertIn(title, html, title)

    def test_the_fine_grouping_can_be_changed(self):
        for period in ('day', 'week', 'month'):
            r = self.client.get('/admin-portal/analytics/?period=' + period)
            self.assertEqual(r.context['period'], period)
            self.assertIn('per %s' % period, r.context['penalties']['chart']['title'])

    def test_a_nonsense_period_falls_back_to_day(self):
        r = self.client.get('/admin-portal/analytics/?period=fortnight')
        self.assertEqual(r.context['period'], 'day')

    def test_it_is_administrator_only(self):
        staff = User.objects.create(
            fullname='Staff', email='an-analytics-staff@example.invalid',
            password_hash=hash_password('SmokeTest123'),
            role='Staff', account_status='Active', modules='logs')
        r = _signed_in(staff).get('/admin-portal/analytics/')
        self.assertIn(r.status_code, (302, 403))

    def test_log_management_carries_the_peak_hour_chart(self):
        """The chart the research adviser asked for, where the visits are."""
        html = self.client.get('/admin-portal/log-management/').content.decode('utf-8', 'replace')
        self.assertIn('Arrivals by hour', html)


class CardNumberTests(TestCase):
    """Seven digits that have to survive being read off a card and typed."""

    def test_a_minted_number_validates(self):
        from library.cardnumbers import LENGTH, generate, is_valid

        number = generate(exists=lambda candidate: False)
        self.assertEqual(len(number), LENGTH)
        self.assertTrue(number.startswith('7'))
        self.assertTrue(is_valid(number))

    def test_a_mistyped_digit_is_rejected_without_a_lookup(self):
        from library.cardnumbers import generate, is_valid

        number = generate(exists=lambda candidate: False)
        wrong = list(number)
        wrong[3] = str((int(wrong[3]) + 1) % 10)
        self.assertFalse(is_valid(''.join(wrong)))

    def test_two_swapped_digits_are_rejected(self):
        """The reason it is Luhn and not a plain sum."""
        from library.cardnumbers import is_valid

        # Built rather than generated, so the two digits being swapped differ.
        from library.cardnumbers import _check_digit
        body = '712345'
        number = body + _check_digit(body)
        swapped = '712435' + number[-1]
        self.assertTrue(is_valid(number))
        self.assertFalse(is_valid(swapped))

    def test_punctuation_and_the_printed_prefix_do_not_change_the_number(self):
        from library.cardnumbers import _check_digit, is_valid, normalise

        number = '712345' + _check_digit('712345')
        for typed in (number, '  ' + number + ' ', '712-3455'[:3] + '-' + number[3:],
                      'AYLA-' + number):
            self.assertEqual(normalise(typed), number, typed)
            self.assertTrue(is_valid(typed), typed)

    def test_a_card_number_is_shown_as_seven_unbroken_digits(self):
        from library.cardnumbers import _check_digit, format_card

        number = '712345' + _check_digit('712345')
        self.assertEqual(format_card(number), number)
        self.assertNotIn('-', format_card(number))

    def test_every_member_is_given_one_and_visitors_are_not(self):
        member = Patron.objects.create(
            fullname='Card Holder', email='card-holder@example.invalid',
            password_hash=hash_password('x'), account_status='Active')
        visitor = Patron.objects.create(
            fullname='Walk In', password_hash=hash_password(None),
            account_status='Visitor')

        self.assertTrue(member.card_number)
        self.assertIsNone(visitor.card_number)

        # Made a member later: the number is minted then, not never.
        visitor.account_status = 'Active'
        visitor.save()
        self.assertTrue(visitor.card_number)

    def test_no_two_patrons_share_a_number(self):
        made = set()
        for index in range(25):
            patron = Patron.objects.create(
                fullname='Member %d' % index,
                email='member-%d@example.invalid' % index,
                password_hash=hash_password('x'), account_status='Active')
            self.assertNotIn(patron.card_number, made)
            made.add(patron.card_number)


class DeskKioskTests(TestCase):
    """The screen a patron stands in front of: identify, then state a direction."""

    def setUp(self):
        from .desk import DESK_SESSION_KEY

        self.user = _admin(modules='logs,patrons')
        self.client = _signed_in(self.user)
        session = self.client.session
        session[DESK_SESSION_KEY] = True
        session.save()

        self.patron = Patron.objects.create(
            fullname='Juan Dela Cruz', first_name='Juan', last_name='Dela Cruz',
            email='juan@example.invalid', password_hash=hash_password('x'),
            account_status='Active', qr_code='card-juan')

    def _post(self, path, **data):
        return self.client.post(path, data).json()

    def test_the_armed_desk_shows_the_kiosk_and_no_table(self):
        response = self.client.get('/admin-portal/log-management/')
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('I have a library card', body)
        # The kiosk must not show other visitors.
        self.assertNotIn('<table', body)
        self.assertNotIn(self.patron.fullname, body)

    def test_a_card_number_identifies_its_holder(self):
        r = self._post('/desk/identify/', value=self.patron.card_display)
        self.assertTrue(r['success'])
        self.assertEqual(r['name'], 'Juan Dela Cruz')
        self.assertFalse(r['inside'])
        self.assertEqual(r['suggest'], 'entry')

    def test_an_email_identifies_its_holder(self):
        r = self._post('/desk/identify/', value='JUAN@example.invalid')
        self.assertTrue(r['success'])
        self.assertEqual(r['name'], 'Juan Dela Cruz')

    def test_a_mistyped_card_number_says_so_rather_than_not_found(self):
        wrong = list(self.patron.card_number)
        wrong[2] = str((int(wrong[2]) + 1) % 10)
        r = self._post('/desk/identify/', value=''.join(wrong))
        self.assertFalse(r['success'])
        self.assertIn('not quite right', r['error'])

    def test_an_unknown_email_is_pointed_at_the_visitor_form(self):
        r = self._post('/desk/identify/', value='nobody@example.invalid')
        self.assertFalse(r['success'])
        self.assertIn('visitor', r['error'].lower())

    def test_a_suspended_card_is_sent_to_the_librarian_and_says_nothing_else(self):
        self.patron.account_status = 'Suspended'
        self.patron.save()
        r = self._post('/desk/identify/', value=self.patron.card_display)
        self.assertFalse(r['success'])
        self.assertIn('librarian', r['error'])
        self.assertNotIn('Juan', r['error'])

    def test_entry_then_exit_records_one_visit(self):
        self._post('/desk/identify/', value=self.patron.card_display)
        entry = self._post('/desk/visit/', direction='entry', purpose='Study')
        self.assertTrue(entry['success'])
        self.assertEqual(entry['action'], 'entry')

        log = PatronLog.objects.get(patron=self.patron)
        self.assertIsNone(log.exit_time)

        self._post('/desk/identify/', value=self.patron.card_display)
        leaving = self._post('/desk/visit/', direction='exit')
        self.assertTrue(leaving['success'])
        self.assertEqual(leaving['action'], 'exit')
        log.refresh_from_db()
        self.assertIsNotNone(log.exit_time)
        self.assertEqual(PatronLog.objects.filter(patron=self.patron).count(), 1)

    def test_exit_with_nothing_open_is_refused_rather_than_invented(self):
        self._post('/desk/identify/', value=self.patron.card_display)
        r = self._post('/desk/visit/', direction='exit')
        self.assertFalse(r['success'])
        self.assertEqual(r['action'], 'not_inside')
        self.assertEqual(PatronLog.objects.count(), 0)

    def test_a_direction_cannot_be_recorded_for_somebody_who_did_not_identify(self):
        """The identity lives in the session, never in the page."""
        r = self._post('/desk/visit/', direction='entry')
        self.assertFalse(r['success'])
        self.assertTrue(r.get('expired'))
        self.assertEqual(PatronLog.objects.count(), 0)

    def test_a_walk_in_is_logged_without_an_account(self):
        r = self._post('/desk/visitor/', direction='entry', name='Ana Reyes',
                       patron_type='General Visitor', purpose='Reading')
        self.assertTrue(r['success'])
        visitor = Patron.objects.get(fullname='Ana Reyes')
        self.assertEqual(visitor.account_status, 'Visitor')
        self.assertIsNone(visitor.qr_code)
        self.assertIsNone(visitor.card_number)
        self.assertEqual(PatronLog.objects.filter(patron=visitor).count(), 1)

    def test_a_returning_walk_in_is_the_same_person(self):
        self._post('/desk/visitor/', direction='entry', name='Ana Reyes',
                   purpose='Reading')
        PatronLog.objects.update(exit_time=timezone.now())
        self._post('/desk/visitor/', direction='entry', name='ana reyes',
                   purpose='Study')
        self.assertEqual(Patron.objects.filter(fullname__icontains='Ana').count(), 1)

    def test_a_walk_in_cannot_sign_out_on_a_name_alone(self):
        self._post('/desk/visitor/', direction='entry', name='Ana Reyes',
                   contact='ana@example.invalid', purpose='Reading')
        bare = self._post('/desk/visitor/', direction='exit', name='Ana Reyes')
        self.assertFalse(bare['success'])
        self.assertIsNone(PatronLog.objects.get().exit_time)

        with_contact = self._post('/desk/visitor/', direction='exit',
                                  name='Ana Reyes', contact='ana@example.invalid')
        self.assertTrue(with_contact['success'])
        self.assertIsNotNone(PatronLog.objects.get().exit_time)

    def test_a_stranger_cannot_reach_the_kiosk_endpoints(self):
        stranger = Client()
        for path in ('/desk/identify/', '/desk/visit/', '/desk/visitor/'):
            r = stranger.post(path, {'value': 'x', 'direction': 'entry', 'name': 'x'})
            self.assertFalse(r.json()['success'], path)
        self.assertEqual(PatronLog.objects.count(), 0)


class ElementPropertiesTests(TestCase):
    """Everything drawn on a floor plan can be described after it is drawn."""

    def setUp(self):
        self.user = _admin(modules='shelf')
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True)
        self.upstairs = FloorPlan.objects.create(name='First', floor_number=2)
        self.room = Room.objects.create(
            floor_plan=self.plan, name='Reading Room', map_x=50, map_y=50,
            geometry=[[0, 0], [100, 0], [100, 100], [0, 100]])
        self.shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                          map_x=50, map_y=50)

    def _post(self, path, **data):
        return self.client.post(path, data).json()

    # Shelves

    def test_a_shelf_can_be_told_what_it_really_is(self):
        r = self._post('/admin-portal/edit-shelf/', shelf_id=self.shelf.shelf_id,
                       kind='Display', mount='Wall', mount_height_m='1.4')
        self.assertTrue(r['success'], r)
        self.shelf.refresh_from_db()
        self.assertEqual(self.shelf.kind, 'Display')
        self.assertEqual(self.shelf.mount, 'Wall')
        self.assertEqual(self.shelf.mount_height_m, 1.4)

    def test_a_shelf_type_the_model_does_not_know_is_refused(self):
        r = self._post('/admin-portal/edit-shelf/', shelf_id=self.shelf.shelf_id,
                       kind='Hammock')
        self.assertFalse(r['success'])
        self.shelf.refresh_from_db()
        self.assertEqual(self.shelf.kind, 'Shelf')

    def test_an_impossible_mounting_height_is_refused(self):
        r = self._post('/admin-portal/edit-shelf/', shelf_id=self.shelf.shelf_id,
                       mount_height_m='40')
        self.assertFalse(r['success'])
        self.shelf.refresh_from_db()
        self.assertIsNone(self.shelf.mount_height_m)

    def test_a_blank_height_clears_it_rather_than_failing(self):
        self.shelf.mount_height_m = 1.2
        self.shelf.save()
        r = self._post('/admin-portal/edit-shelf/', shelf_id=self.shelf.shelf_id,
                       mount_height_m='')
        self.assertTrue(r['success'], r)
        self.shelf.refresh_from_db()
        self.assertIsNone(self.shelf.mount_height_m)

    def test_a_shelf_can_be_hidden_without_being_unplaced(self):
        r = self._post('/admin-portal/edit-shelf/', shelf_id=self.shelf.shelf_id,
                       is_active='0')
        self.assertTrue(r['success'], r)
        self.shelf.refresh_from_db()
        self.assertFalse(self.shelf.is_active)
        self.assertEqual(self.shelf.map_x, 50)

    # Rooms

    def test_a_room_can_be_hidden_and_kept(self):
        r = self._post('/admin-portal/edit-room/', room_id=self.room.room_id,
                       is_active='0', description='Closed for repairs')
        self.assertTrue(r['success'], r)
        self.room.refresh_from_db()
        self.assertFalse(self.room.is_active)
        self.assertEqual(self.room.description, 'Closed for repairs')
        self.assertTrue(Shelf.objects.filter(shelf_id=self.shelf.shelf_id).exists())

    def test_a_room_edit_that_says_nothing_about_visibility_leaves_it_alone(self):
        """Absent is "not on this form", never "unticked"."""
        self._post('/admin-portal/edit-room/', room_id=self.room.room_id,
                   name='Quiet Room')
        self.room.refresh_from_db()
        self.assertTrue(self.room.is_active)
        self.assertTrue(self.room.patron_access)

    # Doors

    def test_a_door_can_be_named_and_closed_off(self):
        door = Door.objects.create(room=self.room, map_x=50, map_y=0, width=28)
        r = self._post('/admin-portal/edit-door/', door_id=door.door_id,
                       label='Fire exit', is_active='0')
        self.assertTrue(r['success'], r)
        door.refresh_from_db()
        self.assertEqual(door.label, 'Fire exit')
        self.assertFalse(door.is_active)

    # Waypoints

    def test_a_waypoint_can_be_pointed_at_a_shelf_afterwards(self):
        wp = Waypoint.objects.create(floor_plan=self.plan, map_x=10, map_y=10)
        r = self._post('/admin-portal/edit-waypoint/', waypoint_id=wp.waypoint_id,
                       label='Aisle 2', linked_shelf_id=self.shelf.shelf_id)
        self.assertTrue(r['success'], r)
        wp.refresh_from_db()
        self.assertEqual(wp.label, 'Aisle 2')
        self.assertEqual(wp.linked_shelf, self.shelf)

    def test_a_waypoint_link_can_be_cleared(self):
        wp = Waypoint.objects.create(floor_plan=self.plan, map_x=10, map_y=10,
                                     linked_shelf=self.shelf)
        r = self._post('/admin-portal/edit-waypoint/', waypoint_id=wp.waypoint_id,
                       linked_shelf_id='')
        self.assertTrue(r['success'], r)
        wp.refresh_from_db()
        self.assertIsNone(wp.linked_shelf)

    def test_a_waypoint_cannot_stand_in_front_of_a_shelf_upstairs(self):
        other_room = Room.objects.create(floor_plan=self.upstairs, name='Loft',
                                         map_x=10, map_y=10)
        far = Shelf.objects.create(room=other_room, name='Shelf Z',
                                   map_x=10, map_y=10)
        wp = Waypoint.objects.create(floor_plan=self.plan, map_x=10, map_y=10)
        r = self._post('/admin-portal/edit-waypoint/', waypoint_id=wp.waypoint_id,
                       linked_shelf_id=far.shelf_id)
        self.assertFalse(r['success'])
        wp.refresh_from_db()
        self.assertIsNone(wp.linked_shelf)

    # Stairways

    def test_a_stairway_can_be_reshaped_into_a_switchback(self):
        st = Stairway.objects.create(
            floor_plan=self.plan, kind='Stairs', direction='up', bearing=0,
            geometry=[[0, 0], [60, 0], [60, 200], [0, 200]], map_x=30, map_y=100)
        r = self._post('/admin-portal/edit-stairway/', stairway_id=st.stairway_id,
                       shape='half', bearing=0)
        self.assertTrue(r['success'], r)
        st.refresh_from_db()
        # A switchback is more than one run; a straight flight is one or none.
        self.assertGreater(len(st.flights or []), 1)

    def test_a_stairway_can_be_closed_off_without_being_deleted(self):
        st = Stairway.objects.create(
            floor_plan=self.plan, kind='Elevator', direction='both', bearing=0,
            geometry=[[0, 0], [40, 0], [40, 40], [0, 40]], map_x=20, map_y=20)
        r = self._post('/admin-portal/edit-stairway/', stairway_id=st.stairway_id,
                       is_active='0')
        self.assertTrue(r['success'], r)
        st.refresh_from_db()
        self.assertFalse(st.is_active)
        self.assertTrue(Stairway.objects.filter(stairway_id=st.stairway_id).exists())


class CanvasResizeTests(TestCase):
    """The canvas a plan is drawn on, changed after it has been drawn on."""

    def setUp(self):
        self.user = _admin()
        self.client = _signed_in(self.user)
        self.plan = FloorPlan.objects.create(name='Ground', floor_number=1,
                                             is_active=True, canvas_width=1000,
                                             canvas_height=800, pixels_per_meter=50)
        self.room = Room.objects.create(
            floor_plan=self.plan, name='Reading Room', map_x=500, map_y=400,
            geometry=[[400, 300], [600, 300], [600, 500], [400, 500]])
        self.shelf = Shelf.objects.create(room=self.room, name='Shelf A',
                                          map_x=500, map_y=400, width=40, depth=12)
        self.beacon = BLEBeacon.objects.create(
            floor_plan=self.plan, beacon_uuid=str(uuid4()), map_x=900, map_y=700)

    def _resize(self, width, height, scale=False):
        return self.client.post('/admin-portal/set-floorplan-canvas/', {
            'floorplan_id': self.plan.floor_plan_id,
            'canvas_width': width,
            'canvas_height': height,
            'scale_contents': '1' if scale else '0',
        }).json()

    def test_growing_the_canvas_leaves_everything_where_it_was(self):
        r = self._resize(1600, 1200)
        self.assertTrue(r['success'], r)
        self.plan.refresh_from_db()
        self.room.refresh_from_db()
        self.assertEqual(self.plan.canvas_width, 1600)
        self.assertEqual(self.room.geometry[0], [400, 300])
        self.assertEqual(self.plan.pixels_per_meter, 50)

    def test_shrinking_is_refused_when_it_would_strand_something(self):
        r = self._resize(500, 500)
        self.assertFalse(r['success'])
        # The refusal has to say what, or there is nothing to act on.
        self.assertIn('Reading Room', r['error'])
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.canvas_width, 1000)

    def test_shrinking_into_empty_space_is_allowed(self):
        BLEBeacon.objects.all().delete()
        r = self._resize(700, 600)
        self.assertTrue(r['success'], r)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.canvas_width, 700)

    def test_resizing_everything_to_fit_keeps_its_proportions(self):
        r = self._resize(500, 400, scale=True)
        self.assertTrue(r['success'], r)
        self.room.refresh_from_db()
        self.shelf.refresh_from_db()
        self.beacon.refresh_from_db()
        self.assertEqual(self.room.geometry[0], [200, 150])
        self.assertEqual(self.room.map_x, 250)
        # Footprints are distances too, or a shelf ends up inside the wall.
        self.assertEqual(self.shelf.width, 20)
        self.assertEqual(self.shelf.depth, 6)
        self.assertEqual((self.beacon.map_x, self.beacon.map_y), (450, 350))

    def test_the_building_is_still_the_same_size_in_metres_afterwards(self):
        """Halving the drawing without halving the scale moves every beacon fix."""
        self._resize(500, 400, scale=True)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.pixels_per_meter, 25)
        self.room.refresh_from_db()
        width_in_metres = (self.room.geometry[1][0] - self.room.geometry[0][0]) / 25
        self.assertEqual(width_in_metres, 4)     # 200 units at 50/m before, 100 at 25/m now

    def test_scaling_uses_one_factor_so_nothing_is_stretched(self):
        """A square room stays square even when the canvas changes shape."""
        self._resize(2000, 900, scale=True)      # x doubles, y is 1.125
        self.room.refresh_from_db()
        width = self.room.geometry[1][0] - self.room.geometry[0][0]
        height = self.room.geometry[2][1] - self.room.geometry[1][1]
        self.assertEqual(width, height)

    def test_an_absurd_canvas_is_refused(self):
        for width, height in ((10, 800), (1000, 99999)):
            self.assertFalse(self._resize(width, height)['success'])
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.canvas_width, 1000)

    def test_a_stranger_cannot_resize_a_canvas(self):
        r = Client().post('/admin-portal/set-floorplan-canvas/', {
            'floorplan_id': self.plan.floor_plan_id,
            'canvas_width': 4000, 'canvas_height': 4000})
        self.assertEqual(r.status_code, 302)
        self.plan.refresh_from_db()
        self.assertEqual(self.plan.canvas_width, 1000)


class ScannedCodeResolutionTests(TestCase):
    """One scanner at the desk, so one lookup has to say what was scanned."""

    def setUp(self):
        self.user = _admin(modules='transactions')
        self.client = _signed_in(self.user)
        self.patron = Patron.objects.create(
            fullname='Juan Dela Cruz', first_name='Juan', last_name='Dela Cruz',
            email='juan-scan@example.invalid', password_hash=hash_password('x'),
            account_status='Active', qr_code=str(uuid4()))
        self.book = Book.objects.create(title='Noli Me Tangere', author='Rizal',
                                        status='Available', qr_code=str(uuid4()))

    def _resolve(self, code):
        return self.client.get('/admin-portal/resolve-qr/',
                               {'qr_code': code}).json()

    def test_a_card_comes_back_as_a_patron(self):
        r = self._resolve(self.patron.qr_code)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['kind'], 'patron')
        self.assertEqual(r['patron']['fullname'], 'Juan Dela Cruz')
        self.assertIn('eligible', r['patron'])

    def test_a_book_label_comes_back_as_a_book(self):
        r = self._resolve(self.book.qr_code)
        self.assertTrue(r['success'], r)
        self.assertEqual(r['kind'], 'book')
        self.assertEqual(r['book']['title'], 'Noli Me Tangere')

    def test_a_code_that_is_neither_says_so(self):
        r = self._resolve(str(uuid4()))
        self.assertFalse(r['success'])
        self.assertEqual(r['kind'], 'unknown')
        # The error names both expected types.
        self.assertIn('card', r['error'])
        self.assertIn('label', r['error'])

    def test_an_empty_scan_is_refused_without_a_query(self):
        r = self._resolve('   ')
        self.assertFalse(r['success'])
        self.assertEqual(r['kind'], 'unknown')

    def test_the_payloads_match_the_single_kind_endpoints(self):
        """One helper each, so a field cannot be added to one scanner only."""
        unified = self._resolve(self.book.qr_code)['book']
        single = self.client.get('/admin-portal/search-book-by-qr/',
                                 {'qr_code': self.book.qr_code}).json()['book']
        self.assertEqual(unified, single)

        unified = self._resolve(self.patron.qr_code)['patron']
        single = self.client.get('/admin-portal/search-patron-by-qr/',
                                 {'qr_code': self.patron.qr_code}).json()['patron']
        self.assertEqual(unified, single)

    def test_a_stranger_cannot_look_codes_up(self):
        """The card payload names a patron and says whether they may borrow."""
        r = Client().get('/admin-portal/resolve-qr/',
                         {'qr_code': self.patron.qr_code})
        self.assertEqual(r.status_code, 302)


class PaginationTests(TestCase):
    """Table footers show a handful of page links, never every page there is."""

    def setUp(self):
        self.user = _admin(modules='books')
        self.client = _signed_in(self.user)
        # 15 per page gives 14 pages.
        Book.objects.bulk_create([
            Book(title='Book %03d' % i, author='Author', status='Available')
            for i in range(200)
        ])

    def test_the_tag_keeps_the_ends_and_the_neighbours(self):
        from django.core.paginator import Paginator
        from library.templatetags.pagination import elided_pages

        page = Paginator(range(1000), 10).get_page(50)
        pages = elided_pages(page)
        self.assertEqual(pages[0], 1)
        self.assertEqual(pages[-1], 100)
        self.assertIn(Paginator.ELLIPSIS, pages)
        for near in (48, 49, 50, 51, 52):
            self.assertIn(near, pages)
        self.assertLess(len(pages), 12)

    def test_a_middle_page_does_not_list_every_page(self):
        body = self.client.get('/admin-portal/management/?page=7').content.decode()
        self.assertIn('aria-current="page">7<', body)
        self.assertIn('&hellip;', body)
        # Page 12 is outside the window around 7 and not an end, so it has no link.
        self.assertNotIn('?page=12"', body)
        self.assertIn('?page=14', body)

    def test_moving_page_keeps_the_filters(self):
        body = self.client.get('/admin-portal/management/?page=2&q=Book').content.decode()
        self.assertIn('href="?page=3&q=Book"', body)

    def test_a_single_page_renders_no_pager(self):
        Book.objects.filter(title__gt='Book 010').delete()
        body = self.client.get('/admin-portal/management/').content.decode()
        self.assertNotIn('class="ayla-pager', body)


class PatronCatalogTests(TestCase):
    """The catalogue a patron browses: titles, real filters, nothing hidden that exists."""

    def setUp(self):
        self.client = Client()
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Reading Room', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=1, map_y=1)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=1)

    def _book(self, title, **kw):
        kw.setdefault('author', 'Author')
        kw.setdefault('status', 'Available')
        return Book.objects.create(title=title, **kw)

    def _get(self, **params):
        return self.client.get('/patron/catalog/', params)

    def _titles(self, response):
        return [row['title'] for row in response.context['page']]

    def test_an_unshelved_book_is_listed(self):
        self._book('Waiting For A Shelf')
        r = self._get()
        self.assertIn('Waiting For A Shelf', self._titles(r))
        self.assertContains(r, 'Not yet shelved')

    def test_a_book_on_loan_is_listed_rather_than_hidden(self):
        self._book('Out Right Now', status='Borrowed', shelf_level=self.level)
        self.assertIn('Out Right Now', self._titles(self._get()))

    def test_written_off_copies_are_not_listed(self):
        self._book('Gone For Good', status='Lost')
        self._book('Given Away', status='Donated')
        titles = self._titles(self._get())
        self.assertNotIn('Gone For Good', titles)
        self.assertNotIn('Given Away', titles)

    def test_copies_of_one_title_are_one_entry(self):
        for status in ('Available', 'Available', 'Borrowed'):
            self._book('Noli Me Tangere', author='Rizal', status=status,
                       shelf_level=self.level)
        rows = list(self._get().context['page'])
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]['copies'], rows[0]['available']), (3, 2))

    def test_the_card_opens_a_copy_that_is_in_and_shelved(self):
        self._book('Two Copies', status='Borrowed', shelf_level=self.level)
        unshelved = self._book('Two Copies', status='Available')
        shelved = self._book('Two Copies', status='Available', shelf_level=self.level)
        row = list(self._get().context['page'])[0]
        self.assertEqual(row['book'].book_id, shelved.book_id)
        self.assertNotEqual(row['book'].book_id, unshelved.book_id)

    def test_filters_narrow_by_subject_availability_and_location(self):
        self._book('Algebra', genre='Mathematics', shelf_level=self.level)
        self._book('Poems', genre='Poetry', status='Borrowed')
        self._book('Loose', genre='Mathematics')

        self.assertEqual(sorted(self._titles(self._get(genre='Mathematics'))),
                         ['Algebra', 'Loose'])
        self.assertEqual(self._titles(self._get(availability='on_loan')), ['Poems'])
        self.assertEqual(sorted(self._titles(self._get(availability='available'))),
                         ['Algebra', 'Loose'])
        self.assertEqual(self._titles(self._get(shelf=self.shelf.shelf_id)), ['Algebra'])
        self.assertEqual(sorted(self._titles(self._get(shelf='none'))), ['Loose', 'Poems'])

    def test_search_still_finds_by_author(self):
        self._book('Florante at Laura', author='Balagtas')
        self._book('Something Else', author='Nobody')
        self.assertEqual(self._titles(self._get(search='balagtas')), ['Florante at Laura'])

    def test_sorting_is_applied_on_the_server(self):
        self._book('Beta', publication_year=1990)
        self._book('Alpha', publication_year=2020)
        self.assertEqual(self._titles(self._get(sort='newest')), ['Alpha', 'Beta'])
        self.assertEqual(self._titles(self._get(sort='title_desc')), ['Beta', 'Alpha'])

    def test_an_unknown_filter_value_is_ignored_not_an_error(self):
        self._book('Still Here')
        r = self._get(shelf='upstairs', availability='maybe', sort='sideways')
        self.assertEqual(r.status_code, 200)
        self.assertIn('Still Here', self._titles(r))

    def test_the_list_is_paginated_and_keeps_filters(self):
        for i in range(45):
            self._book('Book %02d' % i, genre='Fiction')
        r = self._get(genre='Fiction', page=2)
        self.assertEqual(len(self._titles(r)), 20)
        self.assertContains(r, 'href="?page=3&genre=Fiction"')

    def test_each_active_filter_can_be_removed_on_its_own(self):
        self._book('Algebra', genre='Mathematics', shelf_level=self.level)
        r = self._get(genre='Mathematics', availability='available')
        removes = {f['label']: f['remove'] for f in r.context['active_filters']}
        self.assertEqual(removes['Mathematics'], '?availability=available')
        self.assertEqual(removes['Available now'], '?genre=Mathematics')

    def test_an_unshelved_book_offers_the_desk_instead_of_the_map(self):
        book = self._book('Waiting For A Shelf')
        r = self.client.get('/patron/book-details/%d/' % book.book_id)
        self.assertNotContains(r, 'Navigate on Map')
        self.assertContains(r, 'front desk')


class InventoryPaginationTests(TestCase):
    """Inventory pages three tables on one screen, each under its own parameter."""

    def setUp(self):
        from django.urls import reverse
        from library import views

        self.url = reverse(views.inventory_management)
        self.client = _signed_in(_admin())

    def test_the_stock_table_keeps_its_filters_across_pages(self):
        from library.models import InventoryRecord

        InventoryRecord.objects.bulk_create(
            [InventoryRecord(title_hint='Box %02d' % i) for i in range(45)])
        body = self.client.get(self.url, {'q': 'Box', 'page': 2}).content.decode()
        self.assertIn('class="ayla-pager', body)
        # Encoded by the view.
        self.assertIn('href="?page=3&q=Box"', body)

    def test_the_history_table_pages_under_its_own_parameter_and_tab(self):
        from library.models import InventoryRecord, StockMovement

        record = InventoryRecord.objects.create(title_hint='Box')
        StockMovement.objects.bulk_create(
            [StockMovement(inventory_record=record, action='Received') for _ in range(30)])
        body = self.client.get(self.url, {'tab': 'history'}).content.decode()
        self.assertIn('href="?mpage=2&tab=history"', body)
        self.assertNotIn('href="?page=2&tab=history"', body)

    def test_a_search_with_an_ampersand_survives_the_page_link(self):
        from library.models import InventoryRecord

        InventoryRecord.objects.bulk_create(
            [InventoryRecord(title_hint='Pens & Paper %02d' % i) for i in range(25)])
        body = self.client.get(self.url, {'q': 'Pens & Paper'}).content.decode()
        self.assertIn('q=Pens+%26+Paper', body)


class MapBookSearchTests(TestCase):
    """The map's Find a book sheet runs the catalogue's own search."""

    def setUp(self):
        self.client = Client()
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True)
        room = Room.objects.create(floor_plan=plan, name='Reading Room', map_x=0, map_y=0)
        self.placed = Shelf.objects.create(room=room, name='Shelf A', map_x=10, map_y=10)
        self.unplaced = Shelf.objects.create(room=room, name='Shelf Z')
        self.level = ShelfLevel.objects.create(shelf=self.placed, level_number=1)
        self.hidden_level = ShelfLevel.objects.create(shelf=self.unplaced, level_number=1)

    def _book(self, title, **kw):
        kw.setdefault('author', 'Author')
        kw.setdefault('status', 'Available')
        return Book.objects.create(title=title, **kw)

    def _search(self, **params):
        return self.client.get('/patron/catalog/search/', params).json()

    def test_results_come_back_as_titles_with_a_copy_to_open(self):
        for _ in range(2):
            self._book('Noli Me Tangere', author='Rizal', shelf_level=self.level)
        r = self._search(search='noli')
        self.assertTrue(r['success'])
        self.assertEqual(r['count'], 1)
        hit = r['results'][0]
        self.assertEqual((hit['copies'], hit['available']), (2, 2))
        self.assertTrue(hit['navigable'])
        self.assertIn('Shelf A', hit['location'])

    def test_the_filters_mean_what_they_mean_in_the_catalogue(self):
        self._book('Algebra', genre='Mathematics', shelf_level=self.level)
        self._book('Poems', genre='Poetry', status='Borrowed', shelf_level=self.level)
        self._book('Loose', genre='Mathematics')

        def titles(**p):
            return sorted(x['title'] for x in self._search(**p)['results'])

        self.assertEqual(titles(genre='Mathematics'), ['Algebra', 'Loose'])
        self.assertEqual(titles(availability='on_loan'), ['Poems'])
        self.assertEqual(titles(shelf='none'), ['Loose'])
        self.assertEqual(titles(shelf=self.placed.shelf_id), ['Algebra', 'Poems'])

    def test_an_unshelved_book_is_listed_but_not_offered_a_route(self):
        self._book('Waiting For A Shelf')
        hit = self._search()['results'][0]
        self.assertFalse(hit['shelved'])
        self.assertFalse(hit['navigable'])

    def test_a_shelf_with_no_place_on_the_plan_is_not_offered_a_route(self):
        self._book('Somewhere Unmapped', shelf_level=self.hidden_level)
        hit = self._search()['results'][0]
        self.assertTrue(hit['shelved'])
        self.assertFalse(hit['navigable'])

    def test_written_off_copies_are_not_found(self):
        self._book('Gone', status='Lost')
        self.assertEqual(self._search(search='Gone')['count'], 0)

    def test_filter_choices_come_only_when_asked_for(self):
        self._book('Algebra', genre='Mathematics', shelf_level=self.level)
        self._book('Loose')
        self.assertNotIn('options', self._search())
        options = self._search(options=1)['options']
        self.assertEqual(options['genres'], ['Mathematics'])
        self.assertEqual([s['name'] for s in options['shelves']], ['Shelf A'])
        self.assertTrue(options['has_unshelved'])

    def test_results_page_for_show_more(self):
        for i in range(15):
            self._book('Book %02d' % i, shelf_level=self.level)
        first = self._search()
        self.assertEqual(len(first['results']), 12)
        self.assertTrue(first['has_next'])
        second = self._search(page=2)
        self.assertEqual(len(second['results']), 3)
        self.assertFalse(second['has_next'])

    def test_the_map_offers_the_search_and_the_guide(self):
        body = self.client.get('/patron/map/').content.decode()
        self.assertIn('Find a book', body)
        self.assertIn('/patron/help/live-position/', body)


class LivePositionGuideTests(TestCase):
    def test_the_guide_is_public_and_carries_the_flag_address(self):
        r = Client().get('/patron/help/live-position/')
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'chrome://flags/#enable-experimental-web-platform-features')
        self.assertContains(r, 'Experimental Web Platform features')


class PortalMapBookSearchTests(TestCase):
    """The staff and Administrator indoor maps search with the catalogue's filters."""

    def setUp(self):
        plan = FloorPlan.objects.create(name='Ground', floor_number=1, is_active=True,
                                        canvas_width=1000, canvas_height=800)
        room = Room.objects.create(floor_plan=plan, name='Reading Room', map_x=0, map_y=0)
        self.shelf = Shelf.objects.create(room=room, name='Shelf A', map_x=10, map_y=10)
        self.level = ShelfLevel.objects.create(shelf=self.shelf, level_number=1)

    def test_a_result_carries_the_shelf_the_portal_maps_highlight(self):
        Book.objects.create(title='Algebra', author='Author', status='Available',
                            shelf_level=self.level)
        hit = Client().get('/patron/catalog/search/', {'search': 'Algebra'}).json()['results'][0]
        self.assertEqual(hit['shelf_id'], self.shelf.shelf_id)

    def test_an_unshelved_result_has_no_shelf_to_highlight(self):
        Book.objects.create(title='Loose', author='Author', status='Available')
        hit = Client().get('/patron/catalog/search/').json()['results'][0]
        self.assertIsNone(hit['shelf_id'])

    def test_the_staff_map_offers_the_search(self):
        staff = User.objects.create(
            fullname='Desk Staff', email='desk-staff@example.invalid',
            password_hash=hash_password('SmokeTest123'), role='Staff',
            account_status='Active', modules='indoor_map')
        r = _signed_in(staff).get('/library-staff/indoor-map/')
        self.assertEqual(r.status_code, 200)
        body = r.content.decode()
        self.assertIn('id="findPanel"', body)
        self.assertIn('/patron/catalog/search/', body)

    def test_the_admin_map_offers_the_search(self):
        r = _signed_in(_admin()).get('/admin-portal/indoor-map/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('id="findPanel"', r.content.decode())
